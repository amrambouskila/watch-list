"""Regression tests for constructs a naive row delete or table resize would destroy.

Each of these reproduces a corruption path found by review: openpyxl's `delete_rows` moves cell
values but leaves merges and hyperlinks anchored to their original absolute rows, and a table
resize that assumes one full-width table will damage a second table or absorb stray cells.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from library_choice import a_category_whose_order_column_holds_numbers, a_category_with_at_least
from tv_watchlist.constants import FIRST_DATA_ROW
from tv_watchlist.workbook.reader import read_category
from tv_watchlist.workbook.writer import append_row, delete_row, update_row
from workbook_facts import WorkbookFacts

# Deep enough that the merge, the links and the deleted row are all well clear of each other.
ROWS_A_HAZARD_NEEDS = 12
MERGED_ROW = 10
ROW_BELOW_THE_MERGE = 11
DELETED_ROW = 5
LINKED_ROW = 10
FIRST_LINK = "https://example.com/ten"
SECOND_LINK = "https://example.com/further-down"
A_TITLE_THE_LIBRARY_DOES_NOT_HOLD = "Bonus feature the owner never listed"
TITLE_COLUMN_INDEX = 2
SIDE_TABLE_NAME = "SideLookup"
A_YEAR_THAT_LOOKS_NUMERIC = "2024"


def _spare_columns(facts: WorkbookFacts) -> tuple[str, str]:
    """Two adjacent column letters past the end of the grid, with a gap so neither touches it."""
    first = facts.column_count + 2
    return get_column_letter(first), get_column_letter(first + 1)


def _table_end(path: Path, name: str) -> int:
    return range_boundaries(load_workbook(path).worksheets[0].tables[name].ref)[3]


def test_deleting_a_row_does_not_feed_values_into_a_merged_range(library: Path) -> None:
    facts = a_category_with_at_least(rows=ROWS_A_HAZARD_NEEDS)
    path = facts.copied_into(library)
    left, right = _spare_columns(facts)
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    sheet.merge_cells(f"{left}{MERGED_ROW}:{right}{MERGED_ROW}")
    sheet[f"{left}{ROW_BELOW_THE_MERGE}"] = "keep-left"
    sheet[f"{right}{ROW_BELOW_THE_MERGE}"] = "keep-right"
    workbook.save(path)

    delete_row(path, DELETED_ROW)

    sheet = load_workbook(path).worksheets[0]
    assert sheet[f"{left}{MERGED_ROW}"].value == "keep-left"
    assert sheet[f"{right}{MERGED_ROW}"].value == "keep-right"
    lifted = f"{left}{MERGED_ROW - 1}:{right}{MERGED_ROW - 1}"
    assert lifted in {str(cell_range) for cell_range in sheet.merged_cells.ranges}


def test_deleting_a_row_keeps_hyperlinks_on_their_own_titles(library: Path) -> None:
    facts = a_category_with_at_least(rows=ROWS_A_HAZARD_NEEDS)
    path = facts.copied_into(library)
    far_row = facts.last_data_row
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    before = {row: sheet.cell(row=row, column=TITLE_COLUMN_INDEX).value for row in (LINKED_ROW, far_row)}
    sheet.cell(row=LINKED_ROW, column=TITLE_COLUMN_INDEX).hyperlink = FIRST_LINK
    sheet.cell(row=far_row, column=TITLE_COLUMN_INDEX).hyperlink = SECOND_LINK
    workbook.save(path)

    delete_row(path, DELETED_ROW)

    sheet = load_workbook(path).worksheets[0]
    linked = {
        sheet.cell(row=row, column=TITLE_COLUMN_INDEX).value: sheet.cell(row=row, column=TITLE_COLUMN_INDEX).hyperlink
        for row in range(FIRST_DATA_ROW, sheet.max_row + 1)
        if sheet.cell(row=row, column=TITLE_COLUMN_INDEX).hyperlink is not None
    }
    assert {title: link.target for title, link in linked.items()} == {
        before[LINKED_ROW]: FIRST_LINK,
        before[far_row]: SECOND_LINK,
    }


def test_a_second_table_on_the_sheet_keeps_its_own_range(subject: WorkbookFacts, subject_path: Path) -> None:
    from openpyxl.worksheet.table import Table, TableColumn

    left, right = _spare_columns(subject)
    reference = f"{left}1:{right}2"
    workbook = load_workbook(subject_path)
    sheet = workbook.worksheets[0]
    sheet[f"{left}1"], sheet[f"{right}1"] = "Key", "Value"
    sheet[f"{left}2"], sheet[f"{right}2"] = "a", "1"
    lookup = Table(displayName=SIDE_TABLE_NAME, name=SIDE_TABLE_NAME, ref=reference)
    lookup.tableColumns = [TableColumn(id=1, name="Key"), TableColumn(id=2, name="Value")]
    sheet.add_table(lookup)
    workbook.save(subject_path)

    append_row(subject_path, {subject.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    sheet = load_workbook(subject_path).worksheets[0]
    assert sheet.tables[SIDE_TABLE_NAME].ref == reference
    assert _table_end(subject_path, subject.grid_table_name) == subject.last_data_row + 1


def test_a_stray_cell_beside_the_grid_never_becomes_a_table_column(subject: WorkbookFacts, subject_path: Path) -> None:
    stray, _ = _spare_columns(subject)
    workbook = load_workbook(subject_path)
    workbook.worksheets[0][f"{stray}4"] = "a note to myself"
    workbook.save(subject_path)

    append_row(subject_path, {subject.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    table = load_workbook(subject_path).worksheets[0].tables[subject.grid_table_name]
    assert range_boundaries(table.ref)[3] == subject.last_data_row + 1
    assert "" not in [column.name for column in table.tableColumns]


def test_filling_a_blank_cell_in_a_numeric_column_still_writes_a_number(library: Path) -> None:
    facts = a_category_whose_order_column_holds_numbers()
    path = facts.copied_into(library)
    order_index = facts.index_of(facts.order_key)
    blanked = FIRST_DATA_ROW + 1
    workbook = load_workbook(path)
    workbook.worksheets[0].cell(row=blanked, column=order_index).value = None
    workbook.save(path)

    update_row(path, blanked, {facts.order_key: "2"})

    assert load_workbook(path).worksheets[0].cell(row=blanked, column=order_index).value == 2


def test_free_text_columns_are_never_coerced_to_numbers(subject: WorkbookFacts, subject_path: Path) -> None:
    update_row(subject_path, FIRST_DATA_ROW, {subject.free_text_key: A_YEAR_THAT_LOOKS_NUMERIC})

    written = (
        load_workbook(subject_path)
        .worksheets[0]
        .cell(row=FIRST_DATA_ROW, column=subject.index_of(subject.free_text_key))
    )
    assert written.value == A_YEAR_THAT_LOOKS_NUMERIC
    assert read_category(subject_path).rows[0].cells[subject.free_text_key] == A_YEAR_THAT_LOOKS_NUMERIC


def test_the_first_row_of_an_empty_sheet_does_not_inherit_the_header_fill(tmp_path: Path) -> None:
    path = tmp_path / "Fresh_Master_Watch_Order.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for index, header in enumerate(["Order", "Title"], start=1):
        cell = sheet.cell(row=1, column=index)
        cell.value = header
        cell.fill = PatternFill(start_color="FF123456", end_color="FF123456", fill_type="solid")
    workbook.save(path)

    row = append_row(path, {"title": "First entry"})

    sheet = load_workbook(path).worksheets[0]
    assert sheet.cell(row=row, column=2).fill.patternType is None
    assert sheet.cell(row=row, column=2).alignment.wrap_text is True
