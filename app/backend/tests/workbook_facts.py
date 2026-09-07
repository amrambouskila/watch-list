"""The structural facts of one hand-built workbook, so tests can pick a subject by shape."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import FIRST_DATA_ROW, HEADER_ROW
from tv_watchlist.models.category_detail import CategoryDetail
from tv_watchlist.workbook.reader import read_category

GRID_ANCHOR_COLUMN = 1


@dataclass(frozen=True)
class WorkbookFacts:
    """What a workbook is, stated in properties its owner cannot change by renaming the file."""

    path: Path
    category_id: str
    row_count: int
    marked_row_count: int
    watched_row_count: int
    unwatched_row_count: int
    column_count: int
    order_label: str
    order_key: str
    order_holds_numbers: bool
    title_key: str
    watch_key: str
    free_text_keys: tuple[str, ...]
    dropdown_keys_beyond_the_watch_column: tuple[str, ...]
    reference_sheet_count: int
    range_ends: tuple[int, ...]
    grid_table_name: str
    column_indexes: tuple[tuple[str, int], ...]
    titles: tuple[str, ...]

    def index_of(self, key: str) -> int:
        """Which sheet column a derived column key sits in."""
        return dict(self.column_indexes)[key]

    @property
    def file_name(self) -> str:
        """The name on disk, only ever used to reach the same workbook in a disposable copy."""
        return self.path.name

    @property
    def stem(self) -> str:
        return self.path.stem

    @property
    def mtime(self) -> float:
        return self.path.stat().st_mtime

    @property
    def shadowing_name(self) -> str:
        """
        A display name the filesystem files separately but the app addresses identically.

        A trailing hyphen survives into the filename and drops out of the id, so this holds whatever
        the owner has called their workbooks.
        """
        return f"{self.category_id}-"

    @property
    def free_text_key(self) -> str:
        """A column a test can write arbitrary prose into without a dropdown refusing it."""
        return self.free_text_keys[0]

    @property
    def last_data_row(self) -> int:
        return FIRST_DATA_ROW + self.row_count - 1

    @property
    def ranges_all_end_together(self) -> bool:
        """True when table, dropdown and highlight all stop at the same row, as a clean sheet does."""
        return len(self.range_ends) == 1

    @property
    def is_unmarked(self) -> bool:
        """True when no row carries a watch status yet, so a test's own mark is the only one."""
        return self.marked_row_count == 0

    @property
    def rows_a_deletion_at_the_top_would_retitle(self) -> tuple[int, ...]:
        """
        Row numbers that name a different title once the first data row is deleted.

        Deleting the top row lifts every row below it by one, so these are the rows holding a
        different title from the row beneath. A title repeats freely - a season-per-row workbook
        names the same show a dozen times over - so no fixed row number holds across the library.
        The top row itself is excluded: it is the one being deleted, not one shifted under it.
        """
        return tuple(
            FIRST_DATA_ROW + offset
            for offset in range(1, len(self.titles) - 1)
            if self.titles[offset] != self.titles[offset + 1]
        )

    def copied_into(self, library: Path) -> Path:
        """This same workbook inside a disposable copy of the library."""
        return library / self.file_name


def range_ends(sheet: Worksheet) -> tuple[int, ...]:
    """Every row a table, dropdown or highlight rule stops at; one value means they agree."""
    ends = {range_boundaries(sheet.tables[name].ref)[3] for name in sheet.tables}
    ends |= {area.max_row for rule in sheet.data_validations.dataValidation for area in rule.sqref.ranges}
    ends |= {area.max_row for entry in sheet.conditional_formatting for area in entry.sqref.ranges}
    return tuple(sorted(ends))


def grid_table_name(sheet: Worksheet) -> str:
    """The name of the table anchored at A1, the one that owns the watch order itself."""
    for name in sheet.tables:
        min_col, min_row, _, _ = range_boundaries(sheet.tables[name].ref)
        if min_col == GRID_ANCHOR_COLUMN and min_row == HEADER_ROW:
            return name
    raise AssertionError(f"{sheet.title} carries no table anchored at A1")


def _key_for_role(detail: CategoryDetail, role: str) -> str:
    return next(column.key for column in detail.columns if column.role == role)


def read_facts(path: Path) -> WorkbookFacts:
    """Everything about a workbook a test may want to select it on."""
    detail = read_category(path)
    sheet = load_workbook(path).worksheets[0]
    order_key = _key_for_role(detail, "order")
    order_index = next(column.index for column in detail.columns if column.key == order_key)
    watch_key = _key_for_role(detail, "watch")
    indexes = tuple((column.key, column.index) for column in detail.columns)
    counts = detail.counts
    return WorkbookFacts(
        path=path,
        category_id=detail.id,
        row_count=counts.total,
        marked_row_count=counts.watched + counts.in_progress + counts.skipped,
        watched_row_count=counts.watched,
        unwatched_row_count=counts.unwatched,
        column_count=len(detail.columns),
        order_label=next(column.label for column in detail.columns if column.key == order_key),
        order_key=order_key,
        order_holds_numbers=isinstance(sheet.cell(row=FIRST_DATA_ROW, column=order_index).value, int | float),
        title_key=_key_for_role(detail, "title"),
        watch_key=watch_key,
        free_text_keys=tuple(
            column.key for column in detail.columns if column.kind == "text" and column.role == "other"
        ),
        dropdown_keys_beyond_the_watch_column=tuple(
            column.key for column in detail.columns if column.kind == "choice" and column.key != watch_key
        ),
        reference_sheet_count=len(detail.reference_sheets),
        range_ends=range_ends(sheet),
        grid_table_name=grid_table_name(sheet),
        column_indexes=indexes,
        titles=tuple(row.cells.get(_key_for_role(detail, "title"), "") for row in detail.rows),
    )
