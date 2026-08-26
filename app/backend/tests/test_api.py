from __future__ import annotations

import shutil
from collections.abc import AsyncIterator
from pathlib import Path

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from library_choice import (
    a_name_the_library_does_not_hold,
    another_category,
    every_workbook,
    the_category_a_new_name_could_shadow,
)
from tv_watchlist.api.dependencies import get_catalog
from tv_watchlist.constants import FIRST_DATA_ROW, IN_PROGRESS, SKIPPED, WATCHED, WORKBOOK_SUFFIX
from tv_watchlist.main import create_app
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.naming import workbook_filename
from workbook_facts import WorkbookFacts

BROKEN_FILE = "Broken_Watch_Order.xlsx"
BROKEN_DETAIL_FILE = "Broken_Master_Watch_Order.xlsx"
BROKEN_DETAIL_ID = "broken-master-watch-order"
DROPPED_IN_FILE = "Dropped_In_Master_Watch_Order.xlsx"
DROPPED_IN_NAME = "Dropped_In_Master_Watch_Order"
NEW_CATEGORY_NAME = a_name_the_library_does_not_hold("Cowboy Bebop")
NEW_CATEGORY_STEM = "Cowboy_Bebop"
WIDE_CATEGORY_NAME = a_name_the_library_does_not_hold("Twin Peaks")
OVERWIDE_CATEGORY_NAME = a_name_the_library_does_not_hold("Too Wide")
RENAMED_STEM = a_name_the_library_does_not_hold("Renamed_Franchise_Master_Watch_Order")
RENAMED_ID = "renamed-franchise-master-watch-order"
AN_ACCENT = "#2563EB"
A_TITLE_THE_LIBRARY_DOES_NOT_HOLD = "Fan favourite OVA"
A_VALUE_NO_DROPDOWN_OFFERS = "maybe"
AN_ILLEGAL_STEM = "Alpha/Beta"
A_NOTE_WITH_A_CONTROL_CHARACTER = f"before{chr(7)}after"
A_NOTE_WITHOUT_IT = "beforeafter"
TOO_MANY_COLUMNS = 200
HERO_SUFFIX = ".jpg"


@pytest.fixture
async def client(catalog: Catalog) -> AsyncIterator[AsyncClient]:
    app = create_app()
    app.dependency_overrides[get_catalog] = lambda: catalog
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as async_client:
        yield async_client


async def test_listing_covers_every_workbook(client: AsyncClient, library: Path) -> None:
    body = (await client.get("/api/categories")).json()
    assert len(body["categories"]) == len(list(library.glob(f"*{WORKBOOK_SUFFIX}")))
    assert body["unreadable"] == []
    assert body["shadowed"] == []
    names = {category["name"] for category in body["categories"]}
    assert names == {path.stem for path in library.glob(f"*{WORKBOOK_SUFFIX}")}
    assert {facts.stem for facts in every_workbook()} <= names


async def test_listing_reports_a_file_it_cannot_open(client: AsyncClient, library: Path) -> None:
    (library / BROKEN_FILE).write_bytes(b"not a workbook")
    body = (await client.get("/api/categories")).json()
    assert [item["file_name"] for item in body["unreadable"]] == [BROKEN_FILE]


async def test_listing_reports_a_file_that_another_file_shadows(client: AsyncClient, library: Path) -> None:
    facts = the_category_a_new_name_could_shadow()
    shutil.copy2(facts.copied_into(library), library / workbook_filename(facts.shadowing_name))

    body = (await client.get("/api/categories")).json()

    assert [item["category_id"] for item in body["shadowed"]] == [facts.category_id]
    assert len([item for item in body["categories"] if item["id"] == facts.category_id]) == 1


async def test_cycling_watch_status_writes_through_to_the_workbook(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    row = detail["rows"][0]["row"]

    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{row}",
        json={"cells": {subject.watch_key: WATCHED}, "expected_mtime": detail["mtime"]},
    )

    assert response.status_code == 200
    updated = response.json()
    assert updated["rows"][0]["cells"][subject.watch_key] == WATCHED
    assert updated["counts"]["watched"] == detail["counts"]["watched"] + 1
    assert updated["mtime"] != detail["mtime"]

    reloaded = (await client.get(f"/api/categories/{subject.category_id}")).json()
    assert reloaded["rows"][0]["cells"][subject.watch_key] == WATCHED


async def test_a_stale_write_is_refused_instead_of_clobbering(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    row = detail["rows"][0]["row"]
    await client.patch(
        f"/api/categories/{subject.category_id}/rows/{row}",
        json={"cells": {subject.watch_key: WATCHED}, "expected_mtime": detail["mtime"]},
    )

    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{row}",
        json={"cells": {subject.watch_key: SKIPPED}, "expected_mtime": detail["mtime"]},
    )

    assert response.status_code == 409
    assert response.json()["error"] == "StaleWorkbookError"


async def test_a_workbook_open_in_excel_is_refused(
    client: AsyncClient, subject: WorkbookFacts, subject_path: Path
) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    excel_lock_path(subject_path).write_bytes(b"owner")

    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={"cells": {subject.watch_key: WATCHED}, "expected_mtime": detail["mtime"]},
    )

    assert response.status_code == 423
    assert "open in Excel" in response.json()["message"]


async def test_a_value_outside_the_sheet_dropdown_is_rejected(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={"cells": {subject.watch_key: A_VALUE_NO_DROPDOWN_OFFERS}, "expected_mtime": detail["mtime"]},
    )
    assert response.status_code == 400
    assert response.json()["error"] == "InvalidChoiceError"


async def test_add_and_delete_a_row(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    total = detail["counts"]["total"]

    added = await client.post(
        f"/api/categories/{subject.category_id}/rows",
        json={
            "cells": {subject.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD},
            "expected_mtime": detail["mtime"],
        },
    )
    assert added.status_code == 200
    body = added.json()
    assert body["counts"]["total"] == total + 1
    assert body["rows"][-1]["cells"][subject.title_key] == A_TITLE_THE_LIBRARY_DOES_NOT_HOLD

    removed = await client.delete(
        f"/api/categories/{subject.category_id}/rows/{body['rows'][-1]['row']}",
        params={"expected_mtime": body["mtime"]},
    )
    assert removed.status_code == 200
    assert removed.json()["counts"]["total"] == total


async def test_creating_a_category_makes_it_appear_in_the_listing(client: AsyncClient, library: Path) -> None:
    response = await client.post(
        "/api/categories",
        json={"name": NEW_CATEGORY_NAME, "accent": AN_ACCENT, "titles": [NEW_CATEGORY_NAME]},
    )
    assert response.status_code == 201
    created = response.json()
    assert created["name"] == NEW_CATEGORY_STEM
    assert (library / f"{NEW_CATEGORY_STEM}{WORKBOOK_SUFFIX}").exists()

    listing = (await client.get("/api/categories")).json()
    assert NEW_CATEGORY_STEM in {category["name"] for category in listing["categories"]}


async def test_a_dropped_in_sheet_can_be_given_a_watch_column(client: AsyncClient, library: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for index, header in enumerate(["Order", "Title"], start=1):
        sheet.cell(row=1, column=index).value = header
    sheet.cell(row=FIRST_DATA_ROW, column=1).value = 1
    sheet.cell(row=FIRST_DATA_ROW, column=2).value = "Serial Experiments Lain"
    workbook.save(library / DROPPED_IN_FILE)

    listing = (await client.get("/api/categories")).json()
    dropped = next(item for item in listing["categories"] if item["name"] == DROPPED_IN_NAME)
    assert dropped["has_watch_column"] is False

    response = await client.post(
        f"/api/categories/{dropped['id']}/watch-column", params={"expected_mtime": dropped["mtime"]}
    )
    assert response.status_code == 200
    assert response.json()["has_watch_column"] is True


async def test_an_unknown_category_is_a_404(client: AsyncClient) -> None:
    response = await client.get("/api/categories/does-not-exist")
    assert response.status_code == 404


async def test_the_first_write_snapshots_a_backup(
    client: AsyncClient, catalog: Catalog, subject: WorkbookFacts
) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={"cells": {subject.watch_key: WATCHED}, "expected_mtime": detail["mtime"]},
    )
    backups = list(catalog._settings.backup_dir.glob(f"{subject.stem}__*{WORKBOOK_SUFFIX}"))  # noqa: SLF001
    assert len(backups) == 1


async def test_the_excel_lock_flag_is_live_not_cached(
    client: AsyncClient, subject: WorkbookFacts, subject_path: Path
) -> None:
    category = subject.category_id
    assert (await client.get(f"/api/categories/{category}")).json()["locked_by_excel"] is False

    excel_lock_path(subject_path).write_bytes(b"owner")

    assert (await client.get(f"/api/categories/{category}")).json()["locked_by_excel"] is True
    listing = (await client.get("/api/categories")).json()
    listed = next(item for item in listing["categories"] if item["id"] == category)
    assert listed["locked_by_excel"] is True

    excel_lock_path(subject_path).unlink()
    assert (await client.get(f"/api/categories/{category}")).json()["locked_by_excel"] is False


async def test_control_characters_are_stripped_rather_than_crashing(
    client: AsyncClient, subject: WorkbookFacts
) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={
            "cells": {subject.free_text_key: A_NOTE_WITH_A_CONTROL_CHARACTER},
            "expected_mtime": detail["mtime"],
        },
    )
    assert response.status_code == 200
    assert response.json()["rows"][0]["cells"][subject.free_text_key] == A_NOTE_WITHOUT_IT


async def test_an_unreadable_workbook_is_a_422_not_a_500(client: AsyncClient, library: Path) -> None:
    (library / BROKEN_DETAIL_FILE).write_bytes(b"not a workbook")
    response = await client.get(f"/api/categories/{BROKEN_DETAIL_ID}")
    assert response.status_code == 422
    assert response.json()["error"] == "UnreadableWorkbookError"


async def test_a_new_category_with_repeated_columns_stays_valid(client: AsyncClient) -> None:
    response = await client.post(
        "/api/categories", json={"name": WIDE_CATEGORY_NAME, "columns": ["Order", "Title", "Title", "Notes"]}
    )
    assert response.status_code == 201
    labels = [column["label"] for column in response.json()["columns"]]
    assert labels == ["Order", "Title", "Title 2", "Notes", "Watched?"]
    assert len(set(labels)) == len(labels)


async def test_a_new_category_cannot_request_an_unbounded_column_list(client: AsyncClient) -> None:
    response = await client.post(
        "/api/categories", json={"name": OVERWIDE_CATEGORY_NAME, "columns": [f"c{n}" for n in range(TOO_MANY_COLUMNS)]}
    )
    assert response.status_code == 422


async def test_renaming_a_category_changes_its_id_and_keeps_its_rows(
    client: AsyncClient, subject: WorkbookFacts, subject_path: Path, library: Path
) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()

    response = await client.post(
        f"/api/categories/{subject.category_id}/rename",
        json={"stem": RENAMED_STEM, "expected_mtime": detail["mtime"]},
    )

    assert response.status_code == 200
    renamed = response.json()
    assert renamed["id"] == RENAMED_ID
    assert renamed["name"] == RENAMED_STEM
    assert renamed["counts"]["total"] == detail["counts"]["total"]
    assert (library / f"{RENAMED_STEM}{WORKBOOK_SUFFIX}").exists()
    assert not subject_path.exists()

    listing = (await client.get("/api/categories")).json()
    names = {category["name"] for category in listing["categories"]}
    assert RENAMED_STEM in names
    assert subject.stem not in names


async def test_renaming_onto_an_existing_name_is_a_409(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    response = await client.post(
        f"/api/categories/{subject.category_id}/rename",
        json={"stem": another_category().stem, "expected_mtime": detail["mtime"]},
    )
    assert response.status_code == 409
    assert response.json()["error"] == "DuplicateCategoryError"


async def test_an_illegal_rename_is_a_400_with_a_readable_reason(client: AsyncClient, subject: WorkbookFacts) -> None:
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()
    response = await client.post(
        f"/api/categories/{subject.category_id}/rename",
        json={"stem": AN_ILLEGAL_STEM, "expected_mtime": detail["mtime"]},
    )
    assert response.status_code == 400
    assert "cannot contain /" in response.json()["message"]


async def test_a_dropped_in_hero_image_is_offered_to_the_card(
    client: AsyncClient, catalog: Catalog, subject: WorkbookFacts
) -> None:
    heroes_dir = catalog._settings.heroes_dir  # noqa: SLF001 - fixture wiring
    heroes_dir.mkdir(parents=True, exist_ok=True)
    (heroes_dir / f"{subject.category_id}{HERO_SUFFIX}").write_bytes(b"jpeg-ish")

    listing = (await client.get("/api/categories")).json()
    pictured = next(item for item in listing["categories"] if item["id"] == subject.category_id)
    assert pictured["hero_url"].split("?")[0] == f"/heroes/{subject.category_id}{HERO_SUFFIX}"

    bare = next(item for item in listing["categories"] if item["id"] == another_category().category_id)
    assert bare["hero_url"] is None


async def test_a_row_can_be_marked_in_progress(client: AsyncClient, subject: WorkbookFacts) -> None:
    """The dropdown offers more than watched, and the counts have to split them apart."""
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()

    response = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={"cells": {subject.watch_key: IN_PROGRESS}, "expected_mtime": detail["mtime"]},
    )

    assert response.status_code == 200
    assert response.json()["counts"]["in_progress"] == detail["counts"]["in_progress"] + 1
