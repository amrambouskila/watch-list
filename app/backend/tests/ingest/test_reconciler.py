"""The reconciler against a real Postgres and real workbooks: identity, adoption, refusal, write-back.

Two things are never stood in for. The database is real and skips loudly without one, because the whole
premise under test is a dual store and a mock would prove only that the mock agrees with itself. The
workbook layer is real too - every pass here writes an actual `.xlsx` through `Catalog._write`, on a
throwaway copy of one of the owner's own files, chosen by shape rather than by name (CLAUDE.md 16.3-16.4).

Providers are absent on purpose in most of these: with no TMDB token the resolver is disabled, every
entry stays `unresolved` with the reason on it, and the half of the pass this file is about - identity,
adoption, minting, soft delete, position, the single write-back - runs exactly as it does in production.
"""

from __future__ import annotations

import logging
import os
import uuid
from collections.abc import Sequence
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Final

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from backup_support import restore_points
from library_choice import every_workbook, ordinary_editable_categories
from snapshot import fingerprint
from tv_watchlist.config import Settings
from tv_watchlist.constants import FIRST_DATA_ROW, TEMP_WRITE_PREFIX, WATCHED
from tv_watchlist.data.enums import (
    MediaType,
    ResolutionState,
    WatchEventKind,
    WatchEventSource,
    WatchSubjectKind,
)
from tv_watchlist.data.models import Category, CuratedEntry, Title, WatchEvent
from tv_watchlist.ingest.reconcile_outcome import ReconcileOutcome
from tv_watchlist.ingest.reconciler import (
    MALFORMED_IDENTITY_TEMPLATE,
    NOT_CONFIGURED_NOTE,
    Reconciler,
)
from tv_watchlist.models.row_write import RowWrite
from tv_watchlist.resolver.resolution_request import ResolutionRequest
from tv_watchlist.resolver.title_resolution import TitleResolution
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.services.mirror_outcome import WRITTEN
from tv_watchlist.services.watch_mirror import mirror_category
from tv_watchlist.workbook.errors import WorkbookLockedError
from tv_watchlist.workbook.identity import ensure_identity_column, write_row_ids
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.reader import read_category, read_projection
from workbook_edits import narrow_the_status_dropdown
from workbook_facts import WorkbookFacts

CANONICAL_UUID_LENGTH: Final[int] = 36
A_NEW_TITLE: Final[str] = "a different film entirely"
# Deep enough that "more than half the rows" is a meaningful retitle rather than two rows.
ROWS_ENOUGH_FOR_A_MISMATCH: Final[int] = 8
# What a sheet dropped into the library may offer instead of the four values the nineteen carry.
A_FOREIGN_VOCABULARY: Final[tuple[str, ...]] = ("Yes", "No")


@pytest.fixture
def subject(library: Path) -> WorkbookFacts:
    """An ordinary editable category deep enough for the reordering and mismatch cases below."""
    found = ordinary_editable_categories(lambda facts: facts.row_count >= ROWS_ENOUGH_FOR_A_MISMATCH)
    if not found:
        raise AssertionError(f"the library holds no editable category of at least {ROWS_ENOUGH_FOR_A_MISMATCH} rows")
    return found[0]


@pytest.fixture
def reconciler(db_session: AsyncSession, catalog: Catalog) -> Reconciler:
    """Providers absent: a missing token is a disabled feature, never a failure (section 8.3)."""
    return Reconciler(db_session, catalog, resolver=None, tmdb_ingest=None, order_ingest=None)


async def _entries(session: AsyncSession, slug: str) -> list[CuratedEntry]:
    rows = await session.scalars(
        select(CuratedEntry)
        .join(Category, Category.id == CuratedEntry.category_id)
        .where(Category.slug == slug)
        .order_by(CuratedEntry.position)
    )
    return list(rows)


async def _event_count(session: AsyncSession) -> int:
    """Every watch event in the database, so a pass that quietly appended one can be caught at it."""
    await session.flush()
    return int(await session.scalar(select(func.count()).select_from(WatchEvent)) or 0)


async def _category(session: AsyncSession, slug: str) -> Category:
    found = await session.scalar(select(Category).where(Category.slug == slug))
    if found is None:
        raise AssertionError(f"no category row for {slug}")
    return found


def _sheet_ids(path: Path) -> dict[int, str]:
    return read_projection(path).row_ids


def _swap_rows(path: Path, first: int, second: int) -> None:
    """Swap two whole sheet rows the way a user dragging them in Excel would - identity cell included."""
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    for column in range(1, (sheet.max_column or 0) + 1):
        top, bottom = sheet.cell(row=first, column=column), sheet.cell(row=second, column=column)
        top.value, bottom.value = bottom.value, top.value
    workbook.save(path)


def _retitle(path: Path, title_index: int, rows: Sequence[int]) -> None:
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    for row in rows:
        sheet.cell(row=row, column=title_index).value = f"something else entirely {row}"
    workbook.save(path)


def _clear_identity_column(path: Path) -> None:
    """Delete the whole Row ID column, the way a user tidying their sheet would."""
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    sheet.delete_cols(sheet.max_column)
    workbook.save(path)


async def _an_event(session: AsyncSession, entry: CuratedEntry) -> WatchEvent:
    """A watch fact against one entry, so a later pass can be shown not to have orphaned it."""
    event = WatchEvent(
        occurred_at=datetime.now(UTC),
        subject_kind=WatchSubjectKind.CURATED_ENTRY,
        kind=WatchEventKind.COMPLETED,
        source=WatchEventSource.UI,
        curated_entry_id=entry.id,
        category_id=entry.category_id,
    )
    session.add(event)
    await session.flush()
    return event


async def test_a_first_ingest_mints_a_uuid_for_every_row_and_writes_them_back_once(
    reconciler: Reconciler,
    db_session: AsyncSession,
    settings: Settings,
    subject: WorkbookFacts,
    subject_path: Path,
) -> None:
    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    sheet_ids = _sheet_ids(subject_path)
    category = await _category(db_session, subject.category_id)
    assert outcome == ReconcileOutcome(slug=subject.category_id, minted=subject.row_count, workbook_written=True)
    assert outcome.refreshed == 0
    assert len(entries) == subject.row_count
    assert [entry.position for entry in entries] == list(range(1, subject.row_count + 1))
    assert {len(value) for value in sheet_ids.values()} == {CANONICAL_UUID_LENGTH}
    assert {str(entry.id) for entry in entries} == set(sheet_ids.values())
    # The stamp stored is the one the file carries AFTER our own write, or the next pass re-ingests.
    assert category.workbook_mtime == subject_path.stat().st_mtime
    assert len(restore_points(settings, subject)) == 1
    assert not list(subject_path.parent.glob(f"{TEMP_WRITE_PREFIX}*"))


async def test_a_second_ingest_with_no_change_performs_zero_workbook_writes_and_zero_row_updates(
    reconciler: Reconciler, db_session: AsyncSession, settings: Settings, subject: WorkbookFacts, subject_path: Path
) -> None:
    await reconciler.reconcile(subject.category_id)
    stamped = subject_path.stat().st_mtime_ns
    before = {entry.id: entry.updated_at for entry in await _entries(db_session, subject.category_id)}

    outcome = await reconciler.reconcile(subject.category_id)

    after = {entry.id: entry.updated_at for entry in await _entries(db_session, subject.category_id)}
    assert outcome == ReconcileOutcome(slug=subject.category_id, skipped=True)
    assert subject_path.stat().st_mtime_ns == stamped
    assert after == before
    assert len(restore_points(settings, subject)) == 1


async def test_an_excel_resave_without_an_edit_is_a_digest_no_op(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """Excel rewrites the file on every close; the mtime alone would re-ingest nineteen workbooks."""
    await reconciler.reconcile(subject.category_id)
    stamp = subject_path.stat().st_mtime
    os.utime(subject_path, (stamp + 10, stamp + 10))
    resaved = subject_path.stat().st_mtime
    assert resaved != stamp
    before = {
        entry.id: (entry.updated_at, entry.mirrored_watch_cell)
        for entry in await _entries(db_session, subject.category_id)
    }
    events = await _event_count(db_session)

    outcome = await reconciler.reconcile(subject.category_id)

    after = {
        entry.id: (entry.updated_at, entry.mirrored_watch_cell)
        for entry in await _entries(db_session, subject.category_id)
    }
    category = await _category(db_session, subject.category_id)
    assert outcome.skipped is True
    assert outcome.workbook_written is False
    # Ruling 49: the digest decided the skip, and the stamp Excel moved is recorded anyway - the next
    # pass and every write-back guard compare against it, so it has to agree with the file.
    assert category.workbook_mtime == resaved
    assert subject_path.stat().st_mtime == resaved
    assert after == before
    assert await _event_count(db_session) == events


async def test_reordering_rows_changes_only_position(
    reconciler: Reconciler,
    db_session: AsyncSession,
    catalog: Catalog,
    settings: Settings,
    subject: WorkbookFacts,
    subject_path: Path,
) -> None:
    await reconciler.reconcile(subject.category_id)
    first, second = FIRST_DATA_ROW, FIRST_DATA_ROW + 1
    moved = _sheet_ids(subject_path)
    top, below = uuid.UUID(moved[first]), uuid.UUID(moved[second])
    event = await _an_event(db_session, await db_session.get_one(CuratedEntry, top))
    # The COMPLETED event above is a cell the sheet is owed, and section 5.16 requires the mirror to pay
    # it. It is paid here, BEFORE the reorder, so that the pass under test carries no mirror work: a
    # reorder that wrote the workbook for an unpaid cell would prove nothing about reordering.
    category_id = (await _category(db_session, subject.category_id)).id
    assert await mirror_category(db_session, catalog, category_id) == WRITTEN

    entries = await _entries(db_session, subject.category_id)
    before = {entry.id: entry.position for entry in entries}
    mirrored = {entry.id: entry.mirrored_watch_cell for entry in entries}
    events = await _event_count(db_session)
    snapshots = len(restore_points(settings, subject))
    _swap_rows(subject_path, first, second)
    stamped = subject_path.stat().st_mtime_ns

    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    after = {entry.id: entry.position for entry in entries}
    assert set(after) == set(before)
    assert (after[top], after[below]) == (before[below], before[top])
    assert outcome.reminted == 0
    assert outcome.minted == 0
    assert outcome.matched == subject.row_count
    assert all(entry.removed_at is None for entry in entries)
    # No identity moved and the mirror was already level, so nothing needed writing: the file still
    # carries the stamp the reorder save left on it.
    assert outcome.workbook_written is False
    assert subject_path.stat().st_mtime_ns == stamped
    assert {entry.id: entry.mirrored_watch_cell for entry in entries} == mirrored
    assert await _event_count(db_session) == events
    assert len(restore_points(settings, subject)) == snapshots
    assert (await db_session.get_one(WatchEvent, event.id)).curated_entry_id == top


async def test_a_cleared_row_id_column_is_recovered_by_adoption_with_history_intact(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    await reconciler.reconcile(subject.category_id)
    before = {entry.id: entry.position for entry in await _entries(db_session, subject.category_id)}
    event = await _an_event(db_session, (await _entries(db_session, subject.category_id))[0])
    _clear_identity_column(subject_path)

    outcome = await reconciler.reconcile(subject.category_id)

    after = {entry.id: entry.position for entry in await _entries(db_session, subject.category_id)}
    assert outcome.adopted == subject.row_count
    assert outcome.minted == 0
    assert after == before
    assert read_projection(subject_path).identity_key is not None
    assert {str(entry_id) for entry_id in after} == set(_sheet_ids(subject_path).values())
    assert (await db_session.get_one(WatchEvent, event.id)).curated_entry_id in after


async def test_a_near_total_mismatch_is_refused_with_last_ingest_error_and_no_write(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    await reconciler.reconcile(subject.category_id)
    before = {entry.id: entry.raw_title for entry in await _entries(db_session, subject.category_id)}
    _clear_identity_column(subject_path)
    _retitle(
        subject_path,
        subject.index_of(subject.title_key),
        range(FIRST_DATA_ROW, FIRST_DATA_ROW + subject.row_count // 2 + 1),
    )
    # After the test's own saves, not before them: the claim is that the REFUSED PASS wrote nothing.
    stamped = subject_path.stat().st_mtime_ns

    outcome = await reconciler.reconcile(subject.category_id)

    after = {entry.id: entry.raw_title for entry in await _entries(db_session, subject.category_id)}
    category = await _category(db_session, subject.category_id)
    assert outcome.refused is not None
    assert category.last_ingest_error == outcome.refused
    assert after == before
    assert subject_path.stat().st_mtime_ns == stamped


async def test_a_non_uuid_row_id_cell_refuses_the_workbook_and_writes_nothing(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """Ruling 20: minting over a hand-typed cell would destroy whatever the owner put there."""
    ensure_identity_column(subject_path)
    write_row_ids(subject_path, {FIRST_DATA_ROW: "mine, do not touch"})
    stamped = subject_path.stat().st_mtime_ns

    outcome = await reconciler.reconcile(subject.category_id)

    category = await _category(db_session, subject.category_id)
    assert outcome.refused == MALFORMED_IDENTITY_TEMPLATE.format(rows=str(FIRST_DATA_ROW))
    assert category.last_ingest_error == outcome.refused
    # No stamp was stored, so the next pass retries rather than skipping the file for good.
    assert (category.workbook_mtime, category.workbook_digest) == (None, None)
    assert await _entries(db_session, subject.category_id) == []
    assert subject_path.stat().st_mtime_ns == stamped


async def test_a_duplicated_row_id_keeps_the_lowest_position_and_re_mints_the_rest(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    await reconciler.reconcile(subject.category_id)
    ids = _sheet_ids(subject_path)
    first, later = FIRST_DATA_ROW, FIRST_DATA_ROW + 3
    kept = ids[first]
    write_row_ids(subject_path, {later: kept})

    outcome = await reconciler.reconcile(subject.category_id)

    after = _sheet_ids(subject_path)
    assert outcome.reminted == 1
    assert after[first] == kept
    assert after[later] != kept
    assert len(set(after.values())) == subject.row_count
    assert {str(entry.id) for entry in await _entries(db_session, subject.category_id)} == set(after.values())


async def test_a_row_removed_from_the_sheet_is_soft_deleted_never_hard_deleted(
    reconciler: Reconciler, db_session: AsyncSession, catalog: Catalog, subject: WorkbookFacts, subject_path: Path
) -> None:
    await reconciler.reconcile(subject.category_id)
    entries = await _entries(db_session, subject.category_id)
    last = entries[-1]
    event = await _an_event(db_session, last)
    detail = await catalog.detail(subject.category_id)
    await catalog.delete_row(subject.category_id, detail.rows[-1].row, detail.mtime)

    outcome = await reconciler.reconcile(subject.category_id)

    retired = await db_session.get_one(CuratedEntry, last.id)
    assert outcome.removed == 1
    assert retired.removed_at is not None
    # ON DELETE RESTRICT is never exercised, because nothing in the curated chain is ever deleted.
    assert (await db_session.get_one(WatchEvent, event.id)).curated_entry_id == last.id
    assert (
        await db_session.scalar(
            select(func.count()).select_from(CuratedEntry).where(CuratedEntry.category_id == last.category_id)
        )
        == subject.row_count
    )


async def test_a_confirmed_entry_with_unchanged_cells_is_never_re_resolved(
    db_session: AsyncSession, catalog: Catalog, subject: WorkbookFacts, subject_path: Path
) -> None:
    counting = _CountingResolver()
    reconciler = Reconciler(db_session, catalog, resolver=counting, tmdb_ingest=_NoTitles(), order_ingest=None)
    await reconciler.reconcile(subject.category_id)
    entries = await _entries(db_session, subject.category_id)
    confirmed = entries[0]
    confirmed.resolution_state = ResolutionState.CONFIRMED
    await db_session.flush()
    counting.seen.clear()

    detail = await catalog.detail(subject.category_id)
    await catalog.update_row(
        subject.category_id,
        detail.rows[1].row,
        RowWrite(cells={subject.title_key: A_NEW_TITLE}, expected_mtime=detail.mtime),
    )
    outcome = await reconciler.reconcile(subject.category_id)

    # Every other row is still unresolved and so is searched again; the confirmed one is not, which is
    # exactly one fewer call than there are rows.
    assert outcome.reresolved == subject.row_count - 1
    assert len(counting.seen) == subject.row_count - 1
    assert A_NEW_TITLE in counting.seen


async def test_an_unconfigured_provider_leaves_every_entry_unresolved_with_the_reason_on_it(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts
) -> None:
    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    assert outcome.reresolved == 0
    assert {entry.resolution_state for entry in entries} == {ResolutionState.UNRESOLVED}
    assert {entry.resolution_note for entry in entries} == {NOT_CONFIGURED_NOTE}
    assert {entry.title_id for entry in entries} == {None}


async def test_an_unconfigured_provider_does_not_undo_a_confirmed_entry_whose_cells_moved(
    reconciler: Reconciler, db_session: AsyncSession, catalog: Catalog, subject: WorkbookFacts
) -> None:
    """A human's confirmation is not lost because a token was absent; the cells moving is a reason to
    re-resolve, not a reason to forget the standing answer before anything has been asked."""
    await reconciler.reconcile(subject.category_id)
    entries = await _entries(db_session, subject.category_id)
    confirmed = entries[0]
    standing = uuid.uuid4()
    title = Title(id=standing, media_type=MediaType.MOVIE, tmdb_id=603, name="Max", match_key="max")
    db_session.add(title)
    confirmed.title_id = standing
    confirmed.resolution_state = ResolutionState.CONFIRMED
    confirmed.resolution_confidence = Decimal("0.95")
    await db_session.flush()

    detail = await catalog.detail(subject.category_id)
    await catalog.update_row(
        subject.category_id,
        detail.rows[0].row,
        RowWrite(cells={subject.title_key: A_NEW_TITLE}, expected_mtime=detail.mtime),
    )
    await reconciler.reconcile(subject.category_id)

    kept = await db_session.get_one(CuratedEntry, confirmed.id)
    assert kept.title_id == standing
    assert kept.resolution_state is ResolutionState.CONFIRMED
    assert kept.resolution_confidence == Decimal("0.95")
    assert kept.resolution_note == NOT_CONFIGURED_NOTE
    assert kept.raw_title == A_NEW_TITLE


async def test_a_locked_workbook_rolls_the_pass_back_and_the_next_pass_replays_the_identical_diff(
    reconciler: Reconciler,
    db_session: AsyncSession,
    subject: WorkbookFacts,
    subject_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    excel_lock_path(subject_path).write_bytes(b"owner")

    with caplog.at_level(logging.INFO, logger="tv_watchlist.ingest.reconciler"), pytest.raises(WorkbookLockedError):
        await reconciler.reconcile(subject.category_id)

    # The pass says why it stopped rather than going silent after `reconcile_start`.
    assert "refused=locked" in caplog.records[-1].getMessage()
    assert await _entries(db_session, subject.category_id) == []

    excel_lock_path(subject_path).unlink()
    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    assert outcome.minted == subject.row_count
    assert {str(entry.id) for entry in entries} == set(_sheet_ids(subject_path).values())


async def test_the_write_back_is_one_snapshot_one_replace(
    reconciler: Reconciler, settings: Settings, subject: WorkbookFacts, subject_path: Path
) -> None:
    before = fingerprint(subject_path)
    stamped = subject_path.stat().st_mtime_ns

    await reconciler.reconcile(subject.category_id)

    after = fingerprint(subject_path)
    order = load_workbook(subject_path).worksheets[0].title
    assert len(restore_points(settings, subject)) == 1
    assert subject_path.stat().st_mtime_ns != stamped
    assert not list(subject_path.parent.glob(f"{TEMP_WRITE_PREFIX}*"))
    # One column wider, and nothing else about the sheet disturbed.
    assert after[order]["dimensions"] == (before[order]["dimensions"][0], before[order]["dimensions"][1] + 1)
    assert after[order]["header_style"] == before[order]["header_style"]
    assert after[order]["body_style"] == before[order]["body_style"]
    assert after[order]["conditional_rule_count"] == before[order]["conditional_rule_count"]


async def test_rows_with_only_order_and_title_ingest_as_unresolved_first_class(
    db_session: AsyncSession, catalog: Catalog, library: Path
) -> None:
    """A row that resolves to nothing is still a first-class row (CLAUDE.md section 18.2)."""
    sparse = _the_category_filling_only_order_and_title()
    sparse.copied_into(library)
    reconciler = Reconciler(db_session, catalog, resolver=None, tmdb_ingest=None, order_ingest=None)

    outcome = await reconciler.reconcile(sparse.category_id)

    entries = await _entries(db_session, sparse.category_id)
    assert outcome.minted == sparse.row_count
    assert [entry.position for entry in entries] == list(range(1, sparse.row_count + 1))
    assert {entry.resolution_state for entry in entries} == {ResolutionState.UNRESOLVED}
    assert all(entry.raw_cells for entry in entries)
    assert all(entry.raw_title for entry in entries)


# -- the mirror, on the production path (T25) --------------------------------------------------------


def _type_status(path: Path, index: int, row: int, value: str) -> None:
    """A status the owner typed into Excel: straight through openpyxl, because this is the other writer."""
    workbook = load_workbook(path)
    workbook.worksheets[0].cell(row=row, column=index).value = value
    workbook.save(path)


def _status_of(path: Path, subject: WorkbookFacts, row: int) -> str:
    cells = {sheet_row.row: sheet_row.cells for sheet_row in read_category(path).rows}
    return cells[row].get(subject.watch_key, "").strip()


async def _mirror_events(session: AsyncSession, slug: str) -> list[WatchEvent]:
    rows = await session.scalars(
        select(WatchEvent)
        .join(Category, Category.id == WatchEvent.category_id)
        .where(Category.slug == slug, WatchEvent.source == WatchEventSource.EXCEL_MIRROR)
        .order_by(WatchEvent.occurred_at)
    )
    return list(rows)


async def test_a_status_typed_in_excel_becomes_one_entry_level_excel_mirror_event(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """The dual store's read direction: the cell is never truth, only the signal that a human moved it."""
    await reconciler.reconcile(subject.category_id)
    _type_status(subject_path, subject.index_of(subject.watch_key), FIRST_DATA_ROW, WATCHED)

    await reconciler.reconcile(subject.category_id)

    events = await _mirror_events(db_session, subject.category_id)
    entries = await _entries(db_session, subject.category_id)
    assert len(events) == 1
    assert events[0].subject_kind is WatchSubjectKind.CURATED_ENTRY
    assert events[0].kind is WatchEventKind.COMPLETED
    assert events[0].curated_entry_id == entries[0].id
    assert events[0].episode_id is None
    assert events[0].note is not None and WATCHED in events[0].note
    # The write that closes the loop: what the pass mirrored back is recorded, so the next read is an echo.
    assert entries[0].mirrored_watch_cell == WATCHED
    assert entries[0].observed_watch_cell == WATCHED


async def test_the_mirror_cell_rides_the_single_write_back(
    reconciler: Reconciler, db_session: AsyncSession, settings: Settings, subject: WorkbookFacts, subject_path: Path
) -> None:
    """One snapshot, one temp file, one `os.replace` - the mirror never opens the workbook of its own."""
    await reconciler.reconcile(subject.category_id)
    before = fingerprint(subject_path)
    taken = len(restore_points(settings, subject))
    _type_status(subject_path, subject.index_of(subject.watch_key), FIRST_DATA_ROW, WATCHED)

    outcome = await reconciler.reconcile(subject.category_id)

    category = await _category(db_session, subject.category_id)
    assert outcome.workbook_written is True
    # Every identity was minted by the first pass, so this write-back exists only for the mirror cell.
    assert (outcome.minted, outcome.adopted) == (0, 0)
    assert len(restore_points(settings, subject)) == taken + 1
    assert fingerprint(subject_path) == before
    assert not list(subject_path.parent.glob(f"{TEMP_WRITE_PREFIX}*"))
    assert _status_of(subject_path, subject, FIRST_DATA_ROW) == WATCHED
    assert category.workbook_mtime == subject_path.stat().st_mtime


async def test_a_third_pass_appends_no_event_and_writes_no_cell(
    reconciler: Reconciler, db_session: AsyncSession, settings: Settings, subject: WorkbookFacts, subject_path: Path
) -> None:
    """The echo loop that does not grow: our own mirrored value is read back as ours, not as an edit."""
    await reconciler.reconcile(subject.category_id)
    _type_status(subject_path, subject.index_of(subject.watch_key), FIRST_DATA_ROW, WATCHED)
    await reconciler.reconcile(subject.category_id)
    stamped = subject_path.stat().st_mtime_ns
    taken = len(restore_points(settings, subject))

    outcome = await reconciler.reconcile(subject.category_id)

    assert outcome.skipped is True
    assert outcome.workbook_written is False
    assert len(await _mirror_events(db_session, subject.category_id)) == 1
    assert subject_path.stat().st_mtime_ns == stamped
    assert len(restore_points(settings, subject)) == taken


async def test_a_locked_workbook_leaves_the_mirror_owing_the_identical_diff(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """The mirror is the writer that must lose: a refusal moves neither store, and the diff replays."""
    await reconciler.reconcile(subject.category_id)
    stamped = (await _category(db_session, subject.category_id)).workbook_mtime
    _type_status(subject_path, subject.index_of(subject.watch_key), FIRST_DATA_ROW, WATCHED)
    excel_lock_path(subject_path).write_bytes(b"owner")

    with pytest.raises(WorkbookLockedError):
        await reconciler.reconcile(subject.category_id)

    await db_session.rollback()
    entries = await _entries(db_session, subject.category_id)
    assert entries[0].mirrored_watch_cell == ""
    assert (await _category(db_session, subject.category_id)).workbook_mtime == stamped
    assert await _mirror_events(db_session, subject.category_id) == []

    excel_lock_path(subject_path).unlink()
    await reconciler.reconcile(subject.category_id)

    replayed = await _entries(db_session, subject.category_id)
    assert replayed[0].mirrored_watch_cell == WATCHED
    assert len(await _mirror_events(db_session, subject.category_id)) == 1
    assert _status_of(subject_path, subject, FIRST_DATA_ROW) == WATCHED


async def test_a_dropdown_the_mirror_cannot_satisfy_drops_the_cell_not_the_pass(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """Ruling 41 on the ingest path, where a raised `InvalidChoiceError` would roll back the whole pass.

    A real progress event is recorded first, so the mirror genuinely wants to write `Watched` into a
    column that now offers only `Yes` and `No`. Without the filter this pass raises and takes the whole
    workbook's ingest with it, every time, for as long as that dropdown stands.
    """
    index = subject.index_of(subject.watch_key)
    await reconciler.reconcile(subject.category_id)
    await _an_event(db_session, (await _entries(db_session, subject.category_id))[0])
    narrow_the_status_dropdown(subject_path, index, A_FOREIGN_VOCABULARY)
    stamped = subject_path.stat().st_mtime_ns

    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    # The pass completed: the rows are all still there, and the sheet was not touched at all.
    assert outcome.refused is None
    assert outcome.workbook_written is False
    assert len(entries) == subject.row_count
    assert subject_path.stat().st_mtime_ns == stamped
    # The cell the dropdown refuses is left blank, and nothing claims to have written over it.
    assert _status_of(subject_path, subject, FIRST_DATA_ROW) == ""
    assert entries[0].mirrored_watch_cell == ""


async def test_text_the_dropdown_offers_but_the_app_cannot_read_appends_no_event(
    reconciler: Reconciler, db_session: AsyncSession, subject: WorkbookFacts, subject_path: Path
) -> None:
    """A status outside the four values is recorded and left alone, never guessed at (section 5.16)."""
    index = subject.index_of(subject.watch_key)
    await reconciler.reconcile(subject.category_id)
    narrow_the_status_dropdown(subject_path, index, A_FOREIGN_VOCABULARY)
    _type_status(subject_path, index, FIRST_DATA_ROW, A_FOREIGN_VOCABULARY[0])

    outcome = await reconciler.reconcile(subject.category_id)

    entries = await _entries(db_session, subject.category_id)
    assert outcome.refused is None
    assert await _mirror_events(db_session, subject.category_id) == []
    assert _status_of(subject_path, subject, FIRST_DATA_ROW) == A_FOREIGN_VOCABULARY[0]
    assert entries[0].observed_watch_cell == A_FOREIGN_VOCABULARY[0]
    assert entries[0].mirrored_watch_cell == ""


class _CountingResolver:
    """A `TitleResolver` stand-in at the resolver seam: it counts what it was asked, and resolves nothing."""

    def __init__(self) -> None:
        self.seen: list[str] = []

    async def resolve(self, request: ResolutionRequest) -> TitleResolution:
        self.seen.append(request.raw_title)
        return TitleResolution(
            state=ResolutionState.UNRESOLVED,
            confidence=Decimal("0.00"),
            note="nothing was searched for",
            chosen=None,
            alternatives=(),
            unit_kind=request.span.unit_kind,
        )


class _NoTitles:
    """A `TmdbIngest` stand-in. It is never reached: the counting resolver never chooses a candidate."""

    async def ensure_title(self, media_type: object, tmdb_id: int) -> None:  # pragma: no cover - never chosen
        return None


def _the_category_filling_only_order_and_title() -> WorkbookFacts:
    """The one workbook whose data rows carry nothing but an order number and a title.

    Chosen by that shape rather than by its name, so a rename or a reorder of the library cannot make
    this test silently assert about a different file.
    """
    found = [facts for facts in every_workbook() if _filled_keys(facts) == {facts.order_key, facts.title_key}]
    if len(found) != 1:
        raise AssertionError(f"expected exactly one workbook filling only order and title, found {len(found)}")
    return found[0]


def _filled_keys(facts: WorkbookFacts) -> set[str]:
    detail = read_category(facts.path)
    return {key for row in detail.rows for key, value in row.cells.items() if value.strip()}
