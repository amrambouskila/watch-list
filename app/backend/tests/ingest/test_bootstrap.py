"""The bootstrap run over the whole library: nineteen workbooks, one at a time, nothing shared.

Two halves. The first needs no database at all - argument parsing, the exit-code rule, the arithmetic
the summary is built from, the slug discovery, and the refusal path a workbook that cannot be reached
takes - so the runnable is not a module that only CI has ever executed. The second is the real thing:
nineteen throwaway copies of the owner's own workbooks projected into a real Postgres, which skips
loudly on a workstation with none rather than standing a mock in for the store the whole phase is
about (CLAUDE.md section 16.4).

Providers are absent throughout. That is the test configuration *and* a valid production one: with no
TMDB token every entry stays unresolved with the reason on it, and the half this file is about -
identity, snapshot, write-back, refusal, serialisation - runs exactly as it does for the owner.
"""

from __future__ import annotations

import asyncio
import logging
import time
import uuid
from collections.abc import AsyncIterator, Callable, Iterator
from contextlib import AbstractAsyncContextManager, asynccontextmanager
from pathlib import Path
from typing import Final

import pytest
from mock_responder import MockResponder
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from backup_support import restore_points
from database_support import JoiningSessionFactory
from library_choice import a_category, every_workbook
from tv_watchlist.config import Settings, get_settings
from tv_watchlist.constants import (
    BOOTSTRAP_EXIT_FAILED,
    BOOTSTRAP_EXIT_OK,
    BOOTSTRAP_EXIT_REFUSED,
    HEADER_ROW,
    IDENTITY_HEADER_LABEL,
)
from tv_watchlist.data.models import Category, CuratedEntry
from tv_watchlist.ingest import bootstrap as bootstrap_module
from tv_watchlist.ingest.bootstrap import (
    BootstrapTotals,
    bootstrap,
    exit_code,
    library_slugs,
    main,
    parse_arguments,
    summarise,
)
from tv_watchlist.ingest.reconcile_outcome import ReconcileOutcome
from tv_watchlist.ingest.runtime import IngestRuntime
from tv_watchlist.models.category_detail import CategoryDetail
from tv_watchlist.services import catalog as catalog_module
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.workbook.reader import read_category, read_projection
from workbook_facts import WorkbookFacts

# Today's library, and the figure the owner's own run is expected to print. The equality against the
# live sums below is the assertion that must hold when a row is added; this literal is the record of
# what the library held when T27 landed, and is updated when it stops matching.
LIBRARY_WORKBOOKS: Final[int] = 19
LIBRARY_ROW_COUNT: Final[int] = 996
A_SLUG_NO_WORKBOOK_ANSWERS_TO: Final[str] = "no-workbook-answers-to-this"
NOT_A_WORKBOOK: Final[bytes] = b"this is not a zip container"
A_REFUSAL: Final[str] = "UnreadableWorkbookError: bad zip"
# Stands in for the one thing an infrastructure failure's message carries that a log may not.
A_CREDENTIAL_A_DSN_WOULD_CARRY: Final[str] = "postgresql://tv:hunter2@localhost:5525/tv"


# -- the offline half: no database, no network, no workbook written ----------------------------------


def test_it_parses_the_whole_library_by_default() -> None:
    arguments = parse_arguments([])

    assert arguments.slugs is None
    assert arguments.dry_run is False


def test_only_narrows_the_run_to_the_slugs_it_names() -> None:
    arguments = parse_arguments(["--only", "marvel", "wwii"])

    assert arguments.slugs == ("marvel", "wwii")


def test_dry_run_is_a_flag() -> None:
    assert parse_arguments(["--dry-run"]).dry_run is True


def test_a_run_with_nothing_refused_exits_zero() -> None:
    outcomes = [ReconcileOutcome(slug="marvel", minted=141), ReconcileOutcome(slug="wwii", skipped=True)]

    assert exit_code(outcomes) == BOOTSTRAP_EXIT_OK


def test_an_empty_run_exits_zero() -> None:
    """Nineteen workbooks and none of them refused is success; so is a library with nothing in it."""
    assert exit_code([]) == BOOTSTRAP_EXIT_OK


def test_one_refusal_anywhere_in_the_run_exits_one() -> None:
    outcomes = [
        ReconcileOutcome(slug="marvel", minted=141),
        ReconcileOutcome(slug="pokemon", refused=A_REFUSAL),
    ]

    assert exit_code(outcomes) == BOOTSTRAP_EXIT_REFUSED


def test_the_totals_separate_ingested_from_skipped_from_refused() -> None:
    outcomes = [
        ReconcileOutcome(slug="marvel", minted=141, workbook_written=True),
        ReconcileOutcome(slug="wwii", skipped=True),
        ReconcileOutcome(slug="pokemon", refused=A_REFUSAL),
        ReconcileOutcome(slug="naruto", matched=10, adopted=2, minted=1, workbook_written=True),
    ]

    totals = BootstrapTotals.of(outcomes)

    assert totals.workbooks == 4
    assert totals.ingested == 2
    assert totals.skipped == 1
    assert totals.refused == 1
    assert totals.minted == 142
    # Every row the pass accounted for, whichever way it accounted for it.
    assert totals.rows == 154
    assert totals.written == 2


def test_the_summary_prints_the_three_figures_the_owner_is_told_to_expect() -> None:
    outcomes = [
        ReconcileOutcome(slug=f"category-{n}", minted=52, workbook_written=True) for n in range(LIBRARY_WORKBOOKS)
    ]

    summary = summarise(outcomes, dry_run=False)

    assert f"{LIBRARY_WORKBOOKS} workbooks" in summary
    assert f"{LIBRARY_WORKBOOKS} ingested" in summary
    assert "988 identities minted" in summary
    assert f"{LIBRARY_WORKBOOKS} workbooks written" in summary


def test_the_summary_names_every_refusal_and_its_reason() -> None:
    outcomes = [
        ReconcileOutcome(slug="marvel", minted=141),
        ReconcileOutcome(slug="pokemon", refused=A_REFUSAL),
    ]

    summary = summarise(outcomes, dry_run=False)

    assert "pokemon" in summary
    assert A_REFUSAL in summary


def test_a_dry_run_summary_says_nothing_was_written() -> None:
    summary = summarise([ReconcileOutcome(slug="marvel", minted=141)], dry_run=True)

    assert "dry run" in summary
    assert "nothing was written" in summary


def test_the_slugs_are_the_ids_every_workbook_in_the_library_answers_to(library: Path) -> None:
    slugs = library_slugs(library)

    assert slugs == sorted(slugs)
    assert len(slugs) == LIBRARY_WORKBOOKS
    assert slugs == [facts.category_id for facts in every_workbook()]


def test_a_shadowed_namesake_is_not_run_twice(library: Path) -> None:
    """Two files can derive one id; `discovery.resolve` answers with the first, so the run does too."""
    subject = a_category()
    shadow = library / f"{subject.stem}-.xlsx"
    shadow.write_bytes(subject.copied_into(library).read_bytes())

    slugs = library_slugs(library)

    assert slugs.count(subject.category_id) == 1
    assert len(slugs) == LIBRARY_WORKBOOKS


def test_an_empty_library_is_a_run_of_nothing_that_still_says_so(
    caplog: pytest.LogCaptureFixture, tmp_path: Path
) -> None:
    """A run that changed nothing must report that it changed nothing (CLAUDE.md section 9.2)."""
    runtime = _offline_runtime(tmp_path / "empty", tmp_path)

    with caplog.at_level(logging.INFO, logger="tv_watchlist.ingest.bootstrap"):
        outcomes = asyncio.run(bootstrap(runtime))

    assert outcomes == []
    records = [record.getMessage() for record in caplog.records]
    assert any(record.startswith("bootstrap_start job=bootstrap dry_run=False workbooks=0") for record in records)
    assert any("bootstrap_end" in record and "workbooks=0 ingested=0" in record for record in records)


def test_a_slug_no_workbook_answers_to_is_recorded_and_the_run_carries_on(
    caplog: pytest.LogCaptureFixture, library: Path, tmp_path: Path
) -> None:
    """The refusal path and the loop past it, proved without a database: neither pass opens a transaction.

    That a refusal is followed by a workbook which really does ingest is the database half of this, in
    `test_one_unreadable_workbook_does_not_stop_the_run`.
    """
    runtime = _offline_runtime(library, tmp_path)
    second = f"{A_SLUG_NO_WORKBOOK_ANSWERS_TO}-either"

    with caplog.at_level(logging.INFO, logger="tv_watchlist.ingest.bootstrap"):
        outcomes = asyncio.run(bootstrap(runtime, slugs=[A_SLUG_NO_WORKBOOK_ANSWERS_TO, second]))

    assert [outcome.slug for outcome in outcomes] == [A_SLUG_NO_WORKBOOK_ANSWERS_TO, second]
    assert all("CategoryNotFoundError" in (outcome.refused or "") for outcome in outcomes)
    assert exit_code(outcomes) == BOOTSTRAP_EXIT_REFUSED
    record = next(
        message
        for message in (record.getMessage() for record in caplog.records)
        if message.startswith(f"bootstrap_workbook job=bootstrap category={A_SLUG_NO_WORKBOOK_ANSWERS_TO}")
    )
    assert "refused=refused" in record
    assert "elapsed_ms=" in record
    # The free-text reason is last, so a message carrying spaces cannot be read as another key.
    assert record.index("reason=") > record.index("elapsed_ms=")


def test_main_reports_a_refusal_as_a_nonzero_exit(
    capsys: pytest.CaptureFixture[str], library: Path, tmp_path: Path, settings_from_environment: None
) -> None:
    """`main` end to end, against a copied library and a slug nothing answers to: no database is reached."""
    assert main(["--only", A_SLUG_NO_WORKBOOK_ANSWERS_TO]) == BOOTSTRAP_EXIT_REFUSED

    printed = capsys.readouterr().out
    assert "1 refused" in printed
    assert A_SLUG_NO_WORKBOOK_ANSWERS_TO in printed


def test_a_run_that_breaks_exits_two_and_names_no_credential(
    capsys: pytest.CaptureFixture[str], caplog: pytest.LogCaptureFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An unreachable database is not a refusal: the shell has to be able to tell the two apart."""

    async def _raise(_: object) -> list[ReconcileOutcome]:
        raise OSError(f"could not connect to {A_CREDENTIAL_A_DSN_WOULD_CARRY}")

    monkeypatch.setattr(bootstrap_module, "_run", _raise)

    with caplog.at_level(logging.ERROR, logger="tv_watchlist.ingest.bootstrap"):
        code = main([])

    assert code == BOOTSTRAP_EXIT_FAILED
    assert code not in (BOOTSTRAP_EXIT_OK, BOOTSTRAP_EXIT_REFUSED)
    record = next(
        message for message in (found.getMessage() for found in caplog.records) if "bootstrap_failed" in message
    )
    assert "job=bootstrap" in record
    assert "outcome=failed" in record
    assert "error=OSError" in record
    printed = capsys.readouterr().out
    # The exception's message is what carries a DSN or a header; only its class may be reported.
    assert A_CREDENTIAL_A_DSN_WOULD_CARRY not in record
    assert A_CREDENTIAL_A_DSN_WOULD_CARRY not in printed
    assert "OSError" in printed


def test_a_refusal_carrying_a_newline_stays_one_log_record(
    caplog: pytest.LogCaptureFixture, library: Path, tmp_path: Path
) -> None:
    """A reason reaches us from a filename, a cell or a provider; a newline in it must not forge a record."""
    runtime = _offline_runtime(library, tmp_path)
    outcome = ReconcileOutcome(slug="marvel", refused="UnreadableWorkbookError: bad zip\nbootstrap_end job=forged")

    with caplog.at_level(logging.INFO, logger="tv_watchlist.ingest.bootstrap"):
        bootstrap_module._log_workbook(outcome, time.monotonic(), dry_run=False)

    record = caplog.records[-1].getMessage()
    assert "\n" not in record
    assert "\r" not in record
    assert record.count("bootstrap_workbook") == 1
    assert "bad zip bootstrap_end job=forged" in record
    assert runtime.catalog.library_dir == library


@pytest.fixture
def settings_from_environment(library: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """Point `get_settings()` at the throwaway library, so `main` can never reach the owner's own files."""
    monkeypatch.setenv("TV_LIBRARY_DIR", str(library))
    monkeypatch.setenv("TV_BACKUP_DIR", str(tmp_path / "backups"))
    monkeypatch.setenv("TV_HEROES_DIR", str(tmp_path / "heroes"))
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def _offline_runtime(library_dir: Path, tmp_path: Path) -> IngestRuntime:
    """A runtime whose sessions are never executed against and whose transports reach no network."""
    settings = Settings(
        _env_file=None, library_dir=library_dir, backup_dir=tmp_path / "backups", heroes_dir=tmp_path / "heroes"
    )
    return IngestRuntime.from_settings(
        settings,
        Catalog(settings),
        async_sessionmaker(expire_on_commit=False, autoflush=False),
        transport=MockResponder().transport,
    )


# -- the real run: nineteen workbooks, a real Postgres, real .xlsx writes -----------------------------


@pytest.fixture
def runtime(settings: Settings, catalog: Catalog, db_session: AsyncSession) -> IngestRuntime:
    """Providers unconfigured, sessions joined to the test's transaction, transports mocked shut."""
    return IngestRuntime.from_settings(
        settings, catalog, JoiningSessionFactory(db_session), transport=MockResponder().transport
    )


def _grid(path: Path) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """What the UI renders: the public columns and the rows, with nothing else about the file."""
    detail: CategoryDetail = read_category(path)
    return ([column.model_dump() for column in detail.columns], [row.model_dump() for row in detail.rows])


def _identities(path: Path) -> dict[int, str]:
    return read_projection(path).row_ids


def _copies(library: Path) -> list[tuple[WorkbookFacts, Path]]:
    return [(facts, facts.copied_into(library)) for facts in every_workbook()]


async def test_bootstrap_ingests_all_nineteen_workbooks(
    runtime: IngestRuntime, db_session: AsyncSession, library: Path
) -> None:
    outcomes = await bootstrap(runtime)

    assert len(outcomes) == LIBRARY_WORKBOOKS
    assert [outcome.slug for outcome in outcomes] == library_slugs(library)
    assert [outcome.refused for outcome in outcomes] == [None] * LIBRARY_WORKBOOKS
    assert await db_session.scalar(select(func.count()).select_from(Category)) == LIBRARY_WORKBOOKS


async def test_every_row_carries_a_uuid_and_the_total_is_the_library_row_count(
    runtime: IngestRuntime, db_session: AsyncSession, library: Path
) -> None:
    await bootstrap(runtime)

    identities = {path: _identities(path) for _, path in _copies(library)}
    rows = {path: read_category(path).counts.total for _, path in _copies(library)}
    minted = sum(len(found) for found in identities.values())
    assert {path: len(found) for path, found in identities.items()} == rows
    assert minted == sum(rows.values())
    assert all(uuid.UUID(value) for found in identities.values() for value in found.values())
    assert await db_session.scalar(select(func.count()).select_from(CuratedEntry)) == minted
    assert minted == LIBRARY_ROW_COUNT, f"the library now holds {minted} rows; update LIBRARY_ROW_COUNT"


async def test_every_workbook_gained_exactly_one_hidden_identity_column(runtime: IngestRuntime, library: Path) -> None:
    """One loop rather than one parametrised case per workbook: the run under test is the whole library,
    and parametrising would bootstrap all nineteen workbooks nineteen times over."""
    await bootstrap(runtime)

    for facts, path in _copies(library):
        sheet = load_workbook(path).worksheets[0]
        headers = [sheet.cell(row=HEADER_ROW, column=index).value for index in range(1, sheet.max_column + 1)]
        assert headers.count(IDENTITY_HEADER_LABEL) == 1, facts.category_id
        index = headers.index(IDENTITY_HEADER_LABEL) + 1
        assert sheet.column_dimensions[get_column_letter(index)].hidden is True, facts.category_id


async def test_nineteen_snapshots_exist_in_backups_one_per_workbook(
    runtime: IngestRuntime, settings: Settings, library: Path
) -> None:
    await bootstrap(runtime)

    assert {facts.category_id: len(restore_points(settings, facts)) for facts, _ in _copies(library)} == {
        facts.category_id: 1 for facts, _ in _copies(library)
    }


async def test_every_workbook_still_opens_and_renders_the_same_grid(runtime: IngestRuntime, library: Path) -> None:
    before = {path: _grid(path) for _, path in _copies(library)}

    await bootstrap(runtime)

    assert {path: _grid(path) for _, path in _copies(library)} == before


async def test_a_second_bootstrap_is_a_no_op(runtime: IngestRuntime, settings: Settings, library: Path) -> None:
    await bootstrap(runtime)
    stamps = {path: path.stat().st_mtime for _, path in _copies(library)}

    outcomes = await bootstrap(runtime)

    assert [outcome.skipped for outcome in outcomes] == [True] * LIBRARY_WORKBOOKS
    assert {path: path.stat().st_mtime for _, path in _copies(library)} == stamps
    assert all(len(restore_points(settings, facts)) == 1 for facts, _ in _copies(library))


async def test_one_unreadable_workbook_does_not_stop_the_run(runtime: IngestRuntime, library: Path) -> None:
    subject = a_category()
    subject.copied_into(library).write_bytes(NOT_A_WORKBOOK)

    outcomes = await bootstrap(runtime)

    refused = [outcome for outcome in outcomes if outcome.refused is not None]
    assert [outcome.slug for outcome in refused] == [subject.category_id]
    assert "UnreadableWorkbookError" in (refused[0].refused or "")
    assert len(outcomes) == LIBRARY_WORKBOOKS
    assert sum(1 for outcome in outcomes if outcome.refused is None) == LIBRARY_WORKBOOKS - 1
    assert exit_code(outcomes) == BOOTSTRAP_EXIT_REFUSED


async def test_dry_run_writes_nothing_to_either_store(
    runtime: IngestRuntime, db_session: AsyncSession, settings: Settings, library: Path
) -> None:
    stamps = {path: path.stat().st_mtime for _, path in _copies(library)}

    outcomes = await bootstrap(runtime, dry_run=True)

    assert len(outcomes) == LIBRARY_WORKBOOKS
    # Computed from the copies rather than from the pinned literal: this test is about a dry run
    # reporting a real run's counts, and `LIBRARY_ROW_COUNT` is pinned once, in the row-count test,
    # where the live equalities are asserted before it.
    assert BootstrapTotals.of(outcomes).minted == sum(read_category(path).counts.total for _, path in _copies(library))
    assert all(outcome.workbook_written is False for outcome in outcomes)
    assert {path: path.stat().st_mtime for _, path in _copies(library)} == stamps
    assert all(_identities(path) == {} for _, path in _copies(library))
    assert not any(restore_points(settings, facts) for facts, _ in _copies(library))
    assert await db_session.scalar(select(func.count()).select_from(CuratedEntry)) == 0
    assert await db_session.scalar(select(func.count()).select_from(Category)) == 0


async def test_the_run_is_serial_and_holds_one_workbook_lock_at_a_time(
    runtime: IngestRuntime, library: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tap = _LockTap(catalog_module.lock_for)
    monkeypatch.setattr(catalog_module, "lock_for", tap)

    await bootstrap(runtime)

    assert tap.peak == 1
    assert tap.held == 0
    assert len(set(tap.paths)) == LIBRARY_WORKBOOKS


class _LockTap:
    """A `lock_for` stand-in that counts how many workbook locks are held at the same moment."""

    def __init__(self, real: Callable[[Path], asyncio.Lock]) -> None:
        self._real = real
        self.held = 0
        self.peak = 0
        self.paths: list[Path] = []

    def __call__(self, path: Path) -> AbstractAsyncContextManager[None]:
        return self._counted(path)

    @asynccontextmanager
    async def _counted(self, path: Path) -> AsyncIterator[None]:
        async with self._real(path):
            self.held += 1
            self.peak = max(self.peak, self.held)
            self.paths.append(path)
            try:
                yield
            finally:
                self.held -= 1
