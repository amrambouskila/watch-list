"""The hidden Row ID column: where it lands, what it preserves, and what it refuses."""

from __future__ import annotations

from pathlib import Path
from typing import Final
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from backup_support import restore_points
from library_choice import every_workbook
from snapshot import fingerprint
from tv_watchlist.config import Settings
from tv_watchlist.constants import (
    FIRST_DATA_ROW,
    HEADER_ROW,
    IDENTITY_COLUMN_WIDTH,
    IDENTITY_HEADER_LABEL,
)
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.workbook import writer
from tv_watchlist.workbook.cells import cell_text
from tv_watchlist.workbook.errors import RowNotFoundError, WorkbookLockedError
from tv_watchlist.workbook.identity import ensure_identity_column, write_row_ids
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.naming import slugify
from tv_watchlist.workbook.reader import read_projection
from tv_watchlist.workbook.schema import build_columns, header_width
from workbook_facts import WorkbookFacts, range_ends

CANONICAL_UUID_LENGTH: Final[int] = 36
IDENTITY_COLUMN_KEY: Final[str] = slugify(IDENTITY_HEADER_LABEL)


def _grid_values(path: Path, width: int) -> dict[tuple[int, int], str]:
    """Every cell of the sheet's original columns, so a write past them can be shown harmless."""
    sheet = load_workbook(path).worksheets[0]
    return {
        (row, column): cell_text(sheet.cell(row=row, column=column).value)
        for row in range(HEADER_ROW, sheet.max_row + 1)
        for column in range(1, width + 1)
    }


def _minted(facts: WorkbookFacts) -> dict[int, str]:
    return {row: str(uuid4()) for row in range(FIRST_DATA_ROW, facts.last_data_row + 1)}


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_the_identity_column_lands_at_the_far_right_hidden_at_width_38(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    index = header_width(load_workbook(path).worksheets[0]) + 1

    ensure_identity_column(path)

    sheet = load_workbook(path).worksheets[0]
    dimension = sheet.column_dimensions[get_column_letter(index)]
    assert header_width(sheet) == index
    assert sheet.cell(row=HEADER_ROW, column=index).value == IDENTITY_HEADER_LABEL
    assert dimension.hidden is True
    assert dimension.width == IDENTITY_COLUMN_WIDTH


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_the_fingerprint_is_unchanged_apart_from_the_new_column(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    letter = get_column_letter(facts.column_count + 1)
    before = fingerprint(path)

    ensure_identity_column(path)

    after = fingerprint(path)
    sheet = load_workbook(path).worksheets[0]
    order = sheet.title
    max_row, max_column = before[order]["dimensions"]
    assert {title: entry for title, entry in after.items() if title != order} == {
        title: entry for title, entry in before.items() if title != order
    }
    assert after[order]["dimensions"] == (max_row, max_column + 1)
    assert after[order]["widths"] == before[order]["widths"] | {letter: IDENTITY_COLUMN_WIDTH}
    assert after[order]["header_style"] == before[order]["header_style"]
    assert after[order]["body_style"] == before[order]["body_style"]
    assert after[order]["conditional_rule_count"] == before[order]["conditional_rule_count"]
    assert after[order]["tables"][facts.grid_table_name] == f"A{HEADER_ROW}:{letter}{facts.last_data_row}"
    assert after[order]["table_columns"][facts.grid_table_name] == [
        *before[order]["table_columns"][facts.grid_table_name],
        IDENTITY_HEADER_LABEL,
    ]
    if facts.ranges_all_end_together:
        assert after[order]["validations"] == before[order]["validations"]
        assert after[order]["conditional_ranges"] == before[order]["conditional_ranges"]
    else:
        # The one workbook whose constructs stop short of its grid is stretched to it, exactly as
        # every other write through `ranges.sync` already does.
        assert range_ends(sheet) == (facts.last_data_row,)


def test_ranges_cover_the_new_width(subject: WorkbookFacts, subject_path: Path) -> None:
    letter = get_column_letter(subject.column_count + 1)

    ensure_identity_column(subject_path)

    sheet = load_workbook(subject_path).worksheets[0]
    assert range_ends(sheet) == (subject.last_data_row,)
    assert sheet.tables[subject.grid_table_name].ref.endswith(f"{letter}{subject.last_data_row}")


def test_ensure_is_idempotent_and_skips_the_save_when_present(subject_path: Path) -> None:
    ensure_identity_column(subject_path)
    stamped = subject_path.stat().st_mtime_ns
    width = header_width(load_workbook(subject_path).worksheets[0])

    ensure_identity_column(subject_path)

    assert subject_path.stat().st_mtime_ns == stamped
    assert header_width(load_workbook(subject_path).worksheets[0]) == width


def test_write_row_ids_writes_canonical_36_char_strings_and_touches_no_other_cell(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    minted = _minted(subject)
    ensure_identity_column(subject_path)
    with_column = fingerprint(subject_path)
    untouched = _grid_values(subject_path, subject.column_count)

    write_row_ids(subject_path, minted)

    projection = read_projection(subject_path)
    recorded = projection.row_ids
    assert projection.identity_key == IDENTITY_COLUMN_KEY
    assert recorded == minted
    assert {len(value) for value in recorded.values()} == {CANONICAL_UUID_LENGTH}
    assert _grid_values(subject_path, subject.column_count) == untouched
    assert fingerprint(subject_path) == with_column


def test_a_row_outside_the_grid_is_refused(subject: WorkbookFacts, subject_path: Path) -> None:
    ensure_identity_column(subject_path)
    untouched = subject_path.read_bytes()

    for row in (HEADER_ROW, subject.last_data_row + 1):
        with pytest.raises(RowNotFoundError):
            write_row_ids(subject_path, {row: str(uuid4())})

    assert subject_path.read_bytes() == untouched


def test_the_projection_reads_blanks_as_empty_strings_and_no_key_before_the_column_exists(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    bare = read_projection(subject_path)
    assert (bare.identity_key, bare.row_ids) == (None, {})

    ensure_identity_column(subject_path)

    projection = read_projection(subject_path)
    assert projection.identity_key == IDENTITY_COLUMN_KEY
    # The projection reports the rows it renders, so a blank row the grid skips is not an id to reconcile.
    assert set(projection.row_ids) == {row.row for row in projection.detail.rows}
    assert set(projection.row_ids) == set(range(FIRST_DATA_ROW, subject.last_data_row + 1))
    assert set(projection.row_ids.values()) == {""}


def test_the_header_is_styled_on_a_sheet_with_no_columns_to_copy(tmp_path: Path) -> None:
    path = tmp_path / "Bare_Master_Watch_Order.xlsx"
    workbook = Workbook()
    workbook.active.cell(row=FIRST_DATA_ROW, column=1).value = "orphan row"
    workbook.save(path)

    ensure_identity_column(path)

    sheet = load_workbook(path).worksheets[0]
    assert sheet.cell(row=HEADER_ROW, column=2).value == IDENTITY_HEADER_LABEL
    assert sheet.cell(row=HEADER_ROW, column=2).font.bold is True


def test_the_writer_types_the_identity_column_as_text(subject: WorkbookFacts, subject_path: Path) -> None:
    """A canonical UUID must reach the cell as text, not as something Excel tries to read as a number."""
    ensure_identity_column(subject_path)
    write_row_ids(subject_path, _minted(subject))

    sheet = load_workbook(subject_path).worksheets[0]
    numeric = writer._numeric_columns(sheet, build_columns(sheet), FIRST_DATA_ROW, subject.last_data_row)
    assert IDENTITY_COLUMN_KEY not in numeric


def test_an_excel_lock_refuses_the_write(subject: WorkbookFacts, subject_path: Path) -> None:
    excel_lock_path(subject_path).write_bytes(b"owner")
    untouched = subject_path.read_bytes()

    with pytest.raises(WorkbookLockedError):
        ensure_identity_column(subject_path)
    with pytest.raises(WorkbookLockedError):
        write_row_ids(subject_path, {FIRST_DATA_ROW: str(uuid4())})

    assert subject_path.read_bytes() == untouched


async def test_the_catalog_wrappers_snapshot_unconditionally(
    catalog: Catalog, settings: Settings, subject: WorkbookFacts, subject_path: Path
) -> None:
    untouched = subject_path.read_bytes()
    detail = await catalog.detail(subject.category_id)

    with_column = await catalog.ensure_identity_column(subject.category_id, detail.mtime)
    before_ids = subject_path.read_bytes()
    await catalog.write_row_ids(subject.category_id, _minted(subject), with_column.mtime)

    assert restore_points(settings, subject) == [untouched, before_ids]
