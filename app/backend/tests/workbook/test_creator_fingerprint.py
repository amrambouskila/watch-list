from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from library_choice import a_name_the_library_does_not_hold, every_workbook
from tv_watchlist.constants import FIRST_DATA_ROW, TABLE_STYLE_NAME, WATCHED
from tv_watchlist.models.category_create import CategoryCreate
from tv_watchlist.workbook.creator import create_category
from tv_watchlist.workbook.schema import build_columns, watch_column
from workbook_facts import WorkbookFacts, grid_table_name

BANDING_CATEGORY = a_name_the_library_does_not_hold("Banding Check")
BANDING_TABLE = "BandingCheckMasterOrder"
SHAPE_CATEGORY = a_name_the_library_does_not_hold("Shape Check")
SEED_TITLE = "Seed"


def _sheet(path: Path) -> Worksheet:
    return load_workbook(path).worksheets[0]


def _table_style(sheet: Worksheet) -> tuple[Any, ...]:
    style = sheet.tables[grid_table_name(sheet)].tableStyleInfo
    return (style.name, style.showRowStripes, style.showColumnStripes, style.showFirstColumn, style.showLastColumn)


def _watch_letter(sheet: Worksheet) -> str:
    watch = watch_column(build_columns(sheet))
    assert watch is not None
    return get_column_letter(watch.index)


def _watch_dropdown(sheet: Worksheet) -> tuple[Any, ...]:
    letter = _watch_letter(sheet)
    validation = next(v for v in sheet.data_validations.dataValidation if str(v.sqref).startswith(letter))
    return (validation.type, validation.formula1, validation.allowBlank, validation.showDropDown)


def _watched_highlight(sheet: Worksheet) -> tuple[Any, ...]:
    letter = _watch_letter(sheet)
    entry = next(e for e in sheet.conditional_formatting if str(e.sqref).startswith(letter))
    rule = entry.rules[0]
    return (rule.type, rule.formula, rule.dxf.fill.bgColor.rgb)


def _header_style(sheet: Worksheet) -> tuple[Any, ...]:
    cell = sheet.cell(row=1, column=1)
    return (cell.font.bold, cell.font.name, cell.font.sz, cell.font.color.rgb, cell.fill.patternType)


def test_a_created_table_carries_the_library_banding(library: Path) -> None:
    path = create_category(library, CategoryCreate(name=BANDING_CATEGORY))

    style = _sheet(path).tables[BANDING_TABLE].tableStyleInfo
    assert style is not None
    assert style.name == TABLE_STYLE_NAME
    assert style.showRowStripes is True


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_a_created_workbook_is_shaped_like_a_hand_built_one(library: Path, facts: WorkbookFacts) -> None:
    made = _sheet(create_category(library, CategoryCreate(name=SHAPE_CATEGORY, titles=[SEED_TITLE])))
    hand = _sheet(facts.copied_into(library))

    assert _table_style(made) == _table_style(hand)
    assert _watch_dropdown(made) == _watch_dropdown(hand)
    assert _header_style(made) == _header_style(hand)

    made_type, made_formula, made_fill = _watched_highlight(made)
    hand_type, hand_formula, hand_fill = _watched_highlight(hand)
    assert (made_type, made_fill) == (hand_type, hand_fill)
    assert made_formula == [f'{_watch_letter(made)}{FIRST_DATA_ROW}="{WATCHED}"']
    assert hand_formula == [f'{_watch_letter(hand)}{FIRST_DATA_ROW}="{WATCHED}"']
