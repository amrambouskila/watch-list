"""Player snapshots preserve real workbook structure and hide app-owned positions from editing."""

from __future__ import annotations

from copy import copy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from library_choice import every_workbook
from snapshot import fingerprint
from tv_watchlist.constants import FIRST_DATA_ROW, HEADER_ROW, POSITION_COLUMN_WIDTH, POSITION_HEADER_LABEL, WATCHED
from tv_watchlist.ingest.workbook_digest import digest_grid
from tv_watchlist.models.category_create import CategoryCreate
from tv_watchlist.models.row_change import RowChange
from tv_watchlist.workbook import writer
from tv_watchlist.workbook.errors import InvalidChoiceError, RowNotFoundError, UnknownColumnError, WorkbookLockedError
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.naming import slugify
from tv_watchlist.workbook.position import apply_position_writeback
from tv_watchlist.workbook.reader import read_projection
from tv_watchlist.workbook.schema import _role_for, build_columns, header_width, position_column
from workbook_facts import WorkbookFacts, range_ends

POSITION_SECONDS = 61
POSITION_KEY = slugify(POSITION_HEADER_LABEL)


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_position_is_visible_numeric_and_preserves_every_workbook(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    before = fingerprint(path)
    public_before = read_projection(path).detail
    apply_position_writeback(path, {FIRST_DATA_ROW: POSITION_SECONDS}, {})
    after = fingerprint(path)
    sheet = load_workbook(path).worksheets[0]
    column = position_column(build_columns(sheet))
    assert column is not None
    assert column.index == facts.column_count + 1
    letter = get_column_letter(column.index)
    assert not sheet.column_dimensions[letter].hidden
    assert sheet.column_dimensions[letter].width == pytest.approx(POSITION_COLUMN_WIDTH)
    assert sheet.cell(FIRST_DATA_ROW, column.index).value == POSITION_SECONDS
    assert sheet.cell(FIRST_DATA_ROW, column.index).data_type == "n"
    for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
        assert copy(getattr(sheet.cell(HEADER_ROW, column.index), attribute)) == copy(
            getattr(sheet.cell(HEADER_ROW, 1), attribute)
        )
    assert after[sheet.title]["header_style"] == before[sheet.title]["header_style"]
    assert after[sheet.title]["body_style"] == before[sheet.title]["body_style"]
    assert after[sheet.title]["conditional_rule_count"] == before[sheet.title]["conditional_rule_count"]
    assert after[sheet.title]["widths"] == before[sheet.title]["widths"] | {letter: POSITION_COLUMN_WIDTH}
    assert after[sheet.title]["dimensions"] == (before[sheet.title]["dimensions"][0], column.index)
    assert after[sheet.title]["table_columns"][facts.grid_table_name] == [
        *before[sheet.title]["table_columns"][facts.grid_table_name],
        POSITION_HEADER_LABEL,
    ]
    assert sheet.tables[facts.grid_table_name].ref == f"A1:{letter}{facts.last_data_row}"
    assert range_ends(sheet) == (facts.last_data_row,)
    for name in before:
        if name != sheet.title:
            assert after[name] == before[name]
    if facts.ranges_all_end_together:
        assert after[sheet.title]["validations"] == before[sheet.title]["validations"]
        assert after[sheet.title]["conditional_ranges"] == before[sheet.title]["conditional_ranges"]
    projection = read_projection(path)
    assert projection.detail.columns == public_before.columns
    assert projection.detail.rows == public_before.rows
    assert projection.detail.counts == public_before.counts
    assert projection.position_values[FIRST_DATA_ROW] == str(POSITION_SECONDS)
    assert digest_grid(projection.detail.columns, projection.detail.rows) == digest_grid(
        public_before.columns, public_before.rows
    )


def test_second_snapshot_reuses_column_and_combines_watch_cells(subject_path: Path, subject: WorkbookFacts) -> None:
    apply_position_writeback(subject_path, {FIRST_DATA_ROW: POSITION_SECONDS}, {})
    before = fingerprint(subject_path)
    apply_position_writeback(subject_path, {FIRST_DATA_ROW: 0}, {FIRST_DATA_ROW: {subject.watch_key: WATCHED}})
    projection = read_projection(subject_path)
    assert projection.position_values[FIRST_DATA_ROW] == "0"
    assert projection.detail.rows[0].cells[subject.watch_key] == WATCHED
    assert fingerprint(subject_path) == before


@pytest.mark.parametrize("label", ["Position (s)", "position s", "POSITION (S)"])
def test_position_alias_is_app_owned_and_refused_on_category_create(label: str) -> None:
    assert _role_for(label) == "position"
    with pytest.raises(ValidationError, match="column the app owns"):
        CategoryCreate(name="Position collision", columns=[label])


def test_blank_header_ends_grid_even_with_stray_cells(subject_path: Path, subject: WorkbookFacts) -> None:
    workbook = load_workbook(subject_path)
    workbook.worksheets[0].cell(HEADER_ROW, subject.column_count + 3).value = "Stray cell"
    workbook.save(subject_path)
    apply_position_writeback(subject_path, {FIRST_DATA_ROW: POSITION_SECONDS}, {})
    sheet = load_workbook(subject_path).worksheets[0]
    assert header_width(sheet) == subject.column_count + 1
    assert sheet.cell(HEADER_ROW, subject.column_count + 3).value == "Stray cell"


@pytest.mark.parametrize("target", ["position", "watch"])
@pytest.mark.parametrize("outside", ["header", "after_grid"])
def test_any_out_of_grid_row_refuses_the_entire_snapshot(
    subject_path: Path, subject: WorkbookFacts, target: str, outside: str
) -> None:
    row = HEADER_ROW if outside == "header" else subject.last_data_row + 1
    positions = {row if target == "position" else FIRST_DATA_ROW: POSITION_SECONDS}
    cells = {row if target == "watch" else FIRST_DATA_ROW: {subject.watch_key: WATCHED}}
    before = subject_path.read_bytes()
    with pytest.raises(RowNotFoundError):
        apply_position_writeback(subject_path, positions, cells)
    assert subject_path.read_bytes() == before


def test_invalid_watch_value_refuses_position_too(subject_path: Path, subject: WorkbookFacts) -> None:
    before = subject_path.read_bytes()
    with pytest.raises(InvalidChoiceError):
        apply_position_writeback(
            subject_path, {FIRST_DATA_ROW: POSITION_SECONDS}, {FIRST_DATA_ROW: {subject.watch_key: "invalid"}}
        )
    assert subject_path.read_bytes() == before


@pytest.mark.parametrize("method", ["update", "append", "batch", "snapshot"])
def test_public_writes_cannot_address_position(subject_path: Path, method: str) -> None:
    apply_position_writeback(subject_path, {FIRST_DATA_ROW: POSITION_SECONDS}, {})
    before = subject_path.read_bytes()
    cells = {POSITION_KEY: "=1+1"}
    with pytest.raises(UnknownColumnError):
        if method == "update":
            writer.update_row(subject_path, FIRST_DATA_ROW, cells)
        elif method == "append":
            writer.append_row(subject_path, cells)
        elif method == "batch":
            writer.apply_changes(subject_path, [RowChange(kind="revise", row=FIRST_DATA_ROW, cells=cells)])
        else:
            apply_position_writeback(subject_path, {}, {FIRST_DATA_ROW: cells})
    assert subject_path.read_bytes() == before


def test_locked_file_refuses_snapshot(subject_path: Path) -> None:
    before = subject_path.read_bytes()
    excel_lock_path(subject_path).write_bytes(b"owner")
    with pytest.raises(WorkbookLockedError):
        apply_position_writeback(subject_path, {FIRST_DATA_ROW: POSITION_SECONDS}, {})
    assert subject_path.read_bytes() == before


def test_no_position_column_means_no_internal_position_values(subject_path: Path) -> None:
    assert read_projection(subject_path).position_values == {}
