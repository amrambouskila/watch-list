"""The identity role: what the reader hides, what a row write may never address, what the wire carries."""

from __future__ import annotations

from collections.abc import AsyncIterator
from pathlib import Path
from typing import Final, get_args
from uuid import uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from library_choice import every_workbook
from tv_watchlist.api.dependencies import get_catalog
from tv_watchlist.constants import (
    APP_OWNED_COLUMN_ROLES,
    FIRST_DATA_ROW,
    IDENTITY_COLUMN_ROLE,
    IDENTITY_HEADER_LABEL,
)
from tv_watchlist.main import create_app
from tv_watchlist.models.category_detail import CategoryDetail
from tv_watchlist.models.column_spec import ColumnRole, ColumnSpec
from tv_watchlist.models.row_change import RowChange
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.workbook.errors import UnknownColumnError, UnreadableWorkbookError
from tv_watchlist.workbook.identity import ensure_identity_column, write_row_ids
from tv_watchlist.workbook.naming import category_id, slugify
from tv_watchlist.workbook.reader import read_category, read_projection
from tv_watchlist.workbook.schema import _role_for, build_columns, public_columns
from workbook_facts import WorkbookFacts

IDENTITY_COLUMN_KEY: Final[str] = slugify(IDENTITY_HEADER_LABEL)
# The TypeScript mirror of ColumnRole: the two unions are one contract, so a test reads the file itself.
COLUMN_SPEC_TS: Final[Path] = Path(__file__).resolve().parents[3] / "frontend" / "src" / "types" / "ColumnSpec.ts"
A_ROW_ID: Final[str] = "0f0d3a5e-0000-4000-8000-000000000001"
AN_ID_NO_WRITE_SHOULD_LAND: Final[str] = "an id the owner never asked for"
NOT_A_WORKBOOK: Final[bytes] = b"not a workbook"
EVERY_ROLE: Final[tuple[ColumnRole, ...]] = ("order", "title", "watch", "identity", "other", "position")


@pytest.fixture
async def client(catalog: Catalog) -> AsyncIterator[AsyncClient]:
    app = create_app()
    app.dependency_overrides[get_catalog] = lambda: catalog
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as async_client:
        yield async_client


def _spec(index: int, role: ColumnRole) -> ColumnSpec:
    return ColumnSpec(key=f"column-{index}", label=f"Column {index}", index=index, kind="text", role=role)


def _grid(detail: CategoryDetail) -> str:
    """The columns and rows exactly as they cross the wire."""
    return detail.model_dump_json(include={"columns": True, "rows": True})


def _minted(facts: WorkbookFacts) -> dict[int, str]:
    return {row: str(uuid4()) for row in range(FIRST_DATA_ROW, facts.last_data_row + 1)}


@pytest.mark.parametrize(
    ("label", "role"),
    [
        ("Row ID", "identity"),
        ("RowID", "identity"),
        ("row id", "identity"),
        ("Row-ID", "identity"),
        ("Rows", "other"),
        ("Watched?", "watch"),
        ("Title", "title"),
        ("#", "order"),
    ],
)
def test_row_id_derives_the_identity_role(label: str, role: ColumnRole) -> None:
    assert _role_for(label) == role


def test_public_columns_drops_only_app_owned_roles() -> None:
    columns = [_spec(index, role) for index, role in enumerate(EVERY_ROLE, start=1)]

    public = public_columns(columns)

    assert [column.role for column in public] == ["order", "title", "watch", "other"]
    assert {column.role for column in columns} - {column.role for column in public} == APP_OWNED_COLUMN_ROLES


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_the_reader_hides_the_identity_column_from_columns_and_cells(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    ensure_identity_column(path)
    write_row_ids(path, _minted(facts))

    detail = read_category(path)

    assert build_columns(load_workbook(path).worksheets[0])[-1].role == IDENTITY_COLUMN_ROLE
    assert [column.role for column in detail.columns if column.role == IDENTITY_COLUMN_ROLE] == []
    assert [row.row for row in detail.rows if IDENTITY_COLUMN_KEY in row.cells] == []


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_the_grid_the_ui_renders_is_unchanged_by_the_column(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    before = read_category(path)

    ensure_identity_column(path)
    write_row_ids(path, _minted(facts))

    after = read_category(path)
    assert _grid(after) == _grid(before)
    assert after.counts == before.counts
    assert after.has_watch_column == before.has_watch_column


async def test_the_projection_carries_the_ids_the_reader_hides(
    catalog: Catalog, subject: WorkbookFacts, subject_path: Path
) -> None:
    minted = _minted(subject)
    write_row_ids(subject_path, minted)

    projection = read_projection(subject_path)

    assert projection.identity_key == IDENTITY_COLUMN_KEY
    assert projection.watch_key == subject.watch_key
    assert projection.row_ids == minted
    assert projection.detail == read_category(subject_path)
    assert await catalog.projection(subject.category_id) == projection


def test_the_projection_reports_no_identity_before_the_column_exists(subject_path: Path) -> None:
    projection = read_projection(subject_path)

    assert projection.identity_key is None
    assert projection.row_ids == {}


async def test_a_workbook_that_will_not_open_is_an_unreadable_error_not_a_raw_zip_failure(
    catalog: Catalog, subject_path: Path
) -> None:
    subject_path.write_bytes(NOT_A_WORKBOOK)

    with pytest.raises(UnreadableWorkbookError, match=subject_path.name):
        await catalog.projection(category_id(subject_path))


async def test_a_row_write_cannot_address_the_identity_column(
    client: AsyncClient, subject: WorkbookFacts, subject_path: Path
) -> None:
    write_row_ids(subject_path, {FIRST_DATA_ROW: A_ROW_ID})
    detail = (await client.get(f"/api/categories/{subject.category_id}")).json()

    patched = await client.patch(
        f"/api/categories/{subject.category_id}/rows/{FIRST_DATA_ROW}",
        json={"cells": {IDENTITY_COLUMN_KEY: AN_ID_NO_WRITE_SHOULD_LAND}, "expected_mtime": detail["mtime"]},
    )
    appended = await client.post(
        f"/api/categories/{subject.category_id}/rows",
        json={"cells": {IDENTITY_COLUMN_KEY: AN_ID_NO_WRITE_SHOULD_LAND}, "expected_mtime": detail["mtime"]},
    )

    assert [patched.status_code, appended.status_code] == [400, 400]
    assert {patched.json()["error"], appended.json()["error"]} == {"UnknownColumnError"}
    assert read_projection(subject_path).row_ids[FIRST_DATA_ROW] == A_ROW_ID


async def test_a_batch_change_cannot_address_the_identity_column(
    catalog: Catalog, subject: WorkbookFacts, subject_path: Path
) -> None:
    write_row_ids(subject_path, {FIRST_DATA_ROW: A_ROW_ID})
    detail = await catalog.detail(subject.category_id)

    with pytest.raises(UnknownColumnError):
        await catalog.apply_changes(
            subject.category_id,
            [RowChange(kind="revise", row=FIRST_DATA_ROW, cells={IDENTITY_COLUMN_KEY: AN_ID_NO_WRITE_SHOULD_LAND})],
            detail.mtime,
        )

    assert read_projection(subject_path).row_ids[FIRST_DATA_ROW] == A_ROW_ID


def test_the_writer_still_sees_the_identity_column_for_its_own_writes(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    ensure_identity_column(subject_path)
    columns = build_columns(load_workbook(subject_path).worksheets[0])

    write_row_ids(subject_path, {FIRST_DATA_ROW: A_ROW_ID})

    assert columns[-1].role == IDENTITY_COLUMN_ROLE
    assert len(columns) == len(public_columns(columns)) + 1
    assert read_projection(subject_path).row_ids[FIRST_DATA_ROW] == A_ROW_ID


def test_the_typescript_union_mirrors_the_python_literal() -> None:
    declared = " | ".join(f'"{value}"' for value in get_args(ColumnRole))
    lines = COLUMN_SPEC_TS.read_text(encoding="utf-8").splitlines()

    assert f"export type ColumnRole = {declared};" in lines
