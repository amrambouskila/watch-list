from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from library_choice import every_workbook
from tv_watchlist.constants import FIRST_DATA_ROW
from tv_watchlist.models.row_change import RowChange
from tv_watchlist.workbook.errors import InvalidChoiceError, RowNotFoundError, UnknownColumnError
from tv_watchlist.workbook.reader import read_category
from tv_watchlist.workbook.writer import apply_changes
from workbook_facts import WorkbookFacts, range_ends

BEYOND_THE_LAST_ROW = 9999
A_COLUMN_NO_SHEET_HAS = "nonexistent"
A_VALUE_NO_DROPDOWN_OFFERS = "definitely"
ADDED_TITLE = "Coda"
ANCHORED_TITLE = "Interlude"
REVISED_NOTE = "rewatch first"
# Far enough down that reseating it past the second row is a visible move.
MOVED_ROW = 6


def _grid(path: Path) -> list[list[object]]:
    sheet = load_workbook(path).worksheets[0]
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def _titles(facts: WorkbookFacts, path: Path) -> list[str]:
    return [row.cells[facts.title_key] for row in read_category(path).rows]


def _range_ends(path: Path) -> set[int]:
    return set(range_ends(load_workbook(path).worksheets[0]))


def test_an_add_lands_at_the_end_of_the_sheet(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)

    apply_changes(subject_path, [RowChange(kind="add", cells={subject.title_key: ADDED_TITLE})])

    after = read_category(subject_path)
    assert after.counts.total == before.counts.total + 1
    assert after.rows[-1].cells[subject.title_key] == ADDED_TITLE


def test_a_remove_closes_the_gap_it_leaves(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)

    apply_changes(subject_path, [RowChange(kind="remove", row=FIRST_DATA_ROW)])

    after = read_category(subject_path)
    assert after.counts.total == before.counts.total - 1
    assert [row.cells[subject.title_key] for row in after.rows] == [
        row.cells[subject.title_key] for row in before.rows[1:]
    ]


def test_a_revise_rewrites_only_the_cells_it_names(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)

    apply_changes(
        subject_path, [RowChange(kind="revise", row=FIRST_DATA_ROW + 1, cells={subject.free_text_key: REVISED_NOTE})]
    )

    after = read_category(subject_path)
    assert after.rows[1].cells[subject.free_text_key] == REVISED_NOTE
    assert after.rows[1].cells[subject.title_key] == before.rows[1].cells[subject.title_key]
    assert [row.cells[subject.title_key] for row in after.rows] == [row.cells[subject.title_key] for row in before.rows]


def test_a_move_reseats_a_row_directly_after_its_anchor(subject: WorkbookFacts, subject_path: Path) -> None:
    before = _titles(subject, subject_path)

    apply_changes(subject_path, [RowChange(kind="move", row=MOVED_ROW, after_row=FIRST_DATA_ROW)])

    assert _titles(subject, subject_path) == [before[0], before[4], *before[1:4], *before[5:]]


def test_a_move_anchored_to_the_header_row_lands_first(subject: WorkbookFacts, subject_path: Path) -> None:
    before = _titles(subject, subject_path)

    apply_changes(subject_path, [RowChange(kind="move", row=MOVED_ROW, after_row=1)])

    assert _titles(subject, subject_path) == [before[4], *before[:4], *before[5:]]


def test_an_anchored_add_lands_directly_after_its_anchor(subject: WorkbookFacts, subject_path: Path) -> None:
    before = _titles(subject, subject_path)

    apply_changes(
        subject_path,
        [RowChange(kind="add", after_row=FIRST_DATA_ROW, cells={subject.title_key: ANCHORED_TITLE})],
    )

    assert _titles(subject, subject_path) == [before[0], ANCHORED_TITLE, *before[1:]]


def test_a_remove_and_a_move_in_one_batch_both_read_the_original_row_numbers(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    before = _titles(subject, subject_path)

    apply_changes(
        subject_path,
        [RowChange(kind="remove", row=FIRST_DATA_ROW), RowChange(kind="move", row=MOVED_ROW, after_row=3)],
    )

    assert _titles(subject, subject_path) == [before[1], before[4], before[2], before[3], *before[5:]]


@pytest.mark.parametrize(
    ("kind", "row", "after_row"),
    [
        ("revise", BEYOND_THE_LAST_ROW, None),
        ("remove", BEYOND_THE_LAST_ROW, None),
        ("move", BEYOND_THE_LAST_ROW, FIRST_DATA_ROW),
        ("move", FIRST_DATA_ROW, BEYOND_THE_LAST_ROW),
        ("add", None, BEYOND_THE_LAST_ROW),
    ],
    ids=["revise", "remove", "move-source", "move-anchor", "add-anchor"],
)
def test_a_change_pointing_past_the_last_row_is_refused(
    subject: WorkbookFacts, subject_path: Path, kind: str, row: int | None, after_row: int | None
) -> None:
    cells = {subject.title_key: ADDED_TITLE} if kind in {"revise", "add"} else {}
    change = RowChange(kind=kind, row=row, after_row=after_row, cells=cells)

    with pytest.raises(RowNotFoundError):
        apply_changes(subject_path, [change])


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_an_empty_batch_leaves_every_cell_exactly_where_it_was(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    before = _grid(path)

    apply_changes(path, [])

    assert _grid(path) == before


def test_table_dropdown_and_highlight_ranges_all_span_the_new_row_count(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    before = read_category(subject_path)

    apply_changes(
        subject_path,
        [
            RowChange(kind="remove", row=FIRST_DATA_ROW),
            RowChange(kind="add", after_row=4, cells={subject.title_key: ANCHORED_TITLE}),
            RowChange(kind="add", cells={subject.title_key: ADDED_TITLE}),
        ],
    )

    after = read_category(subject_path)
    assert after.counts.total == before.counts.total + 1
    assert _range_ends(subject_path) == {FIRST_DATA_ROW + after.counts.total - 1}


def test_a_batch_holding_an_unknown_column_writes_nothing(subject: WorkbookFacts, subject_path: Path) -> None:
    before = subject_path.read_bytes()
    batch = [
        RowChange(kind="revise", row=FIRST_DATA_ROW, cells={subject.free_text_key: "fine"}),
        RowChange(kind="revise", row=FIRST_DATA_ROW + 1, cells={A_COLUMN_NO_SHEET_HAS: "x"}),
    ]

    with pytest.raises(UnknownColumnError):
        apply_changes(subject_path, batch)

    assert subject_path.read_bytes() == before


def test_a_batch_holding_a_value_the_dropdown_rejects_writes_nothing(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    before = subject_path.read_bytes()
    batch = [
        RowChange(kind="revise", row=FIRST_DATA_ROW, cells={subject.free_text_key: "fine"}),
        RowChange(kind="revise", row=FIRST_DATA_ROW + 1, cells={subject.watch_key: A_VALUE_NO_DROPDOWN_OFFERS}),
    ]

    with pytest.raises(InvalidChoiceError):
        apply_changes(subject_path, batch)

    assert subject_path.read_bytes() == before


def test_removing_every_row_leaves_an_empty_but_intact_sheet(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)

    apply_changes(
        subject_path,
        [RowChange(kind="remove", row=row) for row in range(FIRST_DATA_ROW, len(before.rows) + FIRST_DATA_ROW)],
    )

    after = read_category(subject_path)
    assert after.rows == []
    assert [column.label for column in after.columns] == [column.label for column in before.columns]
    assert _range_ends(subject_path) == {FIRST_DATA_ROW}
