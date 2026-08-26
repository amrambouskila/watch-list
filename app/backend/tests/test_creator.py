from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from library_choice import a_name_the_library_does_not_hold, the_category_a_new_name_could_shadow
from tv_watchlist.models.category_create import CategoryCreate
from tv_watchlist.workbook.creator import create_category, table_name
from tv_watchlist.workbook.discovery import workbook_paths
from tv_watchlist.workbook.errors import DuplicateCategoryError
from tv_watchlist.workbook.reader import read_category
from tv_watchlist.workbook.schema import watch_column
from tv_watchlist.workbook.writer import append_row, update_row

CREATED = a_name_the_library_does_not_hold("Cowboy Bebop")
CREATED_STEM = "Cowboy_Bebop"
SEEDED = a_name_the_library_does_not_hold("Trigun")
WRITABLE = a_name_the_library_does_not_hold("Monster")
BARE = a_name_the_library_does_not_hold("Lain")
REPEATED = a_name_the_library_does_not_hold("Berserk")


def test_new_category_matches_the_library_shape(library: Path) -> None:
    path = create_category(library, CategoryCreate(name=CREATED, accent="#2563EB"))

    assert path.name == f"{CREATED_STEM}.xlsx"
    detail = read_category(path)
    assert detail.name == CREATED_STEM
    assert detail.accent == "#2563EB"
    assert detail.has_watch_column

    sheet = load_workbook(path).worksheets[0]
    assert sheet.title == "Master Watch Order"
    table = sheet.tables[table_name(CREATED)]
    assert [column.name for column in table.tableColumns] == [
        str(sheet.cell(row=1, column=index).value) for index in range(1, sheet.max_column + 1)
    ]
    assert [str(v.sqref) for v in sheet.data_validations.dataValidation] == ["H2"]
    assert [str(entry.sqref) for entry in sheet.conditional_formatting] == ["H2"]
    assert sheet.cell(row=1, column=1).fill.fgColor.rgb == "FF2563EB"


def test_new_category_seeds_the_titles_it_was_given(library: Path) -> None:
    path = create_category(
        library, CategoryCreate(name=SEEDED, titles=["Trigun", "Trigun Stampede", "Badlands Rumble"])
    )
    detail = read_category(path)
    assert [row.cells["title"] for row in detail.rows] == ["Trigun", "Trigun Stampede", "Badlands Rumble"]
    assert [row.cells["order"] for row in detail.rows] == ["1", "2", "3"]
    assert detail.counts.total == 3


def test_new_category_is_immediately_writable(library: Path) -> None:
    path = create_category(library, CategoryCreate(name=WRITABLE, titles=[WRITABLE]))
    watch = watch_column(read_category(path).columns)
    assert watch is not None

    update_row(path, 2, {watch.key: "Watched"})
    row = append_row(path, {"title": "Another Monster"})

    detail = read_category(path)
    assert detail.counts.watched == 1
    assert detail.rows[-1].cells["title"] == "Another Monster"
    assert load_workbook(path).worksheets[0].tables[table_name(WRITABLE)].ref == f"A1:H{row}"


def test_new_category_gains_a_watch_column_even_if_not_requested(library: Path) -> None:
    path = create_category(library, CategoryCreate(name=BARE, columns=["Order", "Title"]))
    detail = read_category(path)
    assert [column.label for column in detail.columns] == ["Order", "Title", "Watched?"]


def test_duplicate_category_is_refused(library: Path) -> None:
    create_category(library, CategoryCreate(name=REPEATED))
    with pytest.raises(DuplicateCategoryError):
        create_category(library, CategoryCreate(name=REPEATED))


def test_a_name_the_library_already_answers_to_is_refused_even_under_a_free_filename(library: Path) -> None:
    """The app addresses a category by the id its filename derives, so a free filename is not enough."""
    shadowed = the_category_a_new_name_could_shadow()
    before = workbook_paths(library)

    with pytest.raises(DuplicateCategoryError) as refusal:
        create_category(library, CategoryCreate(name=shadowed.shadowing_name))

    assert shadowed.file_name in str(refusal.value)
    assert workbook_paths(library) == before


@pytest.mark.parametrize(
    ("display", "expected"),
    [("Cowboy Bebop", "CowboyBebopMasterOrder"), ("86", "_86MasterOrder"), ("Alpha-Beta!", "AlphaBetaMasterOrder")],
)
def test_table_name_is_a_legal_excel_identifier(display: str, expected: str) -> None:
    assert table_name(display) == expected
