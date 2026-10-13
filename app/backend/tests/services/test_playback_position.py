from __future__ import annotations

import uuid
from collections.abc import Mapping
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from mirror_support import moved_on
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from playback_projection import PlaybackProjection
from tv_watchlist.data.enums import WatchEventKind, WatchEventSource, WatchSubjectKind
from tv_watchlist.data.errors import PlaybackStateError
from tv_watchlist.data.models import WatchEvent
from tv_watchlist.ingest.expansion import rebuild_items
from tv_watchlist.models.category_detail import CategoryDetail
from tv_watchlist.models.row_change import RowChange
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.services.playback_events import record_playback_event
from tv_watchlist.services.playback_position import project_playback_position
from tv_watchlist.workbook.errors import StaleWorkbookError, WorkbookLockedError
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.reader import read_projection
from tv_watchlist.workbook.schema import watch_column
from workbook_edits import narrow_the_status_dropdown


@pytest.mark.parametrize("completed", [False, True])
async def test_one_snapshot_writes_position_and_all_affected_watch_cells(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    completed: bool,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    body = projected.snapshot(
        seconds=300 if completed else 60, kind=WatchEventKind.COMPLETED if completed else WatchEventKind.PROGRESSED
    )
    entry = projected.entries[1]
    result = await record_playback_event(db_session, body, entry.id, projected.file.id)
    await db_session.commit()
    detail = await project_playback_position(
        db_session,
        catalog,
        entry.id,
        result.id,
        expected_mtime=projected.category.workbook_mtime,
    )
    projection = read_projection(subject_path)
    assert projection.position_values[entry.sheet_row] == str(body.position_seconds)
    expected = ["Watched"] * 3 if completed else ["", "In progress", ""]
    assert [projected.cell(row) for row in projected.entries] == expected
    assert [row.mirrored_watch_cell for row in projected.entries] == expected
    assert [row.observed_watch_cell for row in projected.entries] == expected
    assert detail.mtime == pytest.approx(subject_path.stat().st_mtime)


@pytest.mark.parametrize("fault", ["unknown", "ui", "category", "subject", "kind", "position"])
async def test_projection_refuses_unrelated_or_nonplayer_events_without_a_write(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    fault: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    selected = projected.entries[1]
    event = WatchEvent(
        subject_kind=WatchSubjectKind.EPISODE,
        episode_id=projected.episode_ids[0 if fault == "subject" else 1],
        kind=WatchEventKind.STARTED if fault in {"kind", "position"} else WatchEventKind.PROGRESSED,
        source=WatchEventSource.UI if fault == "ui" else WatchEventSource.PLAYER,
        category_id=None if fault == "category" else projected.category.id,
        occurred_at=datetime.now(UTC),
        position_seconds=None if fault == "position" else 60,
    )
    if fault == "position":
        event.kind = WatchEventKind.COMPLETED
    db_session.add(event)
    await db_session.flush()
    untouched = subject_path.read_bytes()
    with pytest.raises(PlaybackStateError):
        await project_playback_position(
            db_session,
            catalog,
            selected.id,
            uuid.uuid4() if fault == "unknown" else event.id,
            expected_mtime=projected.category.workbook_mtime,
        )
    assert subject_path.read_bytes() == untouched


@pytest.mark.parametrize("newer", ["occurred", "recorded", "id"])
async def test_an_older_player_event_projects_the_newest_snapshot_without_regression(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    newer: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    stamp = datetime(2026, 9, 24, tzinfo=UTC)
    first = WatchEvent(
        id=uuid.UUID(int=1),
        subject_kind=WatchSubjectKind.EPISODE,
        episode_id=projected.episode_ids[1],
        category_id=projected.category.id,
        source=WatchEventSource.PLAYER,
        kind=WatchEventKind.PROGRESSED,
        position_seconds=60,
        runtime_seconds=300,
        occurred_at=stamp,
        recorded_at=stamp,
    )
    second = WatchEvent(
        id=uuid.UUID(int=2),
        subject_kind=WatchSubjectKind.EPISODE,
        episode_id=projected.episode_ids[1],
        category_id=projected.category.id,
        source=WatchEventSource.PLAYER,
        kind=WatchEventKind.PROGRESSED,
        position_seconds=120,
        runtime_seconds=300,
        occurred_at=stamp + (timedelta(seconds=1) if newer == "occurred" else timedelta()),
        recorded_at=stamp + (timedelta(seconds=1) if newer == "recorded" else timedelta()),
    )
    db_session.add_all([first, second])
    await db_session.flush()
    await project_playback_position(
        db_session, catalog, entry.id, first.id, expected_mtime=projected.category.workbook_mtime
    )
    assert read_projection(subject_path).position_values[entry.sheet_row] == "120"


async def test_late_older_range_completion_repairs_siblings_without_regressing_newer_episode_position(
    db_session: AsyncSession, catalog: Catalog, subject_path: Path
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    entry.span_end = 3
    await rebuild_items(db_session, projected.category.id)
    stamp = datetime(2026, 9, 24, tzinfo=UTC)
    newer = await record_playback_event(
        db_session,
        projected.snapshot(selected=2, seconds=120, occurred_at=stamp + timedelta(minutes=1)),
        entry.id,
        projected.file.id,
    )
    await db_session.commit()
    first = await project_playback_position(
        db_session, catalog, entry.id, newer.id, expected_mtime=projected.category.workbook_mtime
    )
    await db_session.commit()
    assert projected.cell(projected.entries[0]) == ""
    older = await record_playback_event(
        db_session,
        projected.snapshot(seconds=300, kind=WatchEventKind.COMPLETED, occurred_at=stamp),
        entry.id,
        projected.file.id,
    )
    await db_session.commit()
    await project_playback_position(db_session, catalog, entry.id, older.id, expected_mtime=first.mtime)
    assert read_projection(subject_path).position_values[entry.sheet_row] == "120"
    assert [projected.cell(row) for row in projected.entries] == ["Watched", "In progress", "In progress"]
    assert projected.entries[0].observed_watch_cell == projected.entries[0].mirrored_watch_cell == "Watched"
    assert await db_session.scalar(select(func.count()).select_from(WatchEvent)) == 4


async def test_a_new_snapshot_repairs_multiple_committed_projections_that_failed_in_other_rows(
    db_session: AsyncSession, catalog: Catalog, subject_path: Path
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    stamp = projected.category.workbook_mtime
    excel_lock_path(subject_path).write_bytes(b"owner")
    for selected in (0, 1):
        entry = projected.entries[selected]
        event = await record_playback_event(
            db_session, projected.snapshot(selected=selected), entry.id, projected.file.id
        )
        await db_session.commit()
        with pytest.raises(WorkbookLockedError):
            await project_playback_position(db_session, catalog, entry.id, event.id, expected_mtime=stamp)
        await db_session.rollback()
        await db_session.refresh(projected.file)
        for row in projected.entries:
            await db_session.refresh(row)
    excel_lock_path(subject_path).unlink()
    newest_entry = projected.entries[2]
    newest = await record_playback_event(
        db_session, projected.snapshot(selected=2, seconds=120), newest_entry.id, projected.file.id
    )
    await db_session.commit()
    await project_playback_position(db_session, catalog, newest_entry.id, newest.id, expected_mtime=stamp)
    assert [projected.cell(row) for row in projected.entries] == ["In progress"] * 3
    assert [row.mirrored_watch_cell for row in projected.entries] == ["In progress"] * 3
    assert read_projection(subject_path).position_values[newest_entry.sheet_row] == "120"
    assert await db_session.scalar(select(func.count()).select_from(WatchEvent)) == 3


async def test_unknown_runtime_projects_without_completing_and_calls_one_write(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    body = projected.snapshot().model_copy(update={"runtime_seconds": None})
    event = await record_playback_event(db_session, body, entry.id, projected.file.id)
    writes: list[float] = []
    original = catalog.apply_position_writeback

    async def counted(
        slug: str, seconds: Mapping[int, int], cells: Mapping[int, Mapping[str, str]], expected_mtime: float
    ) -> CategoryDetail:
        writes.append(expected_mtime)
        return await original(slug, seconds, cells, expected_mtime)

    monkeypatch.setattr(catalog, "apply_position_writeback", counted)
    await project_playback_position(
        db_session, catalog, entry.id, event.id, expected_mtime=projected.category.workbook_mtime
    )
    assert len(writes) == 1
    assert projected.cell(entry) == "In progress"
    assert read_projection(subject_path).position_values[entry.sheet_row] == "60"


@pytest.mark.parametrize("fault", ["missing", "duplicate", "moved"])
async def test_first_write_uses_current_row_identity_even_with_a_fresh_queue_stamp(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    fault: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    event = await record_playback_event(db_session, projected.snapshot(), entry.id, projected.file.id)
    first, selected = projected.entries[0].sheet_row, entry.sheet_row
    assert first is not None and selected is not None
    ids = {selected: ""} if fault == "missing" else {first: str(entry.id)}
    if fault == "moved":
        ids[selected] = str(projected.entries[0].id)
    detail = await catalog.write_row_ids(projected.category.slug, ids, projected.category.workbook_mtime)
    untouched = subject_path.read_bytes()
    if fault != "moved":
        with pytest.raises(StaleWorkbookError):
            await project_playback_position(db_session, catalog, entry.id, event.id, expected_mtime=detail.mtime)
        assert subject_path.read_bytes() == untouched
    else:
        await project_playback_position(db_session, catalog, entry.id, event.id, expected_mtime=detail.mtime)
        assert read_projection(subject_path).position_values[first] == "60"
        assert read_projection(subject_path).position_values[selected] == ""


@pytest.mark.parametrize("race", ["stamp", "twice", "position", "identity", "duplicate", "watch"])
async def test_only_one_retry_and_only_when_every_target_is_unchanged(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    race: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    event = await record_playback_event(db_session, projected.snapshot(), entry.id, projected.file.id)
    original = catalog.apply_position_writeback
    attempts: list[float] = []

    async def raced(
        slug: str, seconds: Mapping[int, int], cells: Mapping[int, Mapping[str, str]], expected_mtime: float
    ) -> CategoryDetail:
        attempts.append(expected_mtime)
        if len(attempts) == 1 or race == "twice":
            row = entry.sheet_row
            assert row is not None
            if race == "position":
                await original(slug, {row: 99}, {}, expected_mtime)
            elif race == "watch":
                await catalog.apply_changes(
                    slug, [RowChange(kind="revise", row=row, cells={projected.watch_key: "Skip"})], expected_mtime
                )
            elif race in {"identity", "duplicate"}:
                target = row if race == "identity" else projected.entries[0].sheet_row
                assert target is not None
                await catalog.write_row_ids(
                    slug, {target: str(uuid.uuid4()) if race == "identity" else str(entry.id)}, expected_mtime
                )
            else:
                moved_on(subject_path)
        return await original(slug, seconds, cells, expected_mtime)

    monkeypatch.setattr(catalog, "apply_position_writeback", raced)
    if race == "stamp":
        await project_playback_position(
            db_session, catalog, entry.id, event.id, expected_mtime=projected.category.workbook_mtime
        )
        assert read_projection(subject_path).position_values[entry.sheet_row] == "60"
        assert len(attempts) == 2
    else:
        with pytest.raises(StaleWorkbookError):
            await project_playback_position(
                db_session, catalog, entry.id, event.id, expected_mtime=projected.category.workbook_mtime
            )
        assert len(attempts) == (2 if race == "twice" else 1)
        assert entry.mirrored_watch_cell == ""


@pytest.mark.parametrize("refusal", ["locked", "stale"])
async def test_a_committed_event_survives_workbook_refusal(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    refusal: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    event = await record_playback_event(db_session, projected.snapshot(), entry.id, projected.file.id)
    await db_session.commit()
    if refusal == "locked":
        excel_lock_path(subject_path).write_bytes(b"owner")
    else:
        moved_on(subject_path)
    untouched = subject_path.read_bytes()
    with pytest.raises(WorkbookLockedError if refusal == "locked" else StaleWorkbookError):
        await project_playback_position(
            db_session, catalog, entry.id, event.id, expected_mtime=projected.category.workbook_mtime
        )
    await db_session.rollback()
    assert await db_session.scalar(select(func.count()).select_from(WatchEvent)) == 1
    assert subject_path.read_bytes() == untouched


@pytest.mark.parametrize("condition", ["human", "dropdown"])
async def test_position_updates_without_overwriting_a_human_status_or_invalid_choice(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    condition: str,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    event = await record_playback_event(db_session, projected.snapshot(), entry.id, projected.file.id)
    detail = await catalog.detail(projected.category.slug)
    if condition == "human":
        await catalog.apply_changes(
            detail.id,
            [RowChange(kind="revise", row=entry.sheet_row, cells={projected.watch_key: "Skip"})],
            detail.mtime,
        )
    else:
        column = watch_column(detail.columns)
        assert column is not None
        narrow_the_status_dropdown(subject_path, column.index, ["", "Watched"])
    stamp = subject_path.stat().st_mtime
    await project_playback_position(db_session, catalog, entry.id, event.id, expected_mtime=stamp)
    assert read_projection(subject_path).position_values[entry.sheet_row] == "60"
    assert projected.cell(entry) == ("Skip" if condition == "human" else "")
    assert entry.mirrored_watch_cell == entry.observed_watch_cell == ""


@pytest.mark.parametrize("watch_present", [True, False])
async def test_each_minute_updates_position_even_without_a_pending_status_change(
    db_session: AsyncSession,
    catalog: Catalog,
    subject_path: Path,
    watch_present: bool,
) -> None:
    projected = await PlaybackProjection.create(db_session, catalog, subject_path)
    entry = projected.entries[1]
    if not watch_present:
        book = load_workbook(subject_path)
        sheet = book.worksheets[0]
        column = watch_column(read_projection(subject_path).detail.columns)
        assert column is not None
        sheet.cell(row=1, column=column.index).value = "Personal marker"
        book.save(subject_path)
    for seconds in (60, 120):
        event = await record_playback_event(
            db_session, projected.snapshot(seconds=seconds), entry.id, projected.file.id
        )
        await project_playback_position(
            db_session, catalog, entry.id, event.id, expected_mtime=subject_path.stat().st_mtime
        )
        assert read_projection(subject_path).position_values[entry.sheet_row] == str(seconds)
