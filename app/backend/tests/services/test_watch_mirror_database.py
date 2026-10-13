"""The mirror's session half: the derived state read off the view, and a whole category mirrored.

The database is real and skips loudly without one - the premise under test is a dual store, and a mocked
one would prove only that the mock agrees with itself (CLAUDE.md section 16.4). The workbooks are real
copies of the owner's own files, so a pass here writes an actual `.xlsx` through `Catalog._write`.
"""

from __future__ import annotations

import uuid
from collections.abc import Iterator, Sequence
from datetime import UTC, datetime
from itertools import count
from pathlib import Path
from typing import Final

import pytest
from mirror_support import Projected, row_of
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from tv_watchlist.constants import FIRST_DATA_ROW, IN_PROGRESS, SKIPPED, UNWATCHED, WATCHED
from tv_watchlist.data.enums import (
    MediaType,
    WatchEventKind,
    WatchEventSource,
    WatchSubjectKind,
)
from tv_watchlist.data.models import (
    Category,
    CuratedEntry,
    CuratedEntryItem,
    Episode,
    Season,
    Title,
    WatchEvent,
)
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.services.mirror_outcome import UNCHANGED, WRITTEN
from tv_watchlist.services.watch_mirror import (
    entry_watch_states,
    human_edit_events,
    mirror_category,
    mirror_entry,
)
from tv_watchlist.workbook.identity import ensure_identity_column, write_row_ids
from tv_watchlist.workbook.reader import read_category
from workbook_facts import WorkbookFacts

NOW: Final[datetime] = datetime(2026, 9, 20, 15, 0, tzinfo=UTC)
EARLIER: Final[datetime] = datetime(2026, 9, 20, 12, 0, tzinfo=UTC)
LATER: Final[datetime] = datetime(2026, 9, 20, 13, 0, tzinfo=UTC)
TWO_ROWS: Final[int] = 2
# Every test rolls back, so these only have to be unique inside one transaction.
PROVIDER_IDS: Final[Iterator[int]] = count(1)


async def _projected(session: AsyncSession, path: Path, subject: WorkbookFacts, rows: int = TWO_ROWS) -> Projected:
    """Mint identities into the sheet and the matching entries into Postgres, as a first ingest would.

    `raw_cells` is a synthetic `"row {row}"` marker rather than the sheet's real cells: unlike
    `mirror_support.entries_for`'s `path=` form, nothing here is compared against what those cells
    actually say, only against which row an entry sits on - so a fixed, readable marker is enough, and
    reading the real cells off `path` a second time (the reader has already been called once for
    `detail_before`) would buy nothing.
    """
    detail_before = read_category(path)
    category = Category(
        id=uuid.uuid4(),
        slug=detail_before.id,
        file_name=detail_before.file_name,
        sheet_title=detail_before.sheet_title,
        accent=detail_before.accent,
    )
    session.add(category)
    await session.flush()
    entries = [
        CuratedEntry(
            id=uuid.uuid4(),
            category_id=category.id,
            position=position,
            sheet_row=row,
            raw_cells={subject.title_key: f"row {row}"},
            mirrored_watch_cell=UNWATCHED,
            observed_watch_cell=UNWATCHED,
        )
        for position, row in enumerate(range(FIRST_DATA_ROW, FIRST_DATA_ROW + rows), start=1)
    ]
    session.add_all(entries)
    await session.flush()
    ensure_identity_column(path)
    write_row_ids(path, {row_of(entry): str(entry.id) for entry in entries})
    category.workbook_mtime = read_category(path).mtime
    return Projected(path, category, entries, subject.watch_key)


@pytest.fixture
async def projected(db_session: AsyncSession, subject_path: Path, subject: WorkbookFacts) -> Projected:
    return await _projected(db_session, subject_path, subject)


async def _record(
    session: AsyncSession,
    entry: CuratedEntry,
    kind: WatchEventKind,
    occurred_at: datetime,
    *,
    episode: Episode | None = None,
) -> WatchEvent:
    """One watch fact, against the entry itself or against one episode inside it."""
    event = WatchEvent(
        occurred_at=occurred_at,
        subject_kind=WatchSubjectKind.EPISODE if episode is not None else WatchSubjectKind.CURATED_ENTRY,
        kind=kind,
        source=WatchEventSource.UI,
        episode_id=episode.id if episode is not None else None,
        curated_entry_id=None if episode is not None else entry.id,
        category_id=entry.category_id,
    )
    session.add(event)
    await session.flush()
    return event


async def _episodes_under(session: AsyncSession, entry: CuratedEntry, how_many: int) -> Sequence[Episode]:
    """A title, a season and `how_many` episodes, flattened into this entry's queue slots."""
    title = Title(media_type=MediaType.TV, name="Probe", match_key="probe", tmdb_id=next(PROVIDER_IDS))
    session.add(title)
    await session.flush()
    season = Season(title_id=title.id, season_number=1)
    session.add(season)
    await session.flush()
    episodes = [
        Episode(
            title_id=title.id,
            season_id=season.id,
            season_number=1,
            episode_number=number,
            tmdb_episode_id=next(PROVIDER_IDS),
        )
        for number in range(1, how_many + 1)
    ]
    session.add_all(episodes)
    await session.flush()
    session.add_all(
        CuratedEntryItem(
            category_id=entry.category_id,
            curated_entry_id=entry.id,
            episode_id=episode.id,
            position_in_entry=slot,
            absolute_position=slot,
        )
        for slot, episode in enumerate(episodes, start=1)
    )
    await session.flush()
    return episodes


def _typed_into_excel(projected: Projected, entry: CuratedEntry, value: str) -> None:
    """A status the owner typed in Excel by hand, and the stamp the ingest that read it would store.

    Written straight through openpyxl on purpose: this is the other writer, not the app, and the point
    of the test is what the app does about a file it did not write.
    """
    detail = read_category(projected.path)
    index = next(column.index for column in detail.columns if column.key == projected.watch_key)
    workbook = load_workbook(projected.path)
    workbook.worksheets[0].cell(row=row_of(entry), column=index).value = value
    workbook.save(projected.path)
    projected.category.workbook_mtime = read_category(projected.path).mtime


async def _excel_mirror_events(session: AsyncSession, entry: CuratedEntry) -> int:
    return (
        await session.scalar(
            select(func.count())
            .select_from(WatchEvent)
            .where(
                WatchEvent.curated_entry_id == entry.id,
                WatchEvent.source == WatchEventSource.EXCEL_MIRROR,
            )
        )
    ) or 0


# -- the view, read as the mirror reads it -----------------------------------------------------------


async def test_an_entry_with_no_events_at_all_derives_as_unset(db_session: AsyncSession, projected: Projected) -> None:
    states = await entry_watch_states(db_session, projected.category.id)

    assert states[projected.entries[0].id] is WatchEventKind.UNSET


async def test_the_state_reflects_events_this_transaction_has_only_staged(
    db_session: AsyncSession, projected: Projected
) -> None:
    """A read that ignored the events just appended would mirror the cell the app is about to contradict."""
    entry = projected.entries[0]
    db_session.add_all(human_edit_events([entry], {entry.id: WATCHED}, now=NOW))

    states = await entry_watch_states(db_session, projected.category.id)

    assert states[entry.id] is WatchEventKind.COMPLETED


async def test_a_later_entry_level_event_wins_over_the_episode_roll_up(
    db_session: AsyncSession, projected: Projected
) -> None:
    """A coarse status typed into Excel sits on top of fine-grained history without erasing it."""
    entry = projected.entries[0]
    episodes = await _episodes_under(db_session, entry, 2)
    await _record(db_session, entry, WatchEventKind.COMPLETED, EARLIER, episode=episodes[0])
    await _record(db_session, entry, WatchEventKind.SKIPPED, LATER)

    states = await entry_watch_states(db_session, projected.category.id)

    assert states[entry.id] is WatchEventKind.SKIPPED
    # The episode history is still there underneath it; nothing was deleted to say "Skip".
    assert await db_session.scalar(select(func.count()).select_from(WatchEvent)) == 2


async def test_completed_plus_skipped_equals_total_derives_as_completed(
    db_session: AsyncSession, projected: Projected
) -> None:
    """The same semantics as `WatchCounts.trackable`, so the two progress readouts cannot disagree."""
    entry = projected.entries[0]
    episodes = await _episodes_under(db_session, entry, 2)
    await _record(db_session, entry, WatchEventKind.COMPLETED, EARLIER, episode=episodes[0])
    await _record(db_session, entry, WatchEventKind.SKIPPED, EARLIER, episode=episodes[1])

    states = await entry_watch_states(db_session, projected.category.id)

    assert states[entry.id] is WatchEventKind.COMPLETED


# -- a whole category mirrored -----------------------------------------------------------------------


async def test_mirror_category_writes_the_cell_the_view_derives(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    entry = projected.entries[0]
    await _record(db_session, entry, WatchEventKind.COMPLETED, NOW)

    outcome = await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)

    assert outcome == WRITTEN
    assert projected.cell(entry) == WATCHED
    assert entry.mirrored_watch_cell == WATCHED
    assert projected.category.workbook_mtime == projected.path.stat().st_mtime
    assert projected.category.last_mirrored_at == NOW


async def test_completed_plus_skipped_equals_total_mirrors_as_watched(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    entry = projected.entries[0]
    episodes = await _episodes_under(db_session, entry, 2)
    await _record(db_session, entry, WatchEventKind.COMPLETED, EARLIER, episode=episodes[0])
    await _record(db_session, entry, WatchEventKind.SKIPPED, EARLIER, episode=episodes[1])

    await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)

    assert projected.cell(entry) == WATCHED


async def test_a_later_entry_level_event_wins_over_the_episode_roll_up_in_the_cell(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    entry = projected.entries[0]
    episodes = await _episodes_under(db_session, entry, 2)
    await _record(db_session, entry, WatchEventKind.COMPLETED, EARLIER, episode=episodes[0])
    await _record(db_session, entry, WatchEventKind.STARTED, LATER)

    await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)

    assert projected.cell(entry) == IN_PROGRESS


async def test_mirror_category_with_nothing_derived_writes_nothing(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    stamped = projected.path.stat().st_mtime_ns

    outcome = await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)

    assert outcome == UNCHANGED
    assert projected.path.stat().st_mtime_ns == stamped


async def test_an_unknown_category_is_a_no_op(db_session: AsyncSession, catalog: Catalog) -> None:
    assert await mirror_category(db_session, catalog, uuid.uuid4(), clock=lambda: NOW) == UNCHANGED


# -- one entry mirrored ------------------------------------------------------------------------------


async def test_mirror_entry_writes_that_entry_and_leaves_the_rest_alone(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    wanted, other = projected.entries
    await _record(db_session, wanted, WatchEventKind.SKIPPED, NOW)
    await _record(db_session, other, WatchEventKind.COMPLETED, NOW)

    outcome = await mirror_entry(db_session, catalog, wanted.id, clock=lambda: NOW)

    assert outcome == WRITTEN
    assert projected.cell(wanted) == SKIPPED
    assert projected.cell(other) == UNWATCHED
    assert other.mirrored_watch_cell == UNWATCHED


async def test_an_unknown_entry_is_a_no_op(db_session: AsyncSession, catalog: Catalog) -> None:
    assert await mirror_entry(db_session, catalog, uuid.uuid4(), clock=lambda: NOW) == UNCHANGED


# -- the echo loop that does not grow ----------------------------------------------------------------


async def test_a_human_edit_costs_exactly_one_confirming_write_and_then_settles(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    """The loop terminates because the write is what closes it, not because the write is skipped.

    Staged the way an ingest leaves it: the owner typed a status into Excel, the pass read it, appended
    the event and re-stamped the category. The status is therefore already in the cell, so the mirror's
    write only confirms it - but `mirrored_watch_cell` means the value this app last WROTE, and it is
    the record of that write that stops the next ingest reading the same cell as a fresh human edit. One
    confirming rewrite per hand edit is the price of an echo defence that never lies (section 5.16).
    """
    entry = projected.entries[0]
    _typed_into_excel(projected, entry, WATCHED)
    db_session.add_all(human_edit_events([entry], {entry.id: WATCHED}, now=NOW))
    entry.observed_watch_cell = WATCHED

    first = await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)
    settled = projected.path.stat().st_mtime_ns
    second = await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)

    assert (first, second) == (WRITTEN, UNCHANGED)
    assert entry.mirrored_watch_cell == WATCHED
    assert projected.cell(entry) == WATCHED
    assert projected.path.stat().st_mtime_ns == settled


async def test_a_second_pass_appends_no_further_excel_mirror_event(
    db_session: AsyncSession, catalog: Catalog, projected: Projected
) -> None:
    """Mirror writes the cell; the next ingest reads it back as our own echo and logs nothing."""
    entry = projected.entries[0]
    await _record(db_session, entry, WatchEventKind.COMPLETED, NOW)
    await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW)
    before = await _excel_mirror_events(db_session, entry)

    db_session.add_all(human_edit_events(projected.entries, projected.observed(), now=LATER))
    await db_session.flush()

    assert before == 0
    assert await _excel_mirror_events(db_session, entry) == 0
    assert await mirror_category(db_session, catalog, projected.category.id, clock=lambda: NOW) == UNCHANGED
