"""The mirror's workbook half: one `Catalog._write`, the retry that loses to the owner, and nothing else.

The workbook layer is never stood in for. Every pass here opens a throwaway copy of one of the owner's
own files, chosen by shape rather than by name, and every refusal is provoked for real - the `~$` owner
file Excel drops, a directory `os.replace` cannot write into, a stamp that moved under the plan
(CLAUDE.md sections 16.3-16.4). No database is needed: the entries and the category row are detached ORM
objects, which is what the session-bound half in `test_watch_mirror_database.py` supplies in production.
"""

from __future__ import annotations

import os
import uuid
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Final

import pytest
from mirror_support import Projected, category_row, entries_for, moved_on, row_of
from openpyxl import Workbook

from backup_support import restore_points
from snapshot import fingerprint
from tv_watchlist.config import Settings
from tv_watchlist.constants import (
    IN_PROGRESS,
    MIRROR_RETRY_ATTEMPTS,
    SKIPPED,
    TEMP_WRITE_PREFIX,
    UNWATCHED,
    WATCHED,
)
from tv_watchlist.data.enums import WatchEventKind
from tv_watchlist.models.row_write import RowWrite
from tv_watchlist.services.catalog import Catalog
from tv_watchlist.services.mirror_outcome import DEFERRED, UNCHANGED, WRITTEN, MirrorOutcome
from tv_watchlist.services.watch_mirror import human_edit_events, pending_mirror_cells, write_mirror_cells
from tv_watchlist.workbook.identity import ensure_identity_column, write_row_ids
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.reader import WorkbookProjection, read_category
from workbook_edits import narrow_the_status_dropdown
from workbook_facts import WorkbookFacts

NOW: Final[datetime] = datetime(2026, 9, 20, 13, 0, tzinfo=UTC)
LATER: Final[datetime] = datetime(2026, 9, 20, 14, 0, tzinfo=UTC)
TWO_ROWS: Final[int] = 2
# A directory a non-root owner may read and traverse but not write into, so `os.replace` refuses.
UNWRITABLE_DIRECTORY: Final[int] = 0o500
WRITABLE_DIRECTORY: Final[int] = 0o700
NO_STATUS_COLUMN_HEADERS: Final[tuple[str, ...]] = ("Order", "Title")
# What a sheet dropped into the library may offer instead of the four values the nineteen carry.
A_FOREIGN_VOCABULARY: Final[tuple[str, ...]] = ("Yes", "No")


class _ARestlessLibrary(Catalog):
    """A real catalog over a real library that something else keeps saving over.

    Not a stand-in for the workbook layer: every read and every write below is the production one. The
    only addition is a stamp bumped after each fresh read, which is what a second writer winning the
    race looks like from in here - and the only way to reach the end of the retries deterministically.
    """

    def __init__(self, settings: Settings, path: Path) -> None:
        super().__init__(settings)
        self._path = path

    async def projection(self, wanted_id: str) -> WorkbookProjection:
        read = await super().projection(wanted_id)
        moved_on(self._path)
        return read


def _a_hand_made_sheet_without_a_status_column(library: Path) -> Path:
    """The shape only a dropped-in file can have: headers, rows, and no `Watched?` column at all."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(NO_STATUS_COLUMN_HEADERS))
    for number in range(1, TWO_ROWS + 1):
        sheet.append([str(number), f"Title {number}"])
    path = library / "Mirrorless.xlsx"
    workbook.save(path)
    return path


def _projected(path: Path, watch_key: str, rows: int = TWO_ROWS) -> Projected:
    """Mint identities into the sheet the way a first ingest does, and build the rows it would have."""
    category_id = uuid.uuid4()
    entries = entries_for(category_id, rows)
    ensure_identity_column(path)
    write_row_ids(path, {row_of(entry): str(entry.id) for entry in entries})
    return Projected(path, category_row(read_category(path), category_id), entries, watch_key)


@pytest.fixture
def projected(subject_path: Path, subject: WorkbookFacts) -> Projected:
    return _projected(subject_path, subject.watch_key)


@pytest.fixture
def unwritable(library: Path) -> Iterator[None]:
    """The library directory as a filesystem that will not accept the atomic swap."""
    library.chmod(UNWRITABLE_DIRECTORY)
    try:
        yield
    finally:
        library.chmod(WRITABLE_DIRECTORY)


async def _mirror(
    catalog: Catalog, projected: Projected, pending: dict[uuid.UUID, str], *, now: datetime = NOW
) -> MirrorOutcome:
    return await write_mirror_cells(catalog, projected.category, projected.entries, pending, now=now)


# -- nothing to say ----------------------------------------------------------------------------------


async def test_an_unchanged_projection_performs_no_workbook_write(
    catalog: Catalog, settings: Settings, projected: Projected, subject: WorkbookFacts
) -> None:
    """The throttle that makes a mirror pass free: no diff, no open, no snapshot, no new stamp."""
    stamped = projected.path.stat().st_mtime_ns
    taken = len(restore_points(settings, subject))

    outcome = await _mirror(catalog, projected, {})

    assert outcome == UNCHANGED
    assert projected.path.stat().st_mtime_ns == stamped
    assert len(restore_points(settings, subject)) == taken
    assert projected.category.last_mirrored_at is None


async def test_a_category_never_ingested_defers_rather_than_writing_unguarded(
    catalog: Catalog, projected: Projected
) -> None:
    """With no stored stamp there is nothing to guard the write against, so it is not made."""
    projected.category.workbook_mtime = None
    entry = projected.entries[0]

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == DEFERRED
    assert projected.cell(entry) == UNWATCHED
    assert entry.mirrored_watch_cell == UNWATCHED


async def test_a_workbook_without_a_watch_column_is_never_mirrored_and_never_errors(
    catalog: Catalog, library: Path
) -> None:
    """A sheet the owner never gave a status column is not a failure; there is nothing to write.

    Built by hand rather than through `create_category`, which appends the standard column when a
    request leaves it out - so the only way such a workbook exists is the way this one arrives: dropped
    into the library folder.
    """
    path = _a_hand_made_sheet_without_a_status_column(library)
    entries = entries_for(uuid.uuid4(), TWO_ROWS)
    category = category_row(read_category(path), entries[0].category_id)
    stamped = path.stat().st_mtime_ns

    outcome = await write_mirror_cells(catalog, category, entries, {entries[0].id: WATCHED}, now=NOW)

    assert outcome == UNCHANGED
    assert path.stat().st_mtime_ns == stamped
    assert entries[0].mirrored_watch_cell == UNWATCHED


# -- the write ---------------------------------------------------------------------------------------


async def test_a_changed_projection_writes_the_cell_and_the_stamp_and_the_mirrored_value_together(
    catalog: Catalog, projected: Projected
) -> None:
    entry = projected.entries[0]

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == WRITTEN
    assert projected.cell(entry) == WATCHED
    assert entry.mirrored_watch_cell == WATCHED
    # The post-write stamp is what gets stored, or the next pass reads our own write as a change.
    assert projected.category.workbook_mtime == projected.path.stat().st_mtime
    assert projected.category.last_mirrored_at == NOW


async def test_only_the_entries_whose_projection_moved_have_their_cells_written(
    catalog: Catalog, projected: Projected
) -> None:
    moved, untouched = projected.entries

    await _mirror(catalog, projected, {moved.id: SKIPPED})

    assert projected.cell(moved) == SKIPPED
    assert projected.cell(untouched) == UNWATCHED
    assert untouched.mirrored_watch_cell == UNWATCHED


async def test_the_write_is_one_snapshot_and_disturbs_no_excel_construct(
    catalog: Catalog, settings: Settings, projected: Projected, subject: WorkbookFacts
) -> None:
    before = fingerprint(projected.path)
    taken = len(restore_points(settings, subject))
    entry = projected.entries[0]

    await _mirror(catalog, projected, {entry.id: IN_PROGRESS})

    assert fingerprint(projected.path) == before
    assert len(restore_points(settings, subject)) == taken + 1
    assert not list(projected.path.parent.glob(f"{TEMP_WRITE_PREFIX}*"))


async def test_the_echo_defence_survives_a_round_trip_through_the_sheet(catalog: Catalog, projected: Projected) -> None:
    """The section 10.1 acceptance line: what the mirror wrote is not read back as a human edit."""
    entry = projected.entries[0]

    await _mirror(catalog, projected, {entry.id: WATCHED})

    assert projected.observed()[entry.id] == WATCHED
    assert human_edit_events(projected.entries, projected.observed(), now=LATER) == []


async def test_a_second_pass_over_the_same_state_writes_nothing_at_all(
    catalog: Catalog, settings: Settings, projected: Projected, subject: WorkbookFacts
) -> None:
    """The echo loop that does not grow: the throttle sees the projection and the cell agree."""
    entry = projected.entries[0]
    states = {entry.id: WatchEventKind.COMPLETED}
    await _mirror(catalog, projected, pending_mirror_cells(projected.entries, states))
    stamped = projected.path.stat().st_mtime_ns
    taken = len(restore_points(settings, subject))

    outcome = await _mirror(catalog, projected, pending_mirror_cells(projected.entries, states))

    assert outcome == UNCHANGED
    assert projected.path.stat().st_mtime_ns == stamped
    assert len(restore_points(settings, subject)) == taken


# -- the refusals ------------------------------------------------------------------------------------


async def test_a_locked_workbook_defers_and_writes_nothing_to_either_store(
    catalog: Catalog, projected: Projected
) -> None:
    excel_lock_path(projected.path).write_bytes(b"owner")
    entry = projected.entries[0]
    stamped = projected.category.workbook_mtime

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    excel_lock_path(projected.path).unlink()
    assert outcome == DEFERRED
    assert projected.cell(entry) == UNWATCHED
    assert entry.mirrored_watch_cell == UNWATCHED
    assert projected.category.workbook_mtime == stamped
    assert projected.category.last_mirrored_at is None


async def test_the_identical_diff_replays_on_the_next_pass_once_the_lock_clears(
    catalog: Catalog, projected: Projected
) -> None:
    entry = projected.entries[0]
    lock = excel_lock_path(projected.path)
    lock.write_bytes(b"owner")
    assert await _mirror(catalog, projected, {entry.id: WATCHED}) == DEFERRED
    lock.unlink()

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == WRITTEN
    assert projected.cell(entry) == WATCHED


@pytest.mark.skipif(os.name == "nt", reason="directory permissions do not refuse a rename on Windows")
async def test_a_refused_swap_updates_neither_the_mirrored_cell_nor_the_stamp(
    catalog: Catalog, projected: Projected, unwritable: None
) -> None:
    """The harder shape than an owner file: the rewrite cannot land, and nothing in either store moves."""
    entry = projected.entries[0]
    stamped = projected.category.workbook_mtime

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == DEFERRED
    assert entry.mirrored_watch_cell == UNWATCHED
    assert projected.category.workbook_mtime == stamped
    assert projected.category.last_mirrored_at is None


# -- losing the race, deliberately -------------------------------------------------------------------


async def test_a_stamp_that_moved_for_another_reason_is_retried_against_the_sheet_as_it_stands(
    catalog: Catalog, projected: Projected
) -> None:
    """A write to some other column bumps the file; the mirror's own plan is still true, so it retries."""
    entry = projected.entries[0]
    moved_on(projected.path)

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == WRITTEN
    assert projected.cell(entry) == WATCHED
    assert projected.category.workbook_mtime == projected.path.stat().st_mtime


async def test_a_human_editing_the_very_cell_wins_and_the_mirror_defers(catalog: Catalog, projected: Projected) -> None:
    """Of the two writers, the mirror is the one that must lose (master plan section 6.3)."""
    entry = projected.entries[0]
    detail = read_category(projected.path)
    await catalog.update_row(
        detail.id, row_of(entry), RowWrite(cells={projected.watch_key: IN_PROGRESS}, expected_mtime=detail.mtime)
    )

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == DEFERRED
    assert projected.cell(entry) == IN_PROGRESS
    assert entry.mirrored_watch_cell == UNWATCHED


async def test_a_row_that_no_longer_carries_the_entry_stops_the_retry(catalog: Catalog, projected: Projected) -> None:
    """Row numbers mean nothing except against the grid they were resolved on."""
    entry = projected.entries[0]
    write_row_ids(projected.path, {row_of(entry): str(uuid.uuid4())})

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == DEFERRED
    assert projected.cell(entry) == UNWATCHED


async def test_a_library_that_keeps_moving_exhausts_the_retries_and_defers(
    settings: Settings, projected: Projected
) -> None:
    """Three attempts, three refusals, and the identical diff left to replay next pass."""
    restless = _ARestlessLibrary(settings, projected.path)
    entry = projected.entries[0]
    moved_on(projected.path)

    outcome = await write_mirror_cells(restless, projected.category, projected.entries, {entry.id: WATCHED}, now=NOW)

    assert outcome == DEFERRED
    assert MIRROR_RETRY_ATTEMPTS >= 1
    assert projected.cell(entry) == UNWATCHED
    assert entry.mirrored_watch_cell == UNWATCHED


# -- the sheet's own dropdown wins (Ruling 41) -------------------------------------------------------


async def test_a_value_the_sheet_s_own_dropdown_refuses_is_dropped_not_written(
    catalog: Catalog, settings: Settings, projected: Projected, subject: WorkbookFacts
) -> None:
    narrow_the_status_dropdown(projected.path, subject.index_of(subject.watch_key), A_FOREIGN_VOCABULARY)
    projected.category.workbook_mtime = read_category(projected.path).mtime
    entry = projected.entries[0]
    stamped = projected.path.stat().st_mtime_ns
    taken = len(restore_points(settings, subject))

    outcome = await _mirror(catalog, projected, {entry.id: WATCHED})

    assert outcome == UNCHANGED
    assert projected.path.stat().st_mtime_ns == stamped
    assert len(restore_points(settings, subject)) == taken
    assert entry.mirrored_watch_cell == UNWATCHED


async def test_the_values_that_dropdown_does_offer_are_still_written(
    catalog: Catalog, projected: Projected, subject: WorkbookFacts
) -> None:
    """One refused cell must not cost the rest of the diff, which is what a raised error would do."""
    index = subject.index_of(subject.watch_key)
    narrow_the_status_dropdown(projected.path, index, (UNWATCHED, WATCHED))
    projected.category.workbook_mtime = read_category(projected.path).mtime
    offered, refused = projected.entries

    outcome = await _mirror(catalog, projected, {offered.id: WATCHED, refused.id: SKIPPED})

    assert outcome == WRITTEN
    assert projected.cell(offered) == WATCHED
    assert projected.cell(refused) == UNWATCHED
    assert (offered.mirrored_watch_cell, refused.mirrored_watch_cell) == (WATCHED, UNWATCHED)
