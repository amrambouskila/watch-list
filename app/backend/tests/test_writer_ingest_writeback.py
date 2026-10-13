"""The ingest's one workbook write: identities and mirror cells in a single open, save and swap."""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from snapshot import fingerprint
from tv_watchlist.constants import FIRST_DATA_ROW, IDENTITY_HEADER_LABEL, TEMP_WRITE_PREFIX
from tv_watchlist.workbook.errors import (
    InvalidChoiceError,
    RowNotFoundError,
    UnknownColumnError,
    WorkbookLockedError,
)
from tv_watchlist.workbook.identity import ensure_identity_column
from tv_watchlist.workbook.locking import excel_lock_path
from tv_watchlist.workbook.naming import slugify
from tv_watchlist.workbook.reader import read_projection
from tv_watchlist.workbook.writer import apply_ingest_writeback
from workbook_facts import WorkbookFacts

IDENTITY_COLUMN_KEY = slugify(IDENTITY_HEADER_LABEL)
A_MADE_UP_COLUMN_KEY = "column-the-owner-never-authored"


def _ids(facts: WorkbookFacts) -> dict[int, str]:
    return {row: str(uuid4()) for row in range(FIRST_DATA_ROW, facts.last_data_row + 1)}


def test_it_mints_the_column_records_every_id_and_leaves_one_restore_shaped_file(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    minted = _ids(subject)

    apply_ingest_writeback(subject_path, minted, {})

    projection = read_projection(subject_path)
    assert projection.identity_key == IDENTITY_COLUMN_KEY
    assert projection.row_ids == minted
    assert not list(subject_path.parent.glob(f"{TEMP_WRITE_PREFIX}*"))


def test_it_writes_ids_and_cells_in_the_same_pass(subject: WorkbookFacts, subject_path: Path) -> None:
    minted = _ids(subject)
    written = "watched by the mirror"

    apply_ingest_writeback(subject_path, minted, {FIRST_DATA_ROW: {subject.free_text_key: written}})

    projection = read_projection(subject_path)
    first = projection.detail.rows[0]
    assert projection.row_ids == minted
    assert first.cells[subject.free_text_key] == written


def test_the_fingerprint_is_undisturbed_apart_from_the_identity_column(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    ensure_identity_column(subject_path)
    before = fingerprint(subject_path)

    apply_ingest_writeback(subject_path, _ids(subject), {})

    assert fingerprint(subject_path) == before


def test_a_second_pass_reuses_the_column_it_already_minted(subject: WorkbookFacts, subject_path: Path) -> None:
    apply_ingest_writeback(subject_path, _ids(subject), {})
    width = load_workbook(subject_path).worksheets[0].max_column
    second = _ids(subject)

    apply_ingest_writeback(subject_path, second, {})

    assert load_workbook(subject_path).worksheets[0].max_column == width
    assert read_projection(subject_path).row_ids == second


def test_an_unknown_column_key_is_refused_and_nothing_is_written(subject: WorkbookFacts, subject_path: Path) -> None:
    untouched = subject_path.read_bytes()

    with pytest.raises(UnknownColumnError):
        apply_ingest_writeback(subject_path, _ids(subject), {FIRST_DATA_ROW: {A_MADE_UP_COLUMN_KEY: "x"}})

    assert subject_path.read_bytes() == untouched


def test_a_value_the_sheets_own_dropdown_refuses_is_refused_here_too(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    untouched = subject_path.read_bytes()

    with pytest.raises(InvalidChoiceError):
        apply_ingest_writeback(subject_path, {}, {FIRST_DATA_ROW: {subject.watch_key: "definitely not an option"}})

    assert subject_path.read_bytes() == untouched


def test_a_cell_row_outside_the_grid_is_refused_and_nothing_is_written(
    subject: WorkbookFacts, subject_path: Path
) -> None:
    untouched = subject_path.read_bytes()
    beyond = subject.last_data_row + 1

    with pytest.raises(RowNotFoundError):
        apply_ingest_writeback(subject_path, {}, {beyond: {subject.free_text_key: "x"}})

    assert subject_path.read_bytes() == untouched


def test_an_open_workbook_refuses_the_write(subject: WorkbookFacts, subject_path: Path) -> None:
    excel_lock_path(subject_path).write_bytes(b"owner")
    untouched = subject_path.read_bytes()

    with pytest.raises(WorkbookLockedError):
        apply_ingest_writeback(subject_path, _ids(subject), {})

    assert subject_path.read_bytes() == untouched


def test_the_app_owned_column_cannot_be_addressed_as_a_cell(subject: WorkbookFacts, subject_path: Path) -> None:
    """A mirror write names the owner's columns; the Row ID column is reached only through `ids_by_row`."""
    ensure_identity_column(subject_path)
    untouched = subject_path.read_bytes()

    with pytest.raises(UnknownColumnError):
        apply_ingest_writeback(subject_path, {}, {FIRST_DATA_ROW: {IDENTITY_COLUMN_KEY: str(uuid4())}})

    assert subject_path.read_bytes() == untouched
