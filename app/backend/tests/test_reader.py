from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from library_choice import (
    a_category_carrying_reference_sheets,
    every_workbook,
    the_category_with_a_stale_formatting_range,
    the_category_with_dropdowns_beyond_the_watch_column,
)
from tv_watchlist.constants import DEFAULT_WATCH_CHOICES, FIRST_DATA_ROW, HEADER_ROW
from tv_watchlist.workbook.reader import read_category
from tv_watchlist.workbook.schema import watch_column
from workbook_facts import WorkbookFacts

ACCENT_LENGTH = len("#000000")
REFERENCE_SHEETS_WANTED = 2


def _last_written_row(sheet: Worksheet, column_index: int) -> int:
    """The deepest row the owner actually filled in, read without consulting any range."""
    return max(
        row for row in range(FIRST_DATA_ROW, sheet.max_row + 1) if sheet.cell(row=row, column=column_index).value
    )


def _inline_options(sheet: Worksheet, column_index: int) -> list[str]:
    """The dropdown options Excel holds against a column, read straight out of its formula."""
    formula = next(
        validation.formula1
        for validation in sheet.data_validations.dataValidation
        if any(area.min_col <= column_index <= area.max_col for area in validation.sqref.ranges)
    )
    return formula.strip().strip('"').split(",")


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_every_library_workbook_reads_with_a_watch_column(facts: WorkbookFacts) -> None:
    detail = read_category(facts.path)
    assert detail.rows, f"{facts.file_name} produced no rows"
    assert detail.has_watch_column
    watch = watch_column(detail.columns)
    assert watch is not None
    assert watch.kind == "choice"
    assert watch.choices == list(DEFAULT_WATCH_CHOICES)
    assert detail.accent.startswith("#")
    assert len(detail.accent) == ACCENT_LENGTH


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_counts_partition_every_row(facts: WorkbookFacts) -> None:
    counts = read_category(facts.path).counts
    assert counts.watched + counts.in_progress + counts.skipped + counts.unwatched == counts.total
    assert counts.trackable == counts.total - counts.skipped


def test_a_category_reads_the_rows_beyond_its_stale_formatting_range(library: Path) -> None:
    """The dropdown a hand-edit left short of the grid must not decide where the data ends."""
    facts = the_category_with_a_stale_formatting_range()
    detail = read_category(facts.copied_into(library))
    sheet = load_workbook(facts.copied_into(library)).worksheets[0]
    stale_end = min(facts.range_ends)

    last_written = _last_written_row(sheet, facts.index_of(facts.title_key))

    assert last_written > stale_end
    assert detail.rows[-1].row == last_written
    assert detail.counts.total == last_written - HEADER_ROW


def test_a_category_exposes_the_dropdowns_its_owner_built_beside_the_watch_column(library: Path) -> None:
    facts = the_category_with_dropdowns_beyond_the_watch_column()
    path = facts.copied_into(library)
    detail = read_category(path)
    sheet = load_workbook(path).worksheets[0]

    keys = facts.dropdown_keys_beyond_the_watch_column
    offered = {key: next(column.choices for column in detail.columns if column.key == key) for key in keys}

    assert offered == {key: _inline_options(sheet, facts.index_of(key)) for key in keys}
    assert all(len(choices) > 1 for choices in offered.values())


def test_reference_sheets_are_carried_through(library: Path) -> None:
    facts = a_category_carrying_reference_sheets(at_least=REFERENCE_SHEETS_WANTED)
    path = facts.copied_into(library)
    workbook = load_workbook(path)

    detail = read_category(path)

    assert [sheet.title for sheet in detail.reference_sheets] == [sheet.title for sheet in workbook.worksheets[1:]]
    first = workbook.worksheets[1]
    assert detail.reference_sheets[0].header == [
        str(first.cell(row=HEADER_ROW, column=column).value or "") for column in range(1, first.max_column + 1)
    ]
    assert all(detail.reference_sheets[0].header)
