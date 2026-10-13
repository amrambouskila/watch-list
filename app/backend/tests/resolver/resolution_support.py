"""Builders for the resolution tests: a request from literal cells, or from a real workbook row."""

from __future__ import annotations

from pathlib import Path
from typing import Final

from openpyxl import load_workbook

from tv_watchlist.models.tmdb.tmdb_search_hit import TmdbMediaType
from tv_watchlist.resolver.candidate import Candidate
from tv_watchlist.resolver.resolution_request import ResolutionRequest
from tv_watchlist.resolver.span import parse_unit

# How the test reads a workbook's own headers. This is a TEST convenience, never a resolver rule: which
# column feeds which `ResolutionRequest` field is T23's job, and the nineteen workbooks spell these
# headers nine different ways ("Unit", "Season / Unit", "Season / Episode Block", "Format").
TITLE_HEADER: Final[str] = "title"
UNIT_HINTS: Final[tuple[str, ...]] = ("unit", "season", "episode", "block")
RELEASE_HINTS: Final[tuple[str, ...]] = ("release",)
TYPE_HINTS: Final[tuple[str, ...]] = ("type", "format")

A_CATEGORY_SLUG: Final[str] = "misc-movies"
A_CATEGORY_NAME: Final[str] = "Misc_Movies"
A_TMDB_ID: Final[int] = 1


def a_request(
    raw_title: str,
    *,
    raw_unit: str = "Movie",
    raw_release: str = "",
    raw_type: str = "",
    category_slug: str = A_CATEGORY_SLUG,
    category_name: str = A_CATEGORY_NAME,
    expected_runtime_minutes: int | None = None,
) -> ResolutionRequest:
    """One row's cells as the resolver sees them, with the span parsed exactly as T17 parses it."""
    return ResolutionRequest(
        category_slug=category_slug,
        category_name=category_name,
        raw_title=raw_title,
        raw_unit=raw_unit,
        raw_release=raw_release,
        raw_type=raw_type,
        span=parse_unit(raw_unit, raw_type=raw_type),
        expected_runtime_minutes=expected_runtime_minutes,
    )


def a_candidate(
    name: str,
    *,
    media_type: TmdbMediaType = "movie",
    tmdb_id: int = A_TMDB_ID,
    year: int | None = None,
    runtime_minutes: int | None = None,
    episode_count: int | None = None,
    also_known_as: tuple[str, ...] = (),
) -> Candidate:
    """One TMDB hit as the resolver scores it."""
    return Candidate(
        media_type=media_type,
        tmdb_id=tmdb_id,
        name=name,
        year=year,
        runtime_minutes=runtime_minutes,
        episode_count=episode_count,
        also_known_as=also_known_as,
    )


def a_workbook_row(library: Path, file_name: str, title: str, *, unit: str | None = None) -> ResolutionRequest:
    """The request a real library row produces, read out of the throwaway copy of that workbook.

    `unit` picks between rows a workbook repeats under one title - One Piece has 28 of them.
    """
    workbook = load_workbook(library / file_name, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    headers = [header.casefold() for header in rows[0]]
    title_at = headers.index(TITLE_HEADER)
    unit_at = _column_with(headers, UNIT_HINTS)
    release_at = _column_with(headers, RELEASE_HINTS)
    type_at = _column_with(headers, TYPE_HINTS)
    matched = next(row for row in rows[1:] if row[title_at] == title and (unit is None or row[unit_at] == unit))
    return a_request(
        matched[title_at],
        raw_unit=matched[unit_at],
        raw_release=matched[release_at],
        raw_type=matched[type_at],
        category_slug=Path(file_name).stem.replace("_", "-").casefold(),
        category_name=Path(file_name).stem,
    )


def _column_with(headers: list[str], hints: tuple[str, ...]) -> int:
    """The first column whose header carries one of these words; the sheets agree on no single spelling."""
    return next(index for index, header in enumerate(headers) if any(hint in header for hint in hints))
