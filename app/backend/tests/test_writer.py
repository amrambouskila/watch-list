from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from library_choice import (
    a_category_whose_order_column_holds_numbers,
    a_name_the_library_does_not_hold,
    every_workbook,
    the_category_with_a_stale_formatting_range,
)
from snapshot import fingerprint
from tv_watchlist.constants import (
    DEFAULT_WATCH_CHOICES,
    FIRST_DATA_ROW,
    IN_PROGRESS,
    TEMP_WRITE_PREFIX,
    WATCH_HEADER_LABEL,
    WATCHED,
)
from tv_watchlist.workbook.errors import (
    InvalidChoiceError,
    RowNotFoundError,
    UnknownColumnError,
)
from tv_watchlist.workbook.reader import read_category
from tv_watchlist.workbook.schema import watch_column
from tv_watchlist.workbook.writer import add_watch_column, append_row, delete_row, update_row
from workbook_facts import WorkbookFacts, range_ends

A_TITLE_THE_LIBRARY_DOES_NOT_HOLD = "Bonus feature the owner never listed"
A_VALUE_NO_DROPDOWN_OFFERS = "definitely"
A_COLUMN_NO_SHEET_HAS = "nonexistent"
A_ROW_NO_SHEET_HAS = 9999
A_NOTE_THE_OWNER_MIGHT_TYPE = "left off mid-arc"
DROPPED_IN_TITLES = ["Serial Experiments Lain", "Texhnolyze"]
DROPPED_IN_STEM = a_name_the_library_does_not_hold("Dropped_In_Show")
TITLE_COLUMN_INDEX = 2


def _highlighted_columns(path: Path) -> list[str]:
    """Which columns carry a highlight rule, ignoring how far down the sheet each one reaches."""
    sheet = load_workbook(path).worksheets[0]
    return sorted(str(entry.sqref).partition(str(FIRST_DATA_ROW))[0] for entry in sheet.conditional_formatting)


def _body_style(path: Path, row: int) -> tuple[bool | None, str | None, str | None]:
    cell = load_workbook(path).worksheets[0].cell(row=row, column=TITLE_COLUMN_INDEX)
    return (cell.alignment.wrap_text, cell.alignment.vertical, cell.font.name)


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_update_preserves_every_excel_construct(library: Path, facts: WorkbookFacts) -> None:
    path = facts.copied_into(library)
    before = fingerprint(path)

    update_row(path, FIRST_DATA_ROW, {facts.watch_key: WATCHED})

    assert fingerprint(path) == before
    assert read_category(path).rows[0].cells[facts.watch_key] == WATCHED


@pytest.mark.parametrize("facts", every_workbook(), ids=lambda facts: facts.category_id)
def test_update_leaves_no_temp_file_behind(library: Path, facts: WorkbookFacts) -> None:
    update_row(facts.copied_into(library), FIRST_DATA_ROW, {facts.watch_key: IN_PROGRESS})
    assert not list(library.glob(f"{TEMP_WRITE_PREFIX}*"))


def test_update_rejects_a_value_the_sheet_dropdown_would_reject(subject: WorkbookFacts, subject_path: Path) -> None:
    with pytest.raises(InvalidChoiceError):
        update_row(subject_path, FIRST_DATA_ROW, {subject.watch_key: A_VALUE_NO_DROPDOWN_OFFERS})


def test_update_rejects_an_unknown_column(subject_path: Path) -> None:
    with pytest.raises(UnknownColumnError):
        update_row(subject_path, FIRST_DATA_ROW, {A_COLUMN_NO_SHEET_HAS: "x"})


@pytest.mark.parametrize("row", [1, A_ROW_NO_SHEET_HAS])
def test_update_rejects_rows_outside_the_data_range(subject: WorkbookFacts, subject_path: Path, row: int) -> None:
    with pytest.raises(RowNotFoundError):
        update_row(subject_path, row, {subject.watch_key: WATCHED})


def test_update_accepts_free_text_in_a_plain_column(subject: WorkbookFacts, subject_path: Path) -> None:
    update_row(subject_path, FIRST_DATA_ROW + 1, {subject.free_text_key: A_NOTE_THE_OWNER_MIGHT_TYPE})
    assert read_category(subject_path).rows[1].cells[subject.free_text_key] == A_NOTE_THE_OWNER_MIGHT_TYPE


def test_append_extends_table_validation_and_highlight_together(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)
    highlighted = _highlighted_columns(subject_path)

    row = append_row(subject_path, {subject.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    assert row == subject.last_data_row + 1
    sheet = load_workbook(subject_path).worksheets[0]
    assert range_ends(sheet) == (row,)
    assert sheet.tables[subject.grid_table_name].ref.endswith(str(row))
    assert _highlighted_columns(subject_path) == highlighted

    after = read_category(subject_path)
    assert after.counts.total == before.counts.total + 1
    assert after.rows[-1].cells[subject.title_key] == A_TITLE_THE_LIBRARY_DOES_NOT_HOLD
    assert after.rows[-1].cells[subject.order_key] == str(after.counts.total)


def test_append_repairs_a_stale_range_it_finds(library: Path) -> None:
    """A construct a hand-edit left short of the grid is stretched to it, not left behind."""
    facts = the_category_with_a_stale_formatting_range()
    path = facts.copied_into(library)
    assert not facts.ranges_all_end_together

    row = append_row(path, {facts.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    assert row == facts.last_data_row + 1
    assert range_ends(load_workbook(path).worksheets[0]) == (row,)


def test_appended_row_inherits_the_sheet_body_style(subject: WorkbookFacts, subject_path: Path) -> None:
    existing = _body_style(subject_path, FIRST_DATA_ROW)

    row = append_row(subject_path, {subject.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    assert _body_style(subject_path, row) == existing


def test_append_keeps_the_order_column_numeric(library: Path) -> None:
    facts = a_category_whose_order_column_holds_numbers()
    path = facts.copied_into(library)
    order_index = next(column.index for column in read_category(path).columns if column.key == facts.order_key)
    original_type = type(load_workbook(path).worksheets[0].cell(row=FIRST_DATA_ROW, column=order_index).value)

    row = append_row(path, {facts.title_key: A_TITLE_THE_LIBRARY_DOES_NOT_HOLD})

    assert type(load_workbook(path).worksheets[0].cell(row=row, column=order_index).value) is original_type


def test_delete_row_shrinks_ranges_and_resequences(subject: WorkbookFacts, subject_path: Path) -> None:
    before = read_category(subject_path)
    second_title = before.rows[1].cells[subject.title_key]

    delete_row(subject_path, FIRST_DATA_ROW)

    after = read_category(subject_path)
    assert after.counts.total == before.counts.total - 1
    assert after.rows[0].cells[subject.title_key] == second_title
    assert [row.cells[subject.order_key] for row in after.rows] == [
        str(number) for number in range(1, after.counts.total + 1)
    ]
    sheet = load_workbook(subject_path).worksheets[0]
    assert range_ends(sheet) == (subject.last_data_row - 1,)
    assert sheet.tables[subject.grid_table_name].ref.endswith(str(subject.last_data_row - 1))


def test_delete_rejects_the_header_row(subject_path: Path) -> None:
    with pytest.raises(RowNotFoundError):
        delete_row(subject_path, 1)


def _sheet_without_watch_column(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Master Watch Order"
    for index, header in enumerate(["Order", "Title", "Notes"], start=1):
        sheet.cell(row=1, column=index).value = header
    for offset, title in enumerate(DROPPED_IN_TITLES):
        sheet.cell(row=FIRST_DATA_ROW + offset, column=1).value = offset + 1
        sheet.cell(row=FIRST_DATA_ROW + offset, column=2).value = title
    workbook.save(path)


def test_add_watch_column_makes_a_plain_sheet_trackable(library: Path) -> None:
    path = library / f"{DROPPED_IN_STEM}.xlsx"
    _sheet_without_watch_column(path)
    assert read_category(path).has_watch_column is False

    add_watch_column(path)

    detail = read_category(path)
    watch = watch_column(detail.columns)
    assert watch is not None
    assert watch.label == WATCH_HEADER_LABEL
    assert watch.choices == list(DEFAULT_WATCH_CHOICES)
    sheet = load_workbook(path).worksheets[0]
    assert sorted(str(entry.sqref) for entry in sheet.conditional_formatting) == ["D2:D3"]

    update_row(path, FIRST_DATA_ROW, {watch.key: WATCHED})
    assert read_category(path).counts.watched == 1


def test_add_watch_column_is_idempotent(subject_path: Path) -> None:
    before = fingerprint(subject_path)
    add_watch_column(subject_path)
    assert fingerprint(subject_path) == before
