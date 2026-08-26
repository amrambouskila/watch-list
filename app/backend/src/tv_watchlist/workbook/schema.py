"""Derive the editable column schema from a sheet's header row and data validations."""

from __future__ import annotations

import re

from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import (
    HEADER_ROW,
    ORDER_HEADER_ALIASES,
    TITLE_HEADER_ALIASES,
    WATCH_HEADER_ALIASES,
)
from tv_watchlist.models.column_spec import ColumnRole, ColumnSpec
from tv_watchlist.workbook.naming import slugify

# "#" survives normalisation because the Miscellaneous sheet uses it as its order header.
_PUNCTUATION = re.compile(r"[^a-z0-9# ]+")
_INLINE_LIST = re.compile(r'^"(.*)"$', re.DOTALL)


def normalize_header(label: str) -> str:
    """Header text reduced to comparable words: lowercase, punctuation stripped."""
    return _PUNCTUATION.sub("", label.lower()).strip()


def _role_for(label: str) -> ColumnRole:
    normalized = normalize_header(label)
    if normalized in WATCH_HEADER_ALIASES:
        return "watch"
    if normalized in TITLE_HEADER_ALIASES:
        return "title"
    if normalized in ORDER_HEADER_ALIASES:
        return "order"
    return "other"


def parse_choices(formula: str | None) -> list[str]:
    """Options from an inline Excel list validation; empty for range-backed validations."""
    if not formula:
        return []
    match = _INLINE_LIST.match(formula.strip())
    if not match:
        return []
    return [option.strip() for option in match.group(1).split(",")]


def validation_choices(sheet: Worksheet) -> dict[int, list[str]]:
    """Map of column index to dropdown options declared on that column."""
    found: dict[int, list[str]] = {}
    for validation in sheet.data_validations.dataValidation:
        if validation.type != "list":
            continue
        options = parse_choices(validation.formula1)
        if not options:
            continue
        for cell_range in validation.sqref.ranges:
            for column in range(cell_range.min_col, cell_range.max_col + 1):
                found.setdefault(column, options)
    return found


def unique_key(label: str, index: int, taken: set[str]) -> str:
    """The column key a header derives, disambiguated against the keys already `taken`."""
    base = slugify(label) or f"column-{index}"
    key = base
    suffix = 2
    while key in taken:
        key = f"{base}-{suffix}"
        suffix += 1
    taken.add(key)
    return key


def header_width(sheet: Worksheet) -> int:
    """Width of the contiguous run of header cells; a blank header ends the grid."""
    width = 0
    for index in range(1, sheet.max_column + 1):
        raw = sheet.cell(row=HEADER_ROW, column=index).value
        if raw is None or not str(raw).strip():
            break
        width = index
    return width


def build_columns(sheet: Worksheet) -> list[ColumnSpec]:
    """Column specs for the sheet's header run, in sheet order."""
    choices_by_column = validation_choices(sheet)
    taken: set[str] = set()
    columns: list[ColumnSpec] = []
    for index in range(1, header_width(sheet) + 1):
        raw = sheet.cell(row=HEADER_ROW, column=index).value
        label = str(raw).strip()
        choices = choices_by_column.get(index, [])
        dimension = sheet.column_dimensions.get(sheet.cell(row=HEADER_ROW, column=index).column_letter)
        columns.append(
            ColumnSpec(
                key=unique_key(label, index, taken),
                label=label,
                index=index,
                kind="choice" if choices else "text",
                role=_role_for(label),
                choices=choices,
                width=dimension.width if dimension is not None else None,
            )
        )
    return columns


def watch_column(columns: list[ColumnSpec]) -> ColumnSpec | None:
    """The column that records watch status, if the sheet has one."""
    return next((column for column in columns if column.role == "watch"), None)
