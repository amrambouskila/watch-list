"""Mutate a watch-order workbook in place, preserving every Excel construct it carries."""

from __future__ import annotations

import os
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import (
    DEFAULT_ACCENT_RGB,
    DEFAULT_WATCH_CHOICES,
    FIRST_DATA_ROW,
    HEADER_ROW,
    REPLACE_RETRIES,
    REPLACE_RETRY_SECONDS,
    TEMP_WRITE_PREFIX,
    TEMP_WRITE_SUFFIX,
    WATCH_COLUMN_WIDTH,
    WATCH_HEADER_LABEL,
)
from tv_watchlist.models.column_spec import ColumnSpec
from tv_watchlist.models.row_change import RowChange
from tv_watchlist.workbook import ranges
from tv_watchlist.workbook.cells import cell_text, last_data_row
from tv_watchlist.workbook.errors import (
    InvalidChoiceError,
    RowNotFoundError,
    UnknownColumnError,
    WorkbookLockedError,
)
from tv_watchlist.workbook.locking import excel_lock_path, is_locked_by_excel
from tv_watchlist.workbook.schema import build_columns, header_width, watch_column
from tv_watchlist.workbook.styling import apply_body_style, apply_header_style, clone_style


def _open(path: Path) -> tuple[Workbook, Worksheet]:
    if is_locked_by_excel(path):
        raise WorkbookLockedError(path.name)
    workbook = load_workbook(path)
    return workbook, workbook.worksheets[0]


def _install(temp: Path, path: Path) -> None:
    """
    Swap the rewritten file in.

    A concurrent read of the same workbook briefly holds a handle on Windows, so a first
    PermissionError is retried before concluding that Excel owns the file.
    """
    for attempt in range(REPLACE_RETRIES):
        try:
            os.replace(temp, path)
            return
        except PermissionError:
            if attempt == REPLACE_RETRIES - 1 or excel_lock_path(path).exists():
                raise
            time.sleep(REPLACE_RETRY_SECONDS)


def _save(workbook: Workbook, path: Path) -> None:
    """Write to a sibling temp file, then swap it in, so a failure never truncates the original."""
    temp = path.parent / f"{TEMP_WRITE_PREFIX}{os.getpid()}-{path.stem}{TEMP_WRITE_SUFFIX}"
    try:
        workbook.save(temp)
        _install(temp, path)
    except PermissionError as error:
        temp.unlink(missing_ok=True)
        raise WorkbookLockedError(path.name) from error
    except BaseException:
        temp.unlink(missing_ok=True)
        raise


def _is_number(value: Any) -> bool:  # noqa: ANN401 - openpyxl cells are genuinely heterogeneous
    return isinstance(value, int | float) and not isinstance(value, bool)


def _numeric_columns(sheet: Worksheet, columns: list[ColumnSpec], first: int, last: int) -> set[str]:
    """
    Column keys whose existing cells are all numbers.

    Typing has to be decided per column rather than per cell: a blank cell in the Order column
    carries no type of its own, and writing "17" into it as text would leave Excel showing a
    number-stored-as-text warning on a column that is otherwise numeric.
    """
    numeric: set[str] = set()
    for column in columns:
        seen = 0
        for row in range(first, last + 1):
            value = sheet.cell(row=row, column=column.index).value
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            if not _is_number(value):
                seen = 0
                break
            seen += 1
        if seen:
            numeric.add(column.key)
    return numeric


def _coerce(text: str, numeric: bool) -> Any:  # noqa: ANN401 - mirrors openpyxl's cell value type
    if text == "":
        return None
    if not numeric:
        return text
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() else number


def _validate(columns: list[ColumnSpec], values: dict[str, str]) -> list[tuple[ColumnSpec, str]]:
    by_key = {column.key: column for column in columns}
    resolved: list[tuple[ColumnSpec, str]] = []
    for key, value in values.items():
        column = by_key.get(key)
        if column is None:
            raise UnknownColumnError(key)
        if column.kind == "choice" and value not in column.choices:
            raise InvalidChoiceError(f"{column.label}: {value!r}")
        resolved.append((column, value))
    return resolved


def _write_cells(sheet: Worksheet, row: int, resolved: list[tuple[ColumnSpec, str]], numeric: set[str]) -> None:
    for column, value in resolved:
        sheet.cell(row=row, column=column.index).value = _coerce(value, column.key in numeric)


def _sequential_orders(sheet: Worksheet, column: ColumnSpec | None, first: int, last: int) -> bool:
    if column is None or last < first:
        return False
    expected = [str(number) for number in range(1, last - first + 2)]
    actual = [cell_text(sheet.cell(row=row, column=column.index).value).strip() for row in range(first, last + 1)]
    return actual == expected


def _resequence(sheet: Worksheet, column: ColumnSpec, first: int, last: int, numeric: bool) -> None:
    for offset, row in enumerate(range(first, last + 1), start=1):
        sheet.cell(row=row, column=column.index).value = _coerce(str(offset), numeric)


def _order_column(columns: list[ColumnSpec]) -> ColumnSpec | None:
    return next((column for column in columns if column.role == "order"), None)


def _grid_width(sheet: Worksheet, columns: list[ColumnSpec]) -> int:
    """The header run's width — never sheet.max_column, which counts stray cells to the right."""
    return max(header_width(sheet), max((column.index for column in columns), default=1))


def _take_hyperlinks(sheet: Worksheet, last_row: int, width: int) -> list[tuple[int, int, Any]]:
    """Detach every hyperlink so it can be re-anchored after rows move."""
    taken: list[tuple[int, int, Any]] = []
    for row in range(HEADER_ROW, last_row + 1):
        for column in range(1, width + 1):
            cell = sheet.cell(row=row, column=column)
            if cell.hyperlink is not None:
                taken.append((row, column, cell.hyperlink.target))
                cell.hyperlink = None
    return taken


def _restore_hyperlinks(sheet: Worksheet, taken: list[tuple[int, int, Any]], deleted_row: int) -> None:
    for row, column, target in taken:
        if row == deleted_row:
            continue
        sheet.cell(row=row - 1 if row > deleted_row else row, column=column).hyperlink = target


def _take_merges(sheet: Worksheet) -> list[CellRange]:
    """Unmerge everything so deleting a row shifts values instead of feeding them into a merge."""
    existing = [CellRange(str(cell_range)) for cell_range in sheet.merged_cells.ranges]
    for cell_range in existing:
        sheet.unmerge_cells(str(cell_range))
    return existing


def _restore_merges(sheet: Worksheet, merges: list[CellRange], deleted_row: int) -> None:
    for merge in merges:
        top = merge.min_row - 1 if merge.min_row > deleted_row else merge.min_row
        bottom = merge.max_row - 1 if merge.max_row >= deleted_row else merge.max_row
        if bottom < top or (top == bottom and merge.min_col == merge.max_col):
            continue
        sheet.merge_cells(start_row=top, start_column=merge.min_col, end_row=bottom, end_column=merge.max_col)


def update_row(path: Path, row: int, values: dict[str, str]) -> None:
    """Set one row's cells."""
    workbook, sheet = _open(path)
    columns = build_columns(sheet)
    resolved = _validate(columns, values)
    last = last_data_row(sheet, _grid_width(sheet, columns))
    if row < FIRST_DATA_ROW or row > last:
        raise RowNotFoundError(str(row))
    _write_cells(sheet, row, resolved, _numeric_columns(sheet, columns, FIRST_DATA_ROW, last))
    _save(workbook, path)


def append_row(path: Path, values: dict[str, str]) -> int:
    """Add a row at the end of the table and return its Excel row number."""
    workbook, sheet = _open(path)
    columns = build_columns(sheet)
    resolved = _validate(columns, values)
    width = _grid_width(sheet, columns)
    last = last_data_row(sheet, width)
    order = _order_column(columns)
    numeric = _numeric_columns(sheet, columns, FIRST_DATA_ROW, last)
    was_sequential = _sequential_orders(sheet, order, FIRST_DATA_ROW, last)

    target = max(last, HEADER_ROW) + 1
    empty_sheet = last < FIRST_DATA_ROW
    for index in range(1, width + 1):
        cell = sheet.cell(row=target, column=index)
        if empty_sheet:
            apply_body_style(cell)
        else:
            clone_style(sheet.cell(row=last, column=index), cell)
    _write_cells(sheet, target, resolved, numeric)
    if order is not None and all(column.key != order.key for column, _ in resolved):
        sheet.cell(row=target, column=order.index).value = _coerce(
            str(target - HEADER_ROW), order.key in numeric or empty_sheet
        )
    if was_sequential and order is not None:
        _resequence(sheet, order, FIRST_DATA_ROW, target, order.key in numeric)
    ranges.sync(sheet, target, width)
    _save(workbook, path)
    return target


def delete_row(path: Path, row: int) -> None:
    """Remove one row, closing the gap and moving every anchored construct with it."""
    workbook, sheet = _open(path)
    columns = build_columns(sheet)
    width = _grid_width(sheet, columns)
    last = last_data_row(sheet, width)
    if row < FIRST_DATA_ROW or row > last:
        raise RowNotFoundError(str(row))
    order = _order_column(columns)
    numeric = _numeric_columns(sheet, columns, FIRST_DATA_ROW, last)
    was_sequential = _sequential_orders(sheet, order, FIRST_DATA_ROW, last)

    hyperlinks = _take_hyperlinks(sheet, last, max(width, sheet.max_column))
    merges = _take_merges(sheet)
    sheet.delete_rows(row, 1)
    _restore_merges(sheet, merges, row)
    _restore_hyperlinks(sheet, hyperlinks, row)

    remaining = last - 1
    if was_sequential and order is not None and remaining >= FIRST_DATA_ROW:
        _resequence(sheet, order, FIRST_DATA_ROW, remaining, order.key in numeric)
    ranges.sync(sheet, max(remaining, FIRST_DATA_ROW), width)
    _save(workbook, path)


def _added_id(position: int) -> int:
    """Sequence id for the change at `position`; negative so it can never collide with an Excel row."""
    return -1 - position


def _place(values: list[Any], resolved: list[tuple[ColumnSpec, str]], numeric: set[str]) -> None:
    for column, value in resolved:
        values[column.index - 1] = _coerce(value, column.key in numeric)


def _reject_missing_rows(changes: list[RowChange], last: int) -> None:
    """Every row and anchor a batch names has to exist on the sheet as it stands now."""
    for change in changes:
        if change.row is not None and not FIRST_DATA_ROW <= change.row <= last:
            raise RowNotFoundError(str(change.row))
        if change.after_row is not None and not HEADER_ROW <= change.after_row <= last:
            raise RowNotFoundError(str(change.after_row))


def _anchored(anchor: int, attachments: dict[int, list[int]], removed: set[int], emitted: set[int]) -> list[int]:
    """Ids seated directly after `anchor`, then whatever is seated after those, in batch order."""
    seated: list[int] = []
    pending = list(attachments.get(anchor, ()))
    while pending:
        identifier = pending.pop(0)
        # Two moves can name each other as anchors; emitting each id once keeps that terminating.
        if identifier in emitted or identifier in removed:
            continue
        emitted.add(identifier)
        seated.append(identifier)
        pending = list(attachments.get(identifier, ())) + pending
    return seated


def _final_sequence(original_rows: list[int], changes: list[RowChange]) -> list[int]:
    """Sequence ids in their final order, every anchor read as an original row number."""
    removed = {change.row for change in changes if change.kind == "remove"}
    moved = {change.row for change in changes if change.kind == "move"}
    attachments: dict[int, list[int]] = {}
    trailing: list[int] = []
    for position, change in enumerate(changes):
        if change.kind == "move":
            attachments.setdefault(change.after_row, []).append(change.row)
        elif change.kind == "add":
            identifier = _added_id(position)
            if change.after_row is None:
                trailing.append(identifier)
            else:
                attachments.setdefault(change.after_row, []).append(identifier)

    emitted: set[int] = set()
    sequence = _anchored(HEADER_ROW, attachments, removed, emitted)
    for row in original_rows:
        if row not in removed and row not in moved and row not in emitted:
            emitted.add(row)
            sequence.append(row)
        sequence.extend(_anchored(row, attachments, removed, emitted))
    return sequence + trailing


def _planned_values(
    sheet: Worksheet,
    columns: list[ColumnSpec],
    changes: list[RowChange],
    numeric: set[str],
    width: int,
    rows: list[int],
) -> dict[int, list[Any]]:
    """Every row's final cell values keyed by sequence id, read before any of them move."""
    planned: dict[int, list[Any]] = {
        row: [sheet.cell(row=row, column=index).value for index in range(1, width + 1)] for row in rows
    }
    for position, change in enumerate(changes):
        if change.kind == "add":
            identifier = _added_id(position)
            planned[identifier] = [None] * width
        elif change.kind == "revise":
            identifier = change.row
        else:
            continue
        _place(planned[identifier], _validate(columns, change.cells), numeric)
    return planned


def _rewrite(sheet: Worksheet, planned: dict[int, list[Any]], sequence: list[int], width: int, last: int) -> int:
    """Lay the planned rows out from the first data row down, and return the new last data row."""
    final_last = FIRST_DATA_ROW + len(sequence) - 1
    empty_sheet = last < FIRST_DATA_ROW
    for row in range(last + 1, final_last + 1):
        for index in range(1, width + 1):
            cell = sheet.cell(row=row, column=index)
            if empty_sheet:
                apply_body_style(cell)
            else:
                clone_style(sheet.cell(row=last, column=index), cell)
    for offset, identifier in enumerate(sequence):
        values = planned[identifier]
        for index in range(1, width + 1):
            sheet.cell(row=FIRST_DATA_ROW + offset, column=index).value = values[index - 1]
    if final_last < last:
        sheet.delete_rows(final_last + 1, last - final_last)
    return final_last


def apply_changes(path: Path, changes: list[RowChange]) -> None:
    """Apply a whole batch in one rewrite, every change resolved against the original row numbers."""
    workbook, sheet = _open(path)
    columns = build_columns(sheet)
    width = _grid_width(sheet, columns)
    last = last_data_row(sheet, width)
    _reject_missing_rows(changes, last)

    order = _order_column(columns)
    numeric = _numeric_columns(sheet, columns, FIRST_DATA_ROW, last)
    was_sequential = _sequential_orders(sheet, order, FIRST_DATA_ROW, last)

    original_rows = list(range(FIRST_DATA_ROW, last + 1))
    planned = _planned_values(sheet, columns, changes, numeric, width, original_rows)
    sequence = _final_sequence(original_rows, changes)

    final_last = _rewrite(sheet, planned, sequence, width, last)
    if was_sequential and order is not None and final_last >= FIRST_DATA_ROW:
        _resequence(sheet, order, FIRST_DATA_ROW, final_last, order.key in numeric)
    ranges.sync(sheet, max(final_last, FIRST_DATA_ROW), width)
    _save(workbook, path)


def add_watch_column(path: Path) -> None:
    """Give a sheet the standard Watched? column, dropdown, and highlight rule."""
    workbook, sheet = _open(path)
    columns = build_columns(sheet)
    if watch_column(columns) is not None:
        return
    index = _grid_width(sheet, columns) + 1
    last = last_data_row(sheet, index - 1)

    header = sheet.cell(row=HEADER_ROW, column=index)
    header.value = WATCH_HEADER_LABEL
    if columns:
        clone_style(sheet.cell(row=HEADER_ROW, column=columns[0].index), header)
    else:
        apply_header_style(header, DEFAULT_ACCENT_RGB)
    for row in range(FIRST_DATA_ROW, last + 1):
        apply_body_style(sheet.cell(row=row, column=index))
    sheet.column_dimensions[get_column_letter(index)].width = WATCH_COLUMN_WIDTH

    ranges.sync(sheet, last, index)
    ranges.add_choice_validation(sheet, index, last, list(DEFAULT_WATCH_CHOICES))
    ranges.add_watched_highlight(sheet, index, last)
    _save(workbook, path)
