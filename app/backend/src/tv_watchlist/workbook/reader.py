"""Read a workbook into the API's category representation."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import (
    DEFAULT_ACCENT_RGB,
    FIRST_DATA_ROW,
    HEADER_ROW,
    IN_PROGRESS,
    SKIPPED,
    TRANSPARENT_ARGB,
    WATCHED,
)
from tv_watchlist.models.category_detail import CategoryDetail
from tv_watchlist.models.column_spec import ColumnSpec
from tv_watchlist.models.reference_sheet import ReferenceSheet
from tv_watchlist.models.watch_counts import WatchCounts
from tv_watchlist.models.watch_row import WatchRow
from tv_watchlist.workbook.cells import cell_text, last_data_row
from tv_watchlist.workbook.locking import is_locked_by_excel
from tv_watchlist.workbook.naming import category_id
from tv_watchlist.workbook.schema import build_columns, identity_column, position_column, public_columns, watch_column


def _accent(sheet: Worksheet) -> str:
    """Accent colour taken from the sheet's own header fill."""
    default = f"#{DEFAULT_ACCENT_RGB[2:]}"
    fill = sheet.cell(row=HEADER_ROW, column=1).fill
    if fill is None or fill.patternType is None:
        return default
    rgb = getattr(fill.fgColor, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) != 8 or rgb.upper() == TRANSPARENT_ARGB:
        return default
    return f"#{rgb[2:]}"


def _count(statuses: list[str]) -> WatchCounts:
    watched = sum(1 for status in statuses if status == WATCHED)
    in_progress = sum(1 for status in statuses if status == IN_PROGRESS)
    skipped = sum(1 for status in statuses if status == SKIPPED)
    total = len(statuses)
    return WatchCounts(
        total=total,
        watched=watched,
        in_progress=in_progress,
        skipped=skipped,
        unwatched=total - watched - in_progress - skipped,
        trackable=total - skipped,
    )


def _reference_sheets(sheets: list[Worksheet]) -> list[ReferenceSheet]:
    references: list[ReferenceSheet] = []
    for sheet in sheets:
        width = sheet.max_column
        header = [cell_text(sheet.cell(row=HEADER_ROW, column=column).value) for column in range(1, width + 1)]
        rows = [
            [cell_text(sheet.cell(row=row, column=column).value) for column in range(1, width + 1)]
            for row in range(FIRST_DATA_ROW, sheet.max_row + 1)
        ]
        references.append(ReferenceSheet(title=sheet.title, header=header, rows=[row for row in rows if any(row)]))
    return references


@dataclass(frozen=True, slots=True)
class WorkbookProjection:
    """One read of a workbook: the category the UI renders, plus the app-owned ids it never sees."""

    detail: CategoryDetail
    identity_key: str | None
    # Keyed by the rows `detail` carries, so a blank row the grid skips is not an entry to reconcile.
    row_ids: dict[int, str]
    watch_key: str | None
    # Used only to detect a raced workbook edit, never to reconstruct authoritative player progress.
    position_values: dict[int, str] = field(default_factory=dict)


def read_projection(path: Path) -> WorkbookProjection:
    """Load one workbook once, keeping the hidden identity cells that the category itself leaves out."""
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    columns: list[ColumnSpec] = build_columns(sheet)
    public = public_columns(columns)
    watch = watch_column(public)
    identity = identity_column(columns)
    position = position_column(columns)
    end_row = last_data_row(sheet, sheet.max_column)

    rows: list[WatchRow] = []
    statuses: list[str] = []
    row_ids: dict[int, str] = {}
    position_values: dict[int, str] = {}
    for row_number in range(FIRST_DATA_ROW, end_row + 1):
        cells = {column.key: cell_text(sheet.cell(row=row_number, column=column.index).value) for column in public}
        # Only the owner's own cells decide this: an id the app minted can never keep a blank row in the grid.
        if not any(value.strip() for value in cells.values()):
            continue
        rows.append(WatchRow(row=row_number, cells=cells))
        statuses.append(cells.get(watch.key, "").strip() if watch else "")
        if identity is not None:
            row_ids[row_number] = cell_text(sheet.cell(row=row_number, column=identity.index).value)
        if position is not None:
            position_values[row_number] = cell_text(sheet.cell(row=row_number, column=position.index).value)

    detail = CategoryDetail(
        id=category_id(path),
        name=path.stem,
        file_name=path.name,
        sheet_title=sheet.title,
        accent=_accent(sheet),
        mtime=path.stat().st_mtime,
        has_watch_column=watch is not None,
        locked_by_excel=is_locked_by_excel(path),
        counts=_count(statuses),
        columns=public,
        rows=rows,
        reference_sheets=_reference_sheets(workbook.worksheets[1:]),
    )
    return WorkbookProjection(
        detail=detail,
        identity_key=identity.key if identity is not None else None,
        row_ids=row_ids,
        watch_key=watch.key if watch is not None else None,
        position_values=position_values,
    )


def read_category(path: Path) -> CategoryDetail:
    """Load one workbook as a fully populated category."""
    return read_projection(path).detail
