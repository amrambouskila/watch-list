"""Conversion between Excel cell values and the plain strings the UI exchanges."""

from __future__ import annotations

import datetime as dt
from typing import Any

from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import FIRST_DATA_ROW


def cell_text(value: Any) -> str:  # noqa: ANN401 - openpyxl cells are genuinely heterogeneous
    """Render any Excel cell value as the string the UI edits."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time.min else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def table_last_row(sheet: Worksheet) -> int | None:
    """Last row covered by the sheet's named table, when it has one."""
    bounds = [range_boundaries(sheet.tables[name].ref) for name in sheet.tables]
    last_rows = [max_row for _, _, _, max_row in bounds if max_row is not None]
    return max(last_rows) if last_rows else None


def is_blank_row(sheet: Worksheet, row: int, column_count: int) -> bool:
    """True when every cell in the row is empty or whitespace."""
    return all(not cell_text(sheet.cell(row=row, column=column).value).strip() for column in range(1, column_count + 1))


def last_data_row(sheet: Worksheet, column_count: int) -> int:
    """Last row holding real data, ignoring trailing blank rows padded by Excel."""
    candidate = max(sheet.max_row, table_last_row(sheet) or 0)
    while candidate >= FIRST_DATA_ROW and is_blank_row(sheet, candidate, column_count):
        candidate -= 1
    return candidate
