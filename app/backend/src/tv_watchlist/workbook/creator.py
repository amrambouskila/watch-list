"""Generate a new category workbook shaped exactly like the ones already in the library."""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tv_watchlist.constants import (
    DEFAULT_COLUMN_WIDTH,
    DEFAULT_WATCH_CHOICES,
    FIRST_DATA_ROW,
    HEADER_ROW,
    NEW_CATEGORY_SHEET_TITLE,
    NEW_CATEGORY_TABLE_SUFFIX,
    NOTES_COLUMN_WIDTH,
    ORDER_COLUMN_WIDTH,
    TEMP_WRITE_PREFIX,
    TEMP_WRITE_SUFFIX,
    TITLE_COLUMN_WIDTH,
    WATCH_COLUMN_WIDTH,
    WATCH_HEADER_LABEL,
    WIDE_HEADER_KEYWORDS,
)
from tv_watchlist.models.category_create import CategoryCreate
from tv_watchlist.models.column_spec import ColumnSpec
from tv_watchlist.workbook import discovery, ranges
from tv_watchlist.workbook.errors import DuplicateCategoryError, UnknownColumnError
from tv_watchlist.workbook.naming import category_id, workbook_filename
from tv_watchlist.workbook.schema import build_columns, normalize_header
from tv_watchlist.workbook.styling import apply_body_style, apply_header_style

_NON_IDENTIFIER = re.compile(r"[^A-Za-z0-9]+")


def table_name(display: str) -> str:
    """Excel-legal table identifier derived from the category name."""
    core = _NON_IDENTIFIER.sub("", display) or "Category"
    if core[0].isdigit():
        core = f"_{core}"
    return f"{core}{NEW_CATEGORY_TABLE_SUFFIX}"


def _column_width(header: str) -> float:
    normalized = normalize_header(header)
    if normalized in {"order", "#"}:
        return ORDER_COLUMN_WIDTH
    if normalized == "title":
        return TITLE_COLUMN_WIDTH
    if normalized == "watched":
        return WATCH_COLUMN_WIDTH
    if any(keyword in normalized for keyword in WIDE_HEADER_KEYWORDS):
        return NOTES_COLUMN_WIDTH
    return DEFAULT_COLUMN_WIDTH


def _headers(requested: list[str]) -> list[str]:
    """Non-empty, distinct headers; Excel treats a table with a repeated column name as corrupt."""
    cleaned: list[str] = []
    seen: set[str] = set()
    for header in requested:
        label = header.strip()
        if not label:
            continue
        normalized = normalize_header(label)
        suffix = 2
        while normalized in seen:
            label = f"{header.strip()} {suffix}"
            normalized = normalize_header(label)
            suffix += 1
        seen.add(normalized)
        cleaned.append(label)
    if "watched" not in seen:
        cleaned.append(WATCH_HEADER_LABEL)
    return cleaned


def _write_headers(sheet: Worksheet, headers: list[str], accent_rgb: str) -> None:
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index)
        cell.value = header
        apply_header_style(cell, accent_rgb)
        sheet.column_dimensions[get_column_letter(index)].width = _column_width(header)


def _title_rows(columns: list[ColumnSpec], titles: list[str]) -> list[dict[str, str]]:
    """Seed titles restated as cells, so the grid writer has one input shape."""
    key = next((column.key for column in columns if normalize_header(column.label) == "title"), None)
    return [{} if key is None else {key: title} for title in (text.strip() for text in titles) if title]


def _write_rows(sheet: Worksheet, columns: list[ColumnSpec], rows: list[dict[str, str]]) -> int:
    """Fill the data region and return its last row; Order is numbered 1..N whatever was supplied."""
    index_by_key = {column.key: column.index for column in columns}
    order = next((column.index for column in columns if column.role == "order"), None)

    seeds = rows or [{}]
    for offset, cells in enumerate(seeds):
        row = FIRST_DATA_ROW + offset
        for column in columns:
            apply_body_style(sheet.cell(row=row, column=column.index))
        for key, text in cells.items():
            if key not in index_by_key:
                raise UnknownColumnError(key)
            if text:
                sheet.cell(row=row, column=index_by_key[key]).value = text
        if order is not None:
            sheet.cell(row=row, column=order).value = offset + 1
    return FIRST_DATA_ROW + len(seeds) - 1


def create_category(library_dir: Path, request: CategoryCreate) -> Path:
    """Write a new workbook for `request` into the library and return its path."""
    destination = library_dir / workbook_filename(request.name)
    # A free filename is not a free category: punctuation drops out of an id, so `Yu Gi Oh` would
    # answer to the same id as a `Yu-Gi-Oh!` already on disk and every later write would land there.
    shadowed = discovery.addressed_by(library_dir, category_id(destination))
    if shadowed is not None:
        raise DuplicateCategoryError(shadowed.name)

    headers = _headers(request.columns)
    accent_rgb = f"FF{request.accent.lstrip('#').upper()}"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = NEW_CATEGORY_SHEET_TITLE
    _write_headers(sheet, headers, accent_rgb)
    columns = build_columns(sheet)
    last_row = _write_rows(sheet, columns, request.rows or _title_rows(columns, request.titles))

    sheet.add_table(ranges.build_table(table_name(request.name), headers, last_row))
    watch_index = next(column.index for column in columns if column.role == "watch")
    ranges.add_choice_validation(sheet, watch_index, last_row, list(DEFAULT_WATCH_CHOICES))
    ranges.add_watched_highlight(sheet, watch_index, last_row)

    handle, name = tempfile.mkstemp(dir=library_dir, prefix=TEMP_WRITE_PREFIX, suffix=TEMP_WRITE_SUFFIX)
    os.close(handle)
    temp = Path(name)
    try:
        workbook.save(temp)
        os.replace(temp, destination)
    except BaseException:
        temp.unlink(missing_ok=True)
        raise
    return destination
